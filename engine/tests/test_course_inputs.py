"""Course Excel inputs, several materials per laminate, compliance results and per-ply answer keys.

Workbooks are built here in the course template's layout with constructed numbers; no
course file is stored in the repository.
"""

from pathlib import Path
import subprocess
import sys

import numpy as np
import pytest
from openpyxl import Workbook

from anisotropy_engine.cases import compare, engine_results, load_case, validate_case_inputs
from anisotropy_engine.lamina import Lamina, reduced_compliance, reduced_stiffness
from anisotropy_engine.laminate import Ply
from anisotropy_engine.response import abd_compliance, abd_matrix
from anisotropy_engine.rotation import transformed_compliance, transformed_reduced_stiffness

ENGINE = Path(__file__).parents[1]
HYBRID = ENGINE / "validation/cases/se142-hw4-p4ab-hybrid-woven-thermal.toml"
MATERIAL = Lamina(E1=140e9, E2=10e9, G12=7e9, nu12=0.3)
TAPE = [1, 27.7e6, 1.44e6, 1.13e6, 0.35, 0.005, -0.5e-6, 12.4e-6, 4.7e5, -2.3e5, 8.9e3, -2.9e4, 1.09e4]
WOVEN = [2, 11.2e6, 10.9e6, 0.94e6, 0.06, 0.010, 1.9e-6, 2.1e-6]
# (ply number, angle, material id, thickness); listed top-down as in the course files.
HYBRID_PLIES = [(11, 0, 2, 0.01), (10, 0, 1, 0.005), (9, 45, 1, 0.005), (8, 90, 1, 0.005), (7, -45, 1, 0.005),
                (6, 90, 1, 0.005), (5, -45, 1, 0.005), (4, 90, 1, 0.005), (3, 45, 1, 0.005), (2, 0, 1, 0.005),
                (1, 0, 2, 0.01)]


def course_workbook(path, materials=(TAPE, WOVEN), loads=(860, 0, -250, 0, 0, 0, -300), plies=HYBRID_PLIES,
                    units="lbf, in, psi, F", stress=None):
    """A workbook laid out like the course input template."""
    book = Workbook()
    sheet = book.active
    sheet["A1"], sheet["A2"], sheet["B2"] = "Input for a constructed test", "Laminate:", "test"
    sheet["A3"], sheet["B3"] = "Units:", units
    sheet["A5"] = "Material Properties:"
    for c, header in enumerate(["Matl ID", "E1", "E2", "G12", "nu12", "Ply Thk", "alpha1", "alpha2",
                                "X", "X'", "Y", "Y'", "S"], 1):
        sheet.cell(6, c, header)
    for r, row in enumerate(materials, 7):
        for c, value in enumerate(row, 1):
            sheet.cell(r, c, value)
    if stress is None:
        sheet["A10"] = "Loading:"
        for c, header in enumerate(["Nx", "Ny", "Nxy", "Mx", "My", "Mxy", "DT"], 1):
            sheet.cell(11, c, header)
            sheet.cell(12, c, loads[c - 1])
    else:
        sheet["A10"] = "Stress applied - global"
        for c, header in enumerate(["sigx", "sigy", "sigxy"], 1):
            sheet.cell(11, c, header)
            sheet.cell(12, c, stress[c - 1])
    sheet["A14"] = "Layup - bottom-most is ply 1 (i.e., most -z location)"
    for c, header in enumerate(["Ply No.", "Angle", "Matl ID", "Ply Thk"], 1):
        sheet.cell(15, c, header)
    for r, row in enumerate(plies, 16):
        for c, value in enumerate(row, 1):
            sheet.cell(r, c, value)
    book.save(path)
    return path


# ------------------------------------------------------------------ compliance results

@pytest.mark.parametrize("angle", [0, 20, 45, -30, 90])
def test_transformed_compliance_inverts_transformed_stiffness(angle):
    product = transformed_compliance(MATERIAL, angle) @ transformed_reduced_stiffness(MATERIAL, angle)
    np.testing.assert_allclose(product, np.eye(3), atol=1e-12)


def test_reduced_compliance_is_the_closed_form_inverse_of_q():
    np.testing.assert_allclose(reduced_compliance(MATERIAL) @ reduced_stiffness(MATERIAL), np.eye(3), atol=1e-12)
    np.testing.assert_allclose(reduced_compliance(MATERIAL)[0, :2], [1 / 140e9, -0.3 / 140e9], rtol=1e-12)


def test_abd_compliance_blocks_invert_the_abd_matrix():
    plies = [Ply(material=MATERIAL, angle_deg=a, thickness_m=0.2e-3) for a in (45, 90, 30, 0)]
    a, b, d = abd_compliance(plies)
    inverse = np.block([[a, b], [b.T, d]])
    np.testing.assert_allclose(inverse @ abd_matrix(plies), np.eye(6), atol=1e-9)
    symmetric = [Ply(material=MATERIAL, angle_deg=x, thickness_m=0.2e-3) for x in (0, 45, 45, 0)]
    np.testing.assert_allclose(abd_compliance(symmetric)[1], 0, atol=1e-15)  # 1/N; the unsymmetric stack above has |b| up to 1.2e-4


# ------------------------------------------------------------------ course Excel files

def test_course_workbook_becomes_a_bottom_first_multi_material_case(tmp_path):
    case = load_case(course_workbook(tmp_path / "hybrid.xlsx"))
    plies = case["laminate"]["plies"]
    assert case["units"] == "psi"
    assert [ply["material"] for ply in plies] == ["2"] + ["1"] * 9 + ["2"]
    assert [ply["angle"] for ply in plies] == [0, 0, 45, 90, -45, 90, -45, 90, 45, 0, 0]
    assert plies[0]["thickness"] == 0.01 and plies[1]["thickness"] == 0.005
    assert case["applied"] == {"loads": [860, 0, -250, 0, 0, 0], "delta_T": -300}
    # Compressive strengths are negative in the course file and positive magnitudes in a case.
    assert case["materials"]["1"]["F1c"] == 2.3e5 and case["materials"]["1"]["F2c"] == 2.9e4
    assert "F1t" not in case["materials"]["2"]


def test_course_workbook_and_equivalent_toml_case_agree(tmp_path):
    from_excel = engine_results(load_case(course_workbook(tmp_path / "hybrid.xlsx")))
    from_toml = engine_results(load_case(HYBRID))
    for key in ("A11", "A66", "N_T", "strain_global", "ply_stress_local"):
        np.testing.assert_allclose(from_excel[key][0], from_toml[key][0], rtol=1e-12)


def test_course_workbook_with_an_applied_stress(tmp_path):
    path = course_workbook(tmp_path / "stress.xlsx", materials=[TAPE], plies=[(1, 20, 1, 0.005)],
                           stress=(50000, -3500, 12000))
    case = load_case(path)
    assert case["applied"] == {"stress": [50000, -3500, 12000]}
    np.testing.assert_allclose(engine_results(case)["stress_local"][0][0], 51455.1, rtol=1e-5)


@pytest.mark.parametrize("changes, message", [
    ({"loads": (860, "TBD", 0, 0, 0, 0, None)}, "Load Ny is 'TBD'"),
    ({"loads": (860, 0, 0, 0, 0, 0, "cold")}, "DT is 'cold'"),
    ({"plies": [(2, 0, 1, 0.005), (2, 90, 1, 0.005)]}, "repeated \\[2\\]"),
    ({"plies": [(1, 0, 3, 0.005)]}, "material 3"),
    ({"plies": [(1, 0, 1, 0.007)]}, "differs"),
    ({"materials": [TAPE[:9] + [2.3e5] + TAPE[10:]]}, "X' must be entered as a negative number"),
    ({"materials": [TAPE[:9] + [None] * 4]}, "all five strengths"),
    ({"materials": [TAPE[:7]]}, "both alpha1 and alpha2"),
    ({"materials": [["(no material in this file)"]]}, "no numeric E1"),
    ({"units": "N, m, Pa, C"}, "Units cell"),
])
def test_course_workbook_errors_name_the_problem(tmp_path, changes, message):
    with pytest.raises(ValueError, match=message):
        load_case(course_workbook(tmp_path / "bad.xlsx", **changes))


def test_placeholder_material_rows_are_skipped_until_a_ply_uses_them(tmp_path):
    placeholder = [2, "XXYYZZ", "Can add", "more ", "materials"]
    case = load_case(course_workbook(tmp_path / "one.xlsx", materials=[TAPE, placeholder],
                                     plies=[(1, 0, 1, 0.005)]))
    assert list(case["materials"]) == ["1"]


# ------------------------------------------------------------------ several materials in TOML

def hybrid_case():
    return load_case(HYBRID)


@pytest.mark.parametrize("edit, message", [
    (lambda c: c["laminate"]["plies"][0].update(material="Kevlar"), "'Kevlar' is not in \\[materials\\]"),
    (lambda c: c["laminate"]["plies"][0].pop("thickness"), "missing fields \\['thickness'\\]"),
    (lambda c: c["laminate"]["plies"][0].update(thickness=0), "thickness must be a positive number"),
    (lambda c: c["laminate"].update(plies=[]), "nonempty list"),
    (lambda c: c["materials"]["AGP370"].pop("alpha1"), "both alpha1 and alpha2"),
    (lambda c: [c["materials"]["AGP370"].pop(k) for k in ("alpha1", "alpha2")], "needs alpha1 and alpha2 in materials \\['AGP370'\\]"),
    (lambda c: c.update(material={"E1": 1, "E2": 1, "G12": 1, "nu12": 0.1}), "exactly one of \\[material\\]"),
    (lambda c: c.pop("materials"), "exactly one of \\[material\\]"),
    (lambda c: c["expected"].pop("tolerance"), "positive relative tolerance"),
])
def test_invalid_multi_material_cases_are_rejected(edit, message):
    case = hybrid_case()
    edit(case)
    with pytest.raises(ValueError, match=message):
        validate_case_inputs(case)


def test_q_and_q_bar_results_are_only_given_where_they_are_unambiguous():
    results = engine_results(hybrid_case())
    assert "Q11" not in results and "Qbar11" not in results  # two materials
    single = load_case(ENGINE / "validation/cases/se142-hw2-p4a-qbar-sbar-20deg.toml")
    assert {"Q11", "S11", "Qbar16", "Sbar16"} <= engine_results(single).keys()


# ------------------------------------------------------------------ answer keys by ply number

def test_answer_keys_can_name_only_some_plies():
    comparisons, problems = compare(hybrid_case())
    assert not problems
    labels = [c.key for c in comparisons]
    assert "ply_stress_local[ply 9][0]" in labels and "ply_stress_local[ply 2][0]" not in labels
    assert all(c.agrees for c in comparisons)


@pytest.mark.parametrize("expected, message", [
    ({"ply_stress_local": {"12": [0, 0, 0]}}, "ply 12 does not exist; plies run from 1 \\(bottom\\) to 11"),
    ({"ply_stress_local": {"1": [0, 0]}}, "ply 1 needs exactly 3 values"),
    ({"strain_global": {"1": [0, 0, 0]}}, "only per-ply results can be given by ply number"),
])
def test_bad_per_ply_answer_keys_are_reported(expected, message):
    case = hybrid_case()
    case["expected"] = {"tolerance": 0.001, **expected}
    _, problems = compare(case)
    assert len(problems) == 1 and problems[0].split(": ", 1)[1].startswith(message.replace("\\", ""))


# ------------------------------------------------------------------ calculator

def calculator(*args):
    return subprocess.run([sys.executable, str(ENGINE / "examples/laminate_calculator.py"), *map(str, args)],
                          cwd=ENGINE, capture_output=True, text=True)


def test_calculator_reads_a_course_workbook_and_names_each_material(tmp_path):
    run = calculator(course_workbook(tmp_path / "hybrid.xlsx"))
    assert run.returncode == 0, run.stderr
    assert "Q, fiber axes, material 2" in run.stdout
    assert "ply 11      0 deg   material 2" in run.stdout
    assert "Temperature change delta_T = -300 degF" in run.stdout
    assert "Not every material has strengths, so failure checks are skipped." in run.stdout


def test_calculator_ends_with_the_answer_key_comparison():
    run = calculator(HYBRID)
    assert run.returncode == 0, run.stderr
    assert "Answer key vs engine" in run.stdout
    assert run.stdout.rstrip().endswith("52 of 52 values agree with the answer key.")


def test_calculator_reports_a_course_workbook_error_without_a_traceback(tmp_path):
    run = calculator(course_workbook(tmp_path / "bad.xlsx", loads=(860, "TBD", 0, 0, 0, 0, None)))
    assert run.returncode != 0
    assert "Load Ny is 'TBD'" in run.stderr and "Traceback" not in run.stderr
