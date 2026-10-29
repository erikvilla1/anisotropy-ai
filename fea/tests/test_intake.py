"""The coupon input form (ADR-0046): the blank form writes and reads back, a filled-in form on the public stand-in gives the stand-in's
spec, an element size seeds the tab lines, every refusal names its row, the open-hole block is read, and a coupon's name follows
the rules of the coupon tests. Nothing here touches .local/."""

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from anisotropy_fea import paths
from anisotropy_fea.forms import labelled
from anisotropy_fea.case_pages import read_case
from anisotropy_fea.intake import (DEFINITIONS, FIELDS, FORM_VERSION, GLOSSARY, HOWTO, INPUT, LIBRARY, LIBRARY_HELP, RECORD, IntakeError, build_specs,
                                   check_form, convert, designation_for, fill_form, library_records, read_form, render_toml, spec_data,
                                   write_blank_form)
from anisotropy_fea.femodel import build_model
from anisotropy_fea.spec import CouponSpec, coupon_card, load_spec, node_lines, parse_spec, row_coordinates, spec_material, to_si
from anisotropy_fea.write_abaqus import write_abaqus


def load_spec_text(text: str):
    return to_si(parse_spec(text))
from anisotropy_engine.library import library_material

from public_coupon import DESIGNATION, PUBLIC_COUPON

PUBLIC_FILL = dict(filled_by="A. Engineer", date="2026-10-25", test="unnotched tension", material_code="IM7", layup_code="L1", standard="none",
                   size_from="the public IM7/8552 check fea/checks/unnotched-im7-8552/", team_model="not built yet", test_data="none",
                   length=250.0, width=25.0, tab_length=56.0, angles=[45, 0, -45, 90], symmetric="yes", library="im7-8552", G13=4688.43, G23=3000.0,
                   control="displacement", value=1.0, element_size=None, mesh_nx=10, mesh_partitions="56, 110, 140, 194", mesh_seeds="28, 27, 10, 27, 28",
                   name_part="Strip", name_instance="Strip-1", name_material="IM7-8552-Lamina", name_layup="Layup-A", name_section="Layup-A-1")
OHT_FILL = dict(filled_by="A. Engineer", date="2026-10-25", test="open-hole tension", material_code="IM7", layup_code="L1",
                size_from="the proportions of ASTM D5766", length=300.0, width=36.0, tab_length=50.0, hole_diameter=6.0, angles=[45, 0, -45, 90],
                symmetric="yes", library="im7-8552", G13=4688.43, G23=3000.0, control="displacement", value=1.0, element_size=2.0)


@pytest.fixture(scope="module", autouse=True)
def local_untouched():
    before = sorted(p.name for p in paths.LOCAL.iterdir()) if paths.LOCAL.is_dir() else None
    yield
    after = sorted(p.name for p in paths.LOCAL.iterdir()) if paths.LOCAL.is_dir() else None
    assert before == after, "the intake tests must not write under .local/"


@pytest.fixture(scope="module")
def blank(tmp_path_factory) -> Path:
    return write_blank_form(tmp_path_factory.mktemp("forms") / "coupon-input-form.xlsx")


def filled(blank: Path, folder: Path, values: dict, **changes) -> Path:
    return fill_form(blank, {**values, **changes}, folder / "filled.xlsx")


def specs_of(blank: Path, folder: Path, values: dict, root: Path | None = None, **changes) -> tuple[str, str]:
    form = read_form(filled(blank, folder, values, **changes))
    checked = check_form(form)
    return build_specs(checked, designation_for(checked, root or folder / "no-coupons"), form.path.name)


# --------------------------------------------------------------------------------------------- the blank form

def test_the_blank_form_has_its_sheets_rows_lists_and_record(blank):
    book = load_workbook(blank)
    assert book.sheetnames == [INPUT, HOWTO, DEFINITIONS, LIBRARY, RECORD] and book[RECORD].sheet_state == "hidden"
    assert labelled(book[RECORD])["Form version"] == FORM_VERSION and labelled(book[RECORD])["Library materials listed"] == str(len(library_records()))
    sheet = book[INPUT]
    labels = [sheet.cell(r, 1).value for r in range(1, sheet.max_row + 1) if sheet.cell(r, 1).value in {f.label for f in FIELDS.values()}]
    assert labels == [f.label for f in FIELDS.values()]  # every row once, in the order of the blocks
    rows = {sheet.cell(r, 1).value: r for r in range(1, sheet.max_row + 1)}
    for f in FIELDS.values():
        assert not sheet.cell(rows[f.label], 2).protection.locked and sheet.cell(rows[f.label], 1).protection.locked, f.label
        assert sheet.cell(rows[f.label], 2).value == f.default, f.label
    assert sheet.protection.sheet and sheet.cell(rows["Length (mm)"], 3).value.startswith("required")
    lists = {str(v.sqref): v.formula1 for v in sheet.data_validations.dataValidation if v.type == "list"}
    assert lists[f"B{rows['Library material']}"] == f"'{LIBRARY}'!$A$2:$A${len(library_records()) + 1}"
    assert lists[f"B{rows['Test']}"] == '"unnotched tension,unnotched compression,open-hole tension,in-plane shear,three-point bending,four-point bending,plate"'
    assert str(sheet.cell(rows["Library material"], 3).value).startswith("=IFERROR(INDEX(")
    assert str(sheet.cell(rows["Library material"], 3).value).endswith(f',"{LIBRARY_HELP}")')  # the instruction until a material is picked
    library = book[LIBRARY]
    assert [library.cell(r, 1).value for r in range(2, library.max_row + 1)] == [m.material_id for m in library_records()]
    assert library.cell(1, 1).value == "Id" and library.auto_filter.ref and library.protection.sheet
    for name in (INPUT, HOWTO, DEFINITIONS, LIBRARY):
        for cells in book[name].iter_rows():
            for cell in cells:
                if isinstance(cell.value, str):
                    assert "kernel" not in cell.value.lower() and "ADR" not in cell.value, (name, cell.coordinate)


def test_every_row_says_what_it_means(blank):
    sheet = load_workbook(blank)[INPUT]
    rows = {sheet.cell(r, 1).value: r for r in range(1, sheet.max_row + 1)}
    later_plies = {f"Ply {i}" for i in range(3, 17)}  # one ply per row: the first two rows say how, the rest stay quiet
    for f in FIELDS.values():
        if f.label not in later_plies:
            assert sheet.cell(rows[f.label], 3).value, f"the row {f.label!r} has no explanation beside it"
    ply = sheet.cell(rows["Ply 1 (bottom)"], 3).value
    assert "measured from x" in ply and "90 = fibres along the pull" in ply  # the direction an engineer asks about first


def test_the_definitions_answer_the_first_questions(blank):
    sheet = load_workbook(blank)[DEFINITIONS]
    terms = {sheet.cell(r, 1).value: (sheet.cell(r, 2).value, sheet.cell(r, 3).value) for r in range(5, sheet.max_row + 1)}
    for group, entries in GLOSSARY:
        assert group in terms, group
        for term, meaning, _ in entries:
            assert terms[term][0] == meaning and meaning, term
    for asked in ("Unnotched tension", "Open-hole tension", "In-plane shear", "Ply angle", "Symmetric", "Element size", "Finer mesh",
                  "Your Abaqus mesh", "G13, G23", "Library material", "Empty cell on the sheet Library materials", "Coupling grip",
                  "Where we read results", "First-ply failure"):
        assert asked in terms, asked
    assert "G23 = E2 / 3" in terms["G13, G23"][1] and "MIL-HDBK-17-3F" in terms["G13, G23"][1]
    assert sheet.protection.sheet and sheet.freeze_panes == "A5"


def test_a_form_of_version_0_1_still_reads(blank, tmp_path):
    path = filled(blank, tmp_path, PUBLIC_FILL)
    book = load_workbook(path)
    record = book[RECORD]
    row = next(r for r in range(1, record.max_row + 1) if record.cell(r, 1).value == "Form version")
    record.cell(row, 2, "0.1")  # the form the team may still hold: same labels, older explanations
    book.save(path)
    form = read_form(path)
    assert form.version == "0.1" and check_form(form).library == "im7-8552"


# --------------------------------------------------------------------------------------------- the public stand-in, round trip

def test_the_public_coupon_filled_in_gives_its_spec(blank, tmp_path):
    main, twin = specs_of(blank, tmp_path, PUBLIC_FILL)
    ours, public = parse_spec(main), parse_spec(PUBLIC_COUPON)
    assert ours.coupon == public.coupon and ours.coupon.designation == DESIGNATION
    assert ours.title.startswith(f"{DESIGNATION}, unnotched tension: ")
    assert ours.geometry == public.geometry and ours.mesh == public.mesh and ours.layup == public.layup
    assert (ours.material.library, ours.material.dataset) == (public.material.library, public.material.dataset)
    assert (ours.material.transverse.G13, ours.material.transverse.G23) == (public.material.transverse.G13, public.material.transverse.G23)
    assert ours.grips == public.grips and ours.loading == public.loading and ours.names == public.names and ours.probes == public.probes
    # our defaults, stated on the form's instructions
    assert ours.units == "mm-MPa" and ours.mesh.element_order == 1 and ours.material.transverse.plane_stress and ours.material.transverse.transverse_isotropy
    assert ours.output.fpf_criteria == ["max_stress", "tsai_hill"] and ours.material.f12_star is None
    # the CalculiX twin differs only in its grips and its title
    other = parse_spec(twin)
    skip = {"title", "grips", "source_sha256", "source_path"}
    assert other.grips.kind == "nodal" and "nodal grips" in other.title
    assert other.model_dump(exclude=skip) == ours.model_dump(exclude=skip)
    assert "# IM7-UNT-L1, unnotched tension: built from the coupon input form filled.xlsx, filled in by A. Engineer on 2026-10-25." in main


def test_the_writer_and_the_parser_agree(blank, tmp_path):
    form = read_form(filled(blank, tmp_path, PUBLIC_FILL))
    values = check_form(form)
    data, comments, header = spec_data(values, designation_for(values, tmp_path / "none"), values.grips, form.path.name)
    rendered = render_toml(data, comments, header)
    skip = {"source_sha256", "source_path"}
    assert parse_spec(rendered).model_dump(exclude=skip) == CouponSpec(**data).model_dump(exclude=skip)
    assert rendered.count("[material.transverse]") == 1 and "abaqus = { field_preselect = true" in rendered


# --------------------------------------------------------------------------------------------- an element size seeds the tab lines

@pytest.mark.parametrize("size, seeds", [(2.0, [28, 69, 28]), (3.0, [19, 46, 19])])
def test_an_element_size_puts_a_node_row_on_each_tab_line(blank, tmp_path, size, seeds):
    main, _ = specs_of(blank, tmp_path, PUBLIC_FILL, element_size=size, mesh_nx=None, mesh_partitions=None, mesh_seeds=None)
    spec = parse_spec(main)
    assert spec.mesh.y_partitions == [56.0, 194.0] and spec.mesh.y_seeds == seeds and spec.mesh.nx == max(1, round(25.0 / size))
    rows = row_coordinates(spec)
    for y in node_lines(spec).values():
        assert any(abs(y - row) <= 1e-9 * 250.0 for row in rows)


def test_the_hole_rows_are_ignored_for_the_other_tests(blank, tmp_path):
    spec = parse_spec(specs_of(blank, tmp_path, PUBLIC_FILL, hole_diameter=6.0)[0])
    assert spec.kind == "unnotched_tension" and spec.geometry.hole_diameter is None and spec.mesh.hole is None


def test_without_tabs_the_mesh_is_uniform(blank, tmp_path):
    main, _ = specs_of(blank, tmp_path, PUBLIC_FILL, tab_length=0.0, region="line", element_size=2.0, mesh_nx=None, mesh_partitions=None, mesh_seeds=None)
    spec = parse_spec(main)
    assert spec.mesh.y_seeds is None and spec.mesh.ny == 125 and spec.grips.region == "line"


# --------------------------------------------------------------------------------------------- refusals name their rows

def own_constants() -> dict:
    lib = library_material("im7-8552")
    s = lib.lamina.strengths
    return dict(library=None, own_name="IM7-8552-Lamina", own_E1=lib.lamina.E1 / 1e6, own_E2=lib.lamina.E2 / 1e6, own_G12=lib.lamina.G12 / 1e6,
                own_nu12=lib.lamina.nu12, own_F1t=s.F1t / 1e6, own_F1c=s.F1c / 1e6, own_F2t=s.F2t / 1e6, own_F2c=s.F2c / 1e6, own_F6=s.F6 / 1e6,
                own_source="the public NCAMP report, typed in by the team", ply_thickness=0.18288)


@pytest.mark.parametrize("changes, message", [
    (dict(length=None), "Length (mm): required"),
    (dict(width="wide"), "Width (mm): a number is needed, not 'wide'"),
    (dict(own_E1=100.0), "Library material and Own material: give one of the two, not both"),
    (dict(library=None), "Library material or Own material: one of the two is needed"),
    ({**own_constants(), "own_F6": None}, "Own material strengths: give all five (F1t, F1c, F2t, F2c, F6) or none; given F1t, F1c, F2t, F2c"),
    (dict(angles=[45, None, -45, 90]), "Ply 3: a blank ply row above it"),
    (dict(mesh_seeds="28, 27"), "rows per segment: 2 counts for 4 partition lines"),
    (dict(element_size=2.0), "Element size and Your Abaqus mesh: give one of the two, not both"),
    (dict(G13=0), "G13 (MPa): must be greater than 0, not 0"),
    (dict(test="open-hole tension"), "Hole diameter (mm): required for open-hole tension"),
    (dict(test="open-hole tension", hole_diameter=6.0, hole_circumferential=30), "mesh.hole.circumferential must be a multiple of 4, at least 8 (form rows: Nodes around the hole)"),
    (dict(value=-1.0), "Value (mm or N): must be greater than 0, not -1"),
    (dict(tab_length=0.0), 'grips.region = "tabs" needs a positive geometry.tab_length (form rows: Grip region, Tab length (mm))'),
    (dict(name_part="1bad"), "Abaqus name: part: '1bad' must start with a letter"),
    (dict(test="bending"), "Test: 'bending' is not one of unnotched tension, unnotched compression, open-hole tension, in-plane shear"),
    (dict(material_code="im7 prepreg"), "Material code: 'IM7 PREPREG' is not capital letters and digits"),
    (dict(date="last Tuesday"), "Date: 'last Tuesday' is not a date"),
    (dict(length=25.0, width=250.0), "Length (mm) and Tab length (mm): the two tabs, 112 mm, must be shorter than the length, 25 mm"),
])
def test_a_refusal_names_its_row(blank, tmp_path, changes, message):
    with pytest.raises(IntakeError, match=re.escape(message)):
        specs_of(blank, tmp_path, PUBLIC_FILL, **changes)


def test_several_problems_are_named_at_once(blank, tmp_path):
    with pytest.raises(IntakeError) as refused:
        specs_of(blank, tmp_path, PUBLIC_FILL, length=None, width="wide", symmetric=None, G23=None)
    text = str(refused.value)
    assert all(part in text for part in ("Length (mm): required", "Width (mm): a number is needed", "Symmetric: required; one of yes, no", "G23 (MPa): required"))


def test_a_file_that_is_not_the_form_is_refused(tmp_path):
    (tmp_path / "notes.xlsx").write_text("not a workbook", encoding="utf-8")
    with pytest.raises(IntakeError, match="not a coupon input form"):
        read_form(tmp_path / "notes.xlsx")
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = INPUT
    book.create_sheet(RECORD)
    book.save(tmp_path / "other.xlsx")
    with pytest.raises(IntakeError, match="no form version"):
        read_form(tmp_path / "other.xlsx")


# --------------------------------------------------------------------------------------------- the open-hole block and own constants

def test_an_open_hole_form_meshes_the_hole(blank, tmp_path):
    main, _ = specs_of(blank, tmp_path, OHT_FILL)
    spec = parse_spec(main)
    assert spec.kind == "open_hole_tension" and spec.mesh.mesher == "gmsh" and spec.mesh.element_size == 2.0 and spec.geometry.hole_diameter == 6.0
    assert spec.mesh.hole.model_dump() == {"circumferential": 32, "radial_layers": 6, "ring_ratio": 2.0, "growth": 1.2}
    assert spec.coupon.designation == "IM7-OHT-L1" and "6 mm hole in 36 mm" in spec.title and spec.names is None
    assert set(spec.probes.nodes) == {"centre_top", "hole_side", "gauge"} and set(spec.probes.elements) == {"hole_side", "hole_top", "far"}


def test_own_constants_make_a_spec_without_a_library_record(blank, tmp_path):
    main, _ = specs_of(blank, tmp_path, PUBLIC_FILL, **own_constants())
    spec = parse_spec(main)
    assert spec.material.library is None and spec.material.E1 == pytest.approx(library_material("im7-8552").lamina.E1 / 1e6)
    assert spec.material.F6 is not None and spec.output.fpf_criteria == ["max_stress", "tsai_hill"] and spec.layup.ply_thickness == 0.18288
    assert spec.material.source == "the public NCAMP report, typed in by the team" and spec.names.material == "IM7-8552-Lamina"
    without = parse_spec(specs_of(blank, tmp_path, PUBLIC_FILL, **{**own_constants(), **{k: None for k in ("own_F1t", "own_F1c", "own_F2t", "own_F2c", "own_F6")}})[0])
    assert without.material.F1t is None and without.output.fpf_criteria == []
    # without the Abaqus names, the deck still names the material as the team named it, with our other names
    unnamed = load_spec_text(specs_of(blank, tmp_path, PUBLIC_FILL, **own_constants(), **{k: None for k in ("name_part", "name_instance", "name_material",
                                                                                                         "name_layup", "name_section")})[0])
    assert unnamed.names.material == "IM7-8552-Lamina" and unnamed.names.part == "Coupon"
    material = spec_material(unnamed)
    assert material.rights == "user-supplied" and material.name == "IM7-8552-Lamina"
    heading = write_abaqus(build_model(unnamed))
    assert "the team's own constants: the public NCAMP report, typed in by the team; rights class user-supplied" in heading
    assert dict(coupon_card(unnamed))["Material"] == "IM7-8552-Lamina (the team's own constants; source: the public NCAMP report, typed in by the team)"


# --------------------------------------------------------------------------------------------- the coupon's name

def seeded_root(blank: Path, folder: Path) -> Path:
    root = folder / "root"
    main, _ = specs_of(blank, folder / "seed", PUBLIC_FILL)
    (root / DESIGNATION / "ours").mkdir(parents=True)
    (root / DESIGNATION / "ours" / f"{DESIGNATION}.toml").write_text(main, encoding="utf-8", newline="\n")
    return root


def test_a_code_is_reused_proposed_or_refused_by_what_the_machine_already_has(blank, tmp_path):
    root = seeded_root(blank, tmp_path)
    name = lambda **changes: parse_spec(specs_of(blank, tmp_path / "a", PUBLIC_FILL, root, **changes)[0]).coupon.designation  # noqa: E731
    assert name(material_code=None, layup_code=None) == "IM7-UNT-L1"  # the same material and stacking sequence: both codes reused
    assert name(material_code=None, layup_code=None, angles=[0, 90]) == "IM7-UNT-L2"  # a new stacking sequence: the next free code
    assert name(material_code=None, layup_code=None, angles=[0, 90], test="unnotched compression", value=0.1) == "IM7-UNC-L2"
    with pytest.raises(IntakeError, match="Material code: this material has no code yet; give one"):
        name(material_code=None, library="as4-8552", ply_thickness=0.18)
    with pytest.raises(IntakeError, match="Material code: IM7 already names another material on this machine"):
        name(material_code="IM7", library="as4-8552", ply_thickness=0.18)
    with pytest.raises(IntakeError, match="Material code: this material already has the code IM7; use it"):
        name(material_code="HEXCEL")
    with pytest.raises(IntakeError, match="Layup code: L1 already names another stacking sequence"):
        name(layup_code="L1", angles=[0, 90])
    with pytest.raises(IntakeError, match="Layup code: this stacking sequence already has the code L1; use it"):
        name(layup_code="L3")


# --------------------------------------------------------------------------------------------- filing a form

def test_a_filled_form_is_filed_and_its_coupon_started(blank, tmp_path):
    root = tmp_path / "root"
    form = filled(blank, tmp_path, PUBLIC_FILL)
    intake = convert(form, root=root, day="2026-10-25")
    folder = root / DESIGNATION
    assert intake.folder == folder and sorted(p.relative_to(folder).as_posix() for p in intake.written) == [
        "README.md", f"ours/{DESIGNATION}-nodal.toml", f"ours/{DESIGNATION}.toml", "team/2026-10-25/README.md", "team/2026-10-25/filled.xlsx"]
    assert (folder / "team/2026-10-25/filled.xlsx").read_bytes() == form.read_bytes()
    readme = (folder / "team/2026-10-25/README.md").read_text(encoding="utf-8")
    assert "filled.xlsx" in readme and "A. Engineer on 2026-10-25" in readme and "library record im7-8552" in readme and "rights class open-publisher" in readme
    case = read_case(folder)  # the skeleton is a case page the tools read
    assert case.status == "in progress" and case.attempts == 1 and case.accepted == "01" and case.category == "team test"
    page = (folder / "README.md").read_text(encoding="utf-8")
    assert page.startswith(f"# {DESIGNATION}, unnotched tension (`{DESIGNATION}`)\n") and "<!-- generated: result -->" in page and "## Rebuild" in page
    assert "250 mm long, 25 mm wide, 56 mm tabs at each end; layup [45/0/-45/90]s; material IM7 / 8552" in page
    assert intake.commands[0] == f"uv run python tools/regenerate_decks.py {DESIGNATION} {DESIGNATION}-nodal --ccx" and intake.commands[0] in page
    assert load_spec(folder / "ours" / f"{DESIGNATION}.toml").coupon.designation == DESIGNATION
    # a second reading refuses to overwrite the specs unless asked; then it rewrites the two specs and nothing else
    with pytest.raises(IntakeError, match="already exist"):
        convert(form, root=root, day="2026-10-26")
    (folder / "README.md").write_text(page + "\nHand-written.\n", encoding="utf-8")
    again = convert(form, root=root, day="2026-10-26", replace=True)  # the two specs and the input they came from; no README, no page
    assert sorted(p.relative_to(folder).as_posix() for p in again.written) == [f"ours/{DESIGNATION}-nodal.toml", f"ours/{DESIGNATION}.toml",
                                                                                "team/2026-10-26/filled.xlsx"]
    assert "Hand-written." in (folder / "README.md").read_text(encoding="utf-8") and not (folder / "team/2026-10-26/README.md").exists()
    # a dry run writes nothing
    dry = convert(form, root=tmp_path / "dry", day="2026-10-25", dry_run=True)
    assert dry.written == () and not (tmp_path / "dry").exists() and dry.main == intake.main and dry.twin == intake.twin


def test_a_form_already_in_the_team_folder_is_not_copied_onto_itself(blank, tmp_path):
    root = tmp_path / "root"
    received = root / DESIGNATION / "team" / "2026-10-25"
    received.mkdir(parents=True)
    form = fill_form(blank, PUBLIC_FILL, received / "coupon-input-form.xlsx")
    intake = convert(form, root=root, day="2026-10-25")
    assert form.read_bytes() and sorted(p.name for p in intake.written) == [f"{DESIGNATION}-nodal.toml", f"{DESIGNATION}.toml", "README.md", "README.md"]


def test_the_deck_tool_skips_the_calculix_deck_of_a_coupling_spec():
    """`regenerate_decks.py <coupon> <coupon>-nodal --ccx`, the documented step, must not stop at the coupling spec. The tool is run
    in-process with a writer that writes nothing, so the committed decks stay as they are."""
    import importlib.util

    spec = importlib.util.spec_from_file_location("regenerate_decks", paths.FEA / "tools" / "regenerate_decks.py")
    tool = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(tool)
    written = []
    lines = tool.regenerate(paths.spec_path("unnotched-im7-8552-coupling"), ccx=True, writer=lambda path, text: written.append(path) or "unchanged")
    assert [line.split()[0] for line in lines] == ["unchanged", "skipped"] and "coupled reference node" in lines[1]
    assert [p.name for p in written] == ["unnotched-im7-8552-coupling-abaqus.inp"]
