"""The coupon document and its package folder, docs/models/coupon.py (ADR-0041, ADR-0043), built in a temporary folder: the public
stand-in as a coupon the team has not built, and a fixture coupon with a team comparison. The coupon models themselves are local
(ADR-0042): the command line is run on one when this machine has it."""

import importlib.util
import subprocess
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from public_coupon import DESIGNATION, PUBLIC_COUPON
from test_equivalence import FIXTURE_SPEC, GRIPS_FIXTURE
from test_side_by_side import kernel_reports

from anisotropy_fea import paths  # where the files live (ADR-0036)
from anisotropy_fea.audit import COVER, DIGEST, RUBRIC_COLUMNS, RUBRIC_SHEET, Row, results_digest, verdict_of
from anisotropy_fea.audit import RECORD as RECORD_SHEET  # the form's Record sheet; RECORD below is the CalculiX record
from anisotropy_fea.compare import load_record
from anisotropy_fea.femodel import build_model
from anisotropy_fea.spec import load_spec
from anisotropy_fea.write_abaqus import write_abaqus

ROOT = Path(__file__).parents[2]
COUPON = ROOT / "docs/models/coupon.py"
TEAM = ROOT / "docs/models/team.py"
RECORD = ROOT / "fea/checks/unnotched-im7-8552/records/unnotched-im7-8552-nodal-calculix.toml"
CARD = """[coupon]
designation = "IM7-UNT-L9"
standard = "none"
size_from = "the grips fixture"
team_model = "the grips fixture deck"
test_data = "none"
"""
LIBRARY_SPEC = FIXTURE_SPEC.replace("E1 = 162095.74\nE2 = 8963.18\nG12 = 4688.43\nnu12 = 0.316\nF1t = 2558.51\nF1c = 1731.48\nF2t = 64.05\nF2c = 285.72\nF6 = 91.15\n",
                                    'library = "im7-8552"\ndataset = "ncamp-im7-8552-2019-rtd-measured-v1"\n')
NAMED_SPEC = LIBRARY_SPEC.replace('title = "The grips fixture as a spec"', 'title = "IM7-UNT-L9, unnotched tension: the grips fixture as a spec"', 1) \
    .replace("[geometry]", CARD + "[geometry]", 1)
CHAPTERS = ("The problem", "The physics: how our engine solves it", "Building the part in our engine", "Building the part in Abaqus/CAE",
            "Results from our engine", "Verification against other solvers", "What this shows and what it does not", "The audit")
TODAY = date.today().isoformat()


@pytest.fixture(scope="module")
def coupon():
    spec = importlib.util.spec_from_file_location("anisotropy_models_coupon", COUPON)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def team_sources(coupon, tmp: Path, spec_text: str = NAMED_SPEC):
    """A fixture coupon `grips` with a team comparison: the team's 2 by 8 deck and kernel-made field reports, our spec, and a
    committed CalculiX record standing in for the twin's."""
    (tmp / "team").mkdir(exist_ok=True)
    (tmp / "ours").mkdir(exist_ok=True)
    deck = tmp / "team" / GRIPS_FIXTURE.name
    deck.write_bytes(GRIPS_FIXTURE.read_bytes())
    reports = kernel_reports(GRIPS_FIXTURE, tmp / "team")
    spec = tmp / "ours" / "grips-coupon.toml"
    spec.write_text(spec_text, encoding="utf-8")
    manifest = tmp / "team-comparison.toml"
    manifest.write_text("\n".join([
        'name = "grips"', 'units = "mm-MPa"', "node_tolerance = 1e-5", "agreement = 0.01", "edge_band = 10.0", "map_layer = 10",
        "[locations]", '"centre" = [6.25, 56.25]',  # an element centroid of the 25 by 100 mm fixture
        "[team]", f'deck = "team/{deck.name}"', 'solver = "a fixture"', 'run_by = "the test"', 'run_date = "2026-10-20"',
        "reports = [" + ", ".join(f'"team/{p.name}"' for p in reports) + "]",
        "[ours]", 'spec = "ours/grips-coupon.toml"', f'calculix_record = "{RECORD.as_posix()}"',
    ]) + "\n", encoding="utf-8")
    twin = (RECORD.parent / load_record(RECORD)["spec"]).resolve()
    return coupon.Sources(coupon="grips", spec=spec, twin=twin, record=RECORD, manifest=manifest, documents=tmp / paths.DOCUMENTS,
                          team_files=tmp / paths.TEAM_FILES, latex=tmp / paths.OURS / paths.LATEX)


def files_in(folder: Path) -> set[str]:
    return {p.relative_to(folder).as_posix() for p in folder.rglob("*") if p.is_file()}


def public_sources(coupon, tmp: Path, spec_text: str = PUBLIC_COUPON):
    """The public stand-in as a coupon the team has not built: its spec with its deck in decks/, as in a coupon's folder, and the
    IM7/8552 check's CalculiX record for its twin."""
    spec = tmp / f"{DESIGNATION}.toml"
    spec.write_text(spec_text, encoding="utf-8")
    paths.deck_path(spec, "abaqus").parent.mkdir()
    paths.deck_path(spec, "abaqus").write_text(write_abaqus(build_model(load_spec(spec))), encoding="utf-8", newline="\n")
    twin = (RECORD.parent / load_record(RECORD)["spec"]).resolve()
    return coupon.Sources(coupon=DESIGNATION, spec=spec, twin=twin, record=RECORD, manifest=None, documents=tmp / DESIGNATION / paths.DOCUMENTS,
                          team_files=tmp / DESIGNATION / paths.TEAM_FILES, latex=tmp / DESIGNATION / paths.OURS / paths.LATEX)


OWN_PUBLIC = (PUBLIC_COUPON.replace('library = "im7-8552"\ndataset = "ncamp-im7-8552-2019-rtd-measured-v1"\n',
                                     'E1 = 162095.74\nE2 = 8963.18\nG12 = 4688.43\nnu12 = 0.316\nF1t = 2558.51\nF1c = 1731.48\nF2t = 64.05\n'
                                     'F2c = 285.72\nF6 = 91.15\nsource = "the NCAMP report, typed in by the team"\n')
              .replace('ply_thickness = "library"', "ply_thickness = 0.18288"))
OWN_FIXTURE = NAMED_SPEC.replace('library = "im7-8552"\ndataset = "ncamp-im7-8552-2019-rtd-measured-v1"\n',
                                 "E1 = 162095.74\nE2 = 8963.18\nG12 = 4688.43\nnu12 = 0.316\nF1t = 2558.51\nF1c = 1731.48\nF2t = 64.05\nF2c = 285.72\n"
                                 'F6 = 91.15\nsource = "the fixture\'s own constants"\n')


def test_the_teams_own_constants_make_a_document_and_a_package_without_a_library_record(coupon, tmp_path):
    """A coupon from an input form with the team's own constants (ADR-0046): the material section, the deck heading, the package's
    material file and the audit's material row all say where the values come from; no library citation anywhere."""
    assert "library" not in OWN_PUBLIC.split("[material]")[1].split("[material.transverse]")[0]
    tex_path, package = coupon.build_document(public_sources(coupon, tmp_path, OWN_PUBLIC), pdf=False, theory=False)
    tex = tex_path.read_text(encoding="utf-8")
    assert r"\section{The material, as given by the team}\label{sec:material}" in tex and "the NCAMP report, typed in by the team" in tex
    assert "Rights class user-supplied" in tex and "under the name" in tex and "With the strengths the team gave the engine finds the first ply to fail" in tex
    material = (package / "3-material" / "IM7-8552-Lamina.csv").read_text(encoding="utf-8-sig").splitlines()
    assert material[:3] == ["material,IM7-8552-Lamina,the team's own constants", "source,\"the NCAMP report, typed in by the team\"",
                            "rights class,user-supplied,the team's values: kept with this coupon and never published"]
    assert "E1,162095.74,MPa,as given by the team" in material and "F6,91.15,MPa,as given by the team" in material and "G23,3000,MPa,as given by the team" in material
    assert not (package / "3-material" / "im7-8552.csv").exists()
    deck = (package / "2-our-model" / "abaqus" / f"{DESIGNATION}-abaqus.inp").read_text(encoding="utf-8")
    assert "the team's own constants: the NCAMP report, typed in by the team; rights class user-supplied" in deck
    form = load_workbook(package / f"{package.name}-audit.xlsx")["Rubric"]
    rows = {form.cell(r, 1).value: r for r in range(2, form.max_row + 1) if form.cell(r, 1).value}
    assert form.cell(rows["B6"], RUBRIC_COLUMNS.index("Ours") + 1).value.startswith("Material: the team's own constants: the NCAMP report")
    readme = (package / "README.txt").read_text(encoding="utf-8")
    assert "3-material/IM7-8552-Lamina.csv" in readme and "the team's own constants as given on the coupon input form" in readme


def test_the_teams_own_constants_are_set_beside_their_deck_in_the_team_comparison(coupon, tmp_path):
    src = team_sources(coupon, tmp_path, OWN_FIXTURE)
    tex_path, package = coupon.build_document(src, pdf=False, theory=False, write_summary=False)
    tex = tex_path.read_text(encoding="utf-8")
    assert "The team gave this material's constants on the coupon input form" in tex and "the fixture's own constants" in tex
    material = tex.split(r"\label{sec:material}")[1].split(r"\section")[0]
    assert material.count(r"\textcolor{passgreen}{same}") >= 4 and r"\textbf{differs}" not in material  # the deck types the same constants
    assert (package / "3-material" / "EXPLICIT.csv").is_file()  # the fixture names no material, so the deck's name is ours
    block = coupon.team.title_block(coupon.team.Comparison(src.manifest))  # the team comparison's own title block names the source too
    assert "Material source & the team" in block and "own constants: the fixture" in block


def test_a_coupon_the_team_has_not_built_gets_its_chapters_and_one_package(coupon, tmp_path):
    tex_path, package = coupon.build_document(public_sources(coupon, tmp_path), pdf=False, theory=False)
    documents = tmp_path / DESIGNATION / "documents"
    tex = tex_path.read_text(encoding="utf-8")
    assert all(rf"\chapter{{{title}}}" in tex for title in CHAPTERS)
    assert "The team has not built this coupon yet" in tex and r"\label{sec:results}" not in tex  # no team side by side
    assert package == documents / f"{TODAY}_{DESIGNATION}"
    assert sorted(p.name for p in documents.iterdir()) == [package.name]  # the packages only: no zip, no loose PDF
    assert sorted(p.name for p in (tmp_path / DESIGNATION / paths.OURS / paths.LATEX).iterdir()) == [f"{DESIGNATION}.tex", "figures"]  # the source is ours
    names = files_in(package)
    for name in ("README.txt", f"{package.name}-audit.xlsx", f"2-our-model/abaqus/{DESIGNATION}-abaqus.inp", "2-our-model/calculix/unnotched-im7-8552-nodal-ccx.inp",
                 "2-our-model/calculix/unnotched-im7-8552-nodal-abaqus.inp",  # the twin our engine compared with CalculiX
                 f"2-our-model/spec/{DESIGNATION}.toml", "2-our-model/spec/unnotched-im7-8552-nodal.toml", "3-material/im7-8552.csv",
                 "4-results/ours-all-plies.csv"):
        assert name in names, name
    # what the team reads: numbered figures the text refers to, no decision-record numbers, the hand check where the record compares
    prose = coupon.VERBATIM.sub("", tex)
    assert "ADR-" not in prose and "engine theory, Section" not in prose and "the engine manual, Sections 3 and 4" in prose
    for label in ("fig:model", "fig:solution", "fig:differences"):
        assert rf"\label{{{label}}}" in tex, label
    assert r"Figure~\ref{fig:solution}" in tex and r"Figure~\ref{fig:differences}" in tex
    assert "The kernel's values are read at element" in tex and "the centre probe" in tex
    assert r"\evidence{MANUAL" not in tex  # the guide for the team is Abaqus/CAE only: no PrePoMax tags, not even in the legend
    assert "The full CalculiX comparison table" not in tex  # every group of 69 values is printed in Chapter 6 already
    assert r"of the larger of its CalculiX value and 2\% of its group's largest value" in tex  # the floor rule, said as the record applies it
    assert "Attempt:" in (package / "README.txt").read_text(encoding="utf-8")
    assert not any(name.endswith(".frd") for name in names) and not any(name.startswith("1-team-model") for name in names)  # one layout, no team deck
    readme = (package / "README.txt").read_text(encoding="utf-8").splitlines()
    assert readme[0] == "Anisotropy AI share folder" and readme[1].startswith(f"{DESIGNATION}, unnotched tension")
    assert any(line.startswith(f"{package.name}-audit.xlsx ") for line in readme) and "There is no 1-team-model folder" in "\n".join(readme)
    plies = (package / "4-results/ours-all-plies.csv").read_text(encoding="utf-8").splitlines()
    assert plies[0].startswith("team element,our element,x (mm),y (mm),ply (1 = bottom),angle (deg),S11 (MPa)") and plies[1].startswith(",")  # no team numbers


def test_the_coupon_document_compiles_where_tectonic_is_installed(coupon, tmp_path):
    """The other tests write the LaTeX only; a brace too many in the cover reached the PDF build. Skipped where Tectonic is absent (CI)."""
    if coupon.manual.tectonic() is None:
        pytest.skip("Tectonic is not installed")
    pdf, package = coupon.build_document(public_sources(coupon, tmp_path), pdf=True, theory=False)
    assert pdf == package / f"{DESIGNATION}.pdf" and pdf.is_file()


LOCAL_COUPONS = [d.name for d in paths.model_dirs() if (d / paths.OURS / paths.VERIFICATION / f"{d.name}-calculix.toml").is_file()
                 and load_spec(d / paths.OURS / f"{d.name}.toml").mesh.mesher == "structured"]  # a coupon just started from its form has no record yet


@pytest.mark.skipif(not LOCAL_COUPONS, reason="no coupon model on this machine: the team's coupons are local (.local/fea/<coupon>/, ADR-0042)")
def test_a_local_coupon_builds_from_the_command_line(tmp_path):
    name = LOCAL_COUPONS[0]
    result = subprocess.run([sys.executable, str(COUPON), name, "--root", str(tmp_path), "--no-pdf", "--no-theory"],
                            capture_output=True, text=True, cwd=ROOT)
    assert result.returncode == 0, result.stdout + result.stderr
    documents = tmp_path / name / "documents"
    assert sorted(p.name for p in documents.iterdir()) == [f"{TODAY}_{name}"]
    tex = tmp_path / name / paths.OURS / paths.LATEX / f"{name}.tex"
    assert all(rf"\chapter{{{title}}}" in tex.read_text(encoding="utf-8") for title in CHAPTERS)
    assert (documents / f"{TODAY}_{name}" / f"{TODAY}_{name}-audit.xlsx").is_file()


def test_the_team_comparison_joins_the_document_and_the_package(coupon, tmp_path):
    src = team_sources(coupon, tmp_path)
    tex_path, package = coupon.build_document(src, pdf=False, theory=False)
    tex = tex_path.read_text(encoding="utf-8")
    assert all(rf"\chapter{{{title}}}" in tex for title in CHAPTERS)
    for label in ("sec:models", "sec:results", "sec:material", "sec:files", "sec:rebuild", "sec:summary", "sec:reading", "sec:open", "ch:audit"):
        assert rf"\label{{{label}}}" in tex, label
    assert "IM7-UNT-L9" in tex and "What the team built" in tex and "Step by step in Abaqus/CAE" in tex
    assert "PrePoMax 2.6}" not in tex  # the Abaqus chapter writes the Abaqus/CAE steps only
    assert package == tmp_path / "documents" / f"{TODAY}_grips"
    expected = ["README.txt", f"{package.name}-audit.xlsx", "1-team-model/coupon-2x8-grips-im7-8552.inp", "2-our-model/abaqus/grips-coupon-abaqus.inp",
                "2-our-model/spec/grips-coupon.toml", "3-material/im7-8552.csv", "4-results/ours-S11.csv", "4-results/ours-S22.csv", "4-results/ours-Mises.csv",
                "4-results/team-vs-ours.csv", "4-results/ours-all-plies.csv"]
    names = files_in(package)
    assert set(expected) <= names and not any(name.endswith(".frd") for name in names)
    ours = (package / "4-results/ours-S22.csv").read_text(encoding="utf-8").splitlines()[0]
    assert ours == (tmp_path / "team/Test_S22.csv").read_text(encoding="utf-8").splitlines()[0]  # the team's columns, in the team's order
    readme = (package / "README.txt").read_text(encoding="utf-8").splitlines()
    assert readme[1] == "IM7-UNT-L9, unnotched tension: the team's Abaqus model and ours" and "Test data:     none" in readme
    summary = paths.team_summary(tmp_path / paths.OURS).read_text(encoding="utf-8")  # the figures for the case page, in ours/verification/ (ADR-0039)
    assert 'coupon = "IM7-UNT-L9"' in summary and "[equivalence]" in summary
    form = load_workbook(package / f"{package.name}-audit.xlsx")
    cover = {form["Cover"].cell(r, 1).value: form["Cover"].cell(r, 2).value for r in range(1, form["Cover"].max_row + 1)}
    record = {form[RECORD_SHEET].cell(r, 1).value: form[RECORD_SHEET].cell(r, 2).value for r in range(1, form[RECORD_SHEET].max_row + 1)}
    assert cover["Package"] == package.name and record["Team deck"] == "coupon-2x8-grips-im7-8552.inp"  # the digest and the build: Record sheet
    assert cover["Team's deck"].startswith("coupon-2x8-grips-im7-8552.inp (") and cover["Our deck"].startswith("grips-coupon-abaqus.inp, written from the spec ")
    assert "/" not in cover["Team's deck"] and "/" not in cover["Our deck"]  # file names only, as the founder asked
    rows = {form["Rubric"].cell(r, 1).value: r for r in range(2, form["Rubric"].max_row + 1)}
    result = lambda item: form["Rubric"].cell(rows[item], RUBRIC_COLUMNS.index("Result") + 1).value  # noqa: E731
    assert result("B1") == "Pass" and result("C3") == "Pass" and result("B6") is None  # the same mesh and stresses; the auditor judges B6
    assert result("C2") is None  # the fixture reports no S12: the row waits for the team, as rows C1, C4 and C7 do
    # the fixture deck types the constants to two decimals in MPa, our library carries more digits: a real difference, sized for the auditor
    assert result("B5") == "Fail" and "Largest relative difference" in form["Rubric"].cell(rows["B5"], RUBRIC_COLUMNS.index("Difference") + 1).value
    assert r"Audit} (rubric 0.3" in tex
    assert coupon.build_document(src, pdf=False, theory=False)[1] == package  # a second build replaces the package it wrote


def test_a_document_names_its_coupon(coupon, tmp_path):
    with pytest.raises(SystemExit, match="has no \\[coupon\\] table"):
        coupon.build_document(team_sources(coupon, tmp_path, LIBRARY_SPEC), pdf=False, theory=False)


def test_a_folder_the_project_did_not_write_is_never_replaced(coupon, tmp_path):
    foreign = tmp_path / "documents" / f"{TODAY}_grips"
    foreign.mkdir(parents=True)
    (foreign / "notes.txt").write_text("keep me\n", encoding="utf-8")
    with pytest.raises(SystemExit, match="was not written by this project"):
        coupon.build_document(team_sources(coupon, tmp_path), pdf=False, theory=False)
    assert files_in(foreign) == {"notes.txt"}


def test_the_team_comparison_is_not_a_document_of_its_own():
    result = subprocess.run([sys.executable, str(TEAM)], capture_output=True, text=True, cwd=ROOT)
    assert result.returncode != 0 and "chapters of the coupon document" in result.stdout + result.stderr


# --------------------------------------------------------------------------------------------- after the audit (ADR-0043)

def form_rows(path: Path) -> dict[str, str]:
    sheet = load_workbook(path)[RUBRIC_SHEET]
    return {sheet.cell(r, 1).value: sheet.cell(r, RUBRIC_COLUMNS.index("Result") + 1).value for r in range(2, sheet.max_row + 1)}


def returned_form(package: Path, folder: Path, changes: dict[str, str] | None = None) -> Path:
    """The package's form as an auditor returns it: every row of theirs Pass and the verdict the rule gives; `changes` alters
    our tools' results, with the digest, as an earlier build of the coupon would have given them."""
    book = load_workbook(package / f"{package.name}-audit.xlsx")
    sheet = book[RUBRIC_SHEET]
    column = {name: RUBRIC_COLUMNS.index(name) + 1 for name in RUBRIC_COLUMNS}
    rows = []
    for r in range(2, sheet.max_row + 1):
        item, ours = sheet.cell(r, 1).value, sheet.cell(r, column["Judged by"]).value == "our tools"
        if not item:
            continue  # a section's banner row
        cell = sheet.cell(r, column["Result"])
        cell.value = (changes or {}).get(item, cell.value) if ours else "Pass"
        if cell.value == "Fail":
            sheet.cell(r, column["Cause"]).value = "engine"  # a failed row names its cause
        rows.append(Row(item, sheet.cell(r, column["Judged by"]).value, sheet.cell(r, column["Required"]).value == "yes", cell.value, "", ""))
    cover = book[COVER]
    labels = {cover.cell(r, 1).value: r for r in range(1, cover.max_row + 1)}
    for label, value in (("Auditor", "A. Engineer"), ("Audit date", "2026-10-24"), ("Verdict", verdict_of(tuple(rows)))):
        cover.cell(labels[label], 2).value = value
    record = book[RECORD_SHEET]
    digest_row = next(r for r in range(1, record.max_row + 1) if record.cell(r, 1).value == DIGEST)
    record.cell(digest_row, 2).value = results_digest([(row.id, row.result) for row in rows if row.judged_by == "our tools"])
    folder.mkdir(parents=True, exist_ok=True)
    book.save(folder / f"{package.name}-audit.xlsx")
    return folder / f"{package.name}-audit.xlsx"


def test_a_rebuilt_deck_folder_is_read_while_section_e_is_off_the_rubric(coupon, tmp_path):
    """Rubric 0.3 defers section E (the rebuild check, E2), so a rebuilt deck adds no row; the folder is still read, so two decks there are refused."""
    src = team_sources(coupon, tmp_path)
    rebuilt = tmp_path / paths.TEAM_FILES / "2026-10-24" / "rebuilt"
    rebuilt.mkdir(parents=True)
    (rebuilt / "Job-1.inp").write_text(write_abaqus(build_model(load_spec(src.spec))), encoding="utf-8")  # rebuilt exactly as our deck
    _, package = coupon.build_document(src, pdf=False, theory=False)
    assert not any(item.startswith("E") for item in form_rows(package / f"{package.name}-audit.xlsx") if item)
    (rebuilt / "Job-3.inp").write_bytes(GRIPS_FIXTURE.read_bytes())
    with pytest.raises(SystemExit, match="holds 2 .inp files"):
        coupon.build_document(src, pdf=False, theory=False)


def test_the_regression_check_compares_with_the_newest_returned_audit(coupon, tmp_path, capsys):
    src = team_sources(coupon, tmp_path)
    assert coupon.check(src) == [] and "no audit form has come back" in capsys.readouterr().out
    _, package = coupon.build_document(src, pdf=False, theory=False)
    returned_form(package, tmp_path / paths.TEAM_FILES / "2026-10-24")
    summary = paths.team_summary(tmp_path / paths.OURS)
    summary.unlink()  # the check builds in a scratch folder and writes nothing beside the model
    assert coupon.check(src) == [] and "still pass" in capsys.readouterr().out
    assert not summary.exists()
    returned_form(package, tmp_path / paths.TEAM_FILES / "2026-10-25", changes={"B5": "Pass"})  # an earlier engine passed B5
    assert coupon.check(src) == ["B5"]
    assert "REGRESSION" in capsys.readouterr().out
