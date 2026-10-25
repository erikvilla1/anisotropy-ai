"""The coupon audit (ADR-0043): the committed rubric reads, agrees with the tools it names, and a malformed one is refused; a
package is scored from what our tools found; the audit form locks our results and opens the auditor's cells."""

import dataclasses
import re
from datetime import datetime

import pytest
from openpyxl import load_workbook

from anisotropy_fea import case_pages, paths  # where the files live (ADR-0036)
from anisotropy_fea.abaqus_reader import parse_deck, read_deck
from anisotropy_fea.abaqus_report import read_field_report
from anisotropy_fea.audit import (AUDITOR_BLOCK, AUDITOR_FIELDS, AUTO, CAUSES, COMPARED, COVER, EQUIVALENCE_GROUPS, FONT, GLANCE, HOW_TO, LISTED_ONLY,
                                  NOT_BUILT_BY_TEAM, RECORD, RESULTS_SHEET, RUBRIC_COLUMNS, RUBRIC_SHEET, RULE_FIELD, SECTIONS, TO_SEND, VALUES_SHEET,
                                  AuditError, Evidence, Found, Row, RubricError, _Facts, _laminate, form_tables, items_for, load_rubric, parse_rubric,
                                  read_form, score, verdict_of, write_form)
from anisotropy_fea.report import laminate_check
from anisotropy_fea.equivalence import RELATIVE, deck_equivalence
from anisotropy_fea.membrane import solve_membrane
from anisotropy_fea.model import to_si
from anisotropy_fea.side_by_side import pair_reports

from test_equivalence import FIXTURE_SPEC, GRIPS_FIXTURE, generated
from test_side_by_side import kernel_reports

UNITS = "mm-MPa"
CARD = (("Coupon", "IM7-UNT-L9"), ("Test", "unnotched tension"), ("Standard", "none"), ("Size from", "the grips fixture"),
        ("Team model", "the grips fixture deck"), ("Test data", "none"), ("Material", "IM7/8552"), ("Layup", "[0/30/-30/30/0]s"), ("Size", "25 x 100 mm"))

SMALL = """
version = "0.1"
status = "draft"

[[section]]
id = "A"
title = "Identity"

[[section]]
id = "B"
title = "Model match"

[[item]]
id = "A1"
section = "A"
check = "The coupon card is complete"
acceptance = "No field is blank"
method = "auto"
required = true
source = ["spec.coupon_card"]

[[item]]
id = "B1"
section = "B"
check = "Geometry"
acceptance = "Nodes match"
method = "auto"
required = true
tolerance = { length_m = 1e-8 }
applies_to = ["UNT", "OHT"]
"""


def test_the_committed_rubric_reads():
    rubric = load_rubric(paths.RUBRIC)
    assert [s.id for s in rubric.sections] == ["A", "B", "C", "D"] and all(s.meaning for s in rubric.sections)  # E, documentation, is deferred
    assert not any("ADR" in item.acceptance or "ADR" in item.check for item in rubric.items)  # the rubric's words go to the team unchanged
    assert len({item.id for item in rubric.items}) == len(rubric.items)
    assert rubric.version == "0.3" and rubric.status == "confirmed", "confirmed by the founder on 2026-10-24 (ADR-0043 amendments 1 and 4)"
    assert [item.id for item in rubric.items if item.section == "A"] == ["A1", "A2"]  # the files-identified row left in 0.3: the Record sheet names them
    d1 = next(item for item in rubric.items if item.id == "D1")
    assert d1.tolerance.relative == 0.01 and d1.condition.min_free_length_widths == 3  # the same tolerance as C2; N/A for a short gauge


def test_model_match_uses_the_equivalence_tolerance():
    """B4, B5 and B8 state the tolerance deck_equivalence applies; the two must not drift apart."""
    items = {item.id: item for item in load_rubric().items}
    for item_id in ("B4", "B5", "B8"):
        assert items[item_id].tolerance.relative == RELATIVE, item_id
    assert all(items[f"B{n}"].source == ("equivalence.deck_equivalence",) for n in (1, 2, 3, 4, 5, 7, 8, 9))


def test_an_auditor_item_has_no_tolerance_and_an_unbuilt_check_is_not_required():
    for item in load_rubric().items:
        if item.method == "auditor":
            assert item.tolerance is None, f"{item.id}: the auditor judges it; a tolerance belongs to a computed check"
        if item.method == "auto" and not item.source:
            assert not item.required, f"{item.id}: no function computes it yet, so a Pass verdict cannot need it"


def test_items_follow_the_test():
    rubric = load_rubric()
    unt = {item.id for item in items_for(rubric, "UNT")}
    oht = {item.id for item in items_for(rubric, "OHT")}
    assert "D1" in unt and "D3" not in unt
    assert "D3" in oht and "D4" in oht and "D1" not in oht
    assert {"A1", "B1", "C2", "D2"} <= unt & oht
    with pytest.raises(RubricError, match="not in spec.TEST_CODES"):
        items_for(rubric, "XYZ")


def test_a_small_rubric_reads():
    rubric = parse_rubric(SMALL)
    assert [item.id for item in rubric.items] == ["A1", "B1"]
    assert rubric.items[1].tolerance.length_m == 1e-8
    assert [item.id for item in items_for(rubric, "UNC")] == ["A1"]


@pytest.mark.parametrize("old, new, message", [
    ('id = "B1"', 'id = "B2"', "expected B1"),
    ('id = "B1"\nsection = "B"', 'id = "A2"\nsection = "B"', "its id names A"),
    ('id = "B1"\nsection = "B"', 'id = "C1"\nsection = "C"', "is not one of"),
    ('id = "B"\ntitle', 'id = "A"\ntitle', "used twice"),
    ('method = "auto"\nrequired = true\nsource', 'method = "machine"\nrequired = true\nsource', "method"),
    ("tolerance = { length_m = 1e-8 }", "tolerance = { length_m = 1e-8, relative = 0.01 }", "exactly one"),
    ("tolerance = { length_m = 1e-8 }", "tolerance = { length_m = -1e-8 }", "greater than 0"),
    ('applies_to = ["UNT", "OHT"]', 'applies_to = ["UNT", "XYZ"]', "not in spec.TEST_CODES"),
    ('source = ["spec.coupon_card"]', 'source = ["spec.no_such_function"]', "is not a function of anisotropy_fea"),
    ('source = ["spec.coupon_card"]', 'source = ["coupon_card"]', "is not a function of anisotropy_fea"),
    ('source = ["spec.coupon_card"]', 'source = ["spec.CouponSpec.no_such_method"]', "is not a function of anisotropy_fea"),
    ('source = ["spec.coupon_card"]', 'source = ["spec.TEST_CODES"]', "is not a function of anisotropy_fea"),
    ('status = "draft"', 'status = "approved"', "status"),
    ('check = "Geometry"', 'check = "Geometry"\nowner = "team"', "Extra inputs"),
    ('acceptance = "Nodes match"\n', "", "acceptance"),
])
def test_a_malformed_rubric_is_refused(old, new, message):
    assert SMALL.count(old) == 1, old
    with pytest.raises(RubricError, match=message):
        parse_rubric(SMALL.replace(old, new), "small rubric")


def test_a_section_without_items_is_refused():
    with pytest.raises(RubricError, match=r"sections \['C'\] have no items"):
        parse_rubric(SMALL + '\n[[section]]\nid = "C"\ntitle = "Results"\n')


# --------------------------------------------------------------------------------------------- scoring a package

@pytest.fixture(scope="module")
def grips(tmp_path_factory):
    """The grips fixture as the team's deck, our generated deck of it, and the team's field reports made by our kernel."""
    folder = tmp_path_factory.mktemp("team")
    team = read_deck(GRIPS_FIXTURE, units=UNITS)
    ours = to_si(parse_deck(generated(FIXTURE_SPEC)), UNITS)
    equivalence = deck_equivalence(team, ours, 1e-8, UNITS)
    solution = solve_membrane(ours)
    reports = [read_field_report(path) for path in kernel_reports(GRIPS_FIXTURE, folder)]
    paired = tuple(pair_reports(reports, team, solution, equivalence, UNITS))
    return Evidence(test_code="UNT", units=UNITS, card=CARD, our_deck=ours, solution=solution, material="explicit constants", team_deck=team,
                    team_deck_sha256="0" * 64, equivalence=equivalence, paired=paired, inner_band=(10.0, 90.0), free_band=(0.0, 100.0),
                    excluded="10 mm beside each grip line", fpf=("max_stress", 0.5, 1, 0.0))


def results(evidence: Evidence) -> dict[str, str | None]:
    return {s.item.id: s.result for s in score(load_rubric(), evidence)}


def test_the_same_model_with_the_same_stresses_scores_pass_where_our_tools_judge(grips):
    found = results(grips)
    assert [item_id for item_id, result in found.items() if result == "Pass"] == ["A1", "A2", "B1", "B2", "B3", "B4", "B5", "B7", "B8", "B9", "C3", "D2"]
    # the auditor's: the auditor items, and C1, C2, C4 and C7, whose team values are not all in the export (our values are stated; C2 lacks S12)
    assert {item_id for item_id, result in found.items() if result is None} == {"B6", "C1", "C2", "C4", "C5", "C6", "C7", "D5"}
    assert not any(item_id.startswith("E") for item_id in found)  # section E, documentation, is off the rubric for now
    assert "D3" not in found and "D4" not in found  # open-hole only


def test_the_scores_say_why(grips):
    found = {s.item.id: s.found for s in score(load_rubric(), grips)}
    c2 = found["C2"]
    assert c2.team[1].startswith("S11: 24 values; largest ") and c2.ours[1].startswith("S11: 24 values; largest ") and "within 1 %" in c2.difference[1]
    assert c2.team[-1] == "S12: not in the reports, left blank, for the team to fill in" and c2.ours[-1].startswith("S12: largest ")
    assert " in the band, at our element " in c2.ours[-1] and c2.difference[-1].startswith("S12: not compared; the Result waits until the team has verified S12")
    assert "unbalanced (2 plies at +30, 1 at -30)" not in found["B3"].text  # the full laminate counts both halves
    assert found["B3"].ours[1] == "Balance: unbalanced (4 plies at +30, 2 at -30), symmetric" and found["B3"].difference == ("Layup, bottom first: the same", "Balance: the same")
    assert found["C7"].team == ("left blank, for the team to fill in",) and "ply 1 (0 deg) by max stress, at 0.5 of the applied loading" in found["C7"].ours[0]
    # the columns are the claim the auditor verifies: the team's value beside ours, no node offsets and no "as written"
    b1 = found["B1"]
    assert b1.team[0].startswith("Nodes: ") and " on the part, " in b1.team[0] and b1.team == b1.ours and len(b1.team) == 1
    assert b1.difference == ("Nodes: the same", "Node positions: the same, node for node")
    assert "offset" not in b1.text and "as written" not in found["B2"].text and found["B5"].difference == tuple(f"{c}: the same" for c in ("E1", "E2", "nu12", "G12", "G13", "G23"))
    assert found["A2"] == Found(("mm and MPa",), ("mm and MPa", "This audit: mm, MPa and N"), ("The same",))
    b8 = found["B8"]  # boundary conditions and loads: the team's side in full beside ours
    assert b8.team[0].startswith("Reference point at (") and b8.ours[0].startswith("Reference point at (") and "U2 = " in b8.team[0] and "UR3 = 0" in b8.team[0]
    assert "Concentrated loads: none" in b8.team and "Concentrated loads: none" in b8.ours
    assert found["C1"].team == ("left blank, for the team to fill in",) and found["C1"].ours[0].endswith(" N along the coupon at the top grip") and found["C1"].difference == ()
    assert found["C4"].ours[0].startswith("Largest displacement: U1 ") and found["C4"].ours[1].startswith("Strain at the gauge centre (element ")
    assert found["D1"].team == () and found["D1"].ours[0].startswith("Centre element ") and found["D1"].difference[-1] == "Tolerance 1 %"


def test_a_stress_beyond_the_tolerance_fails(grips):
    worst = next(p for p in grips.paired if p.quantity == "S22" and 10.0 <= p.y <= 90.0)
    changed = tuple(dataclasses.replace(p, reference=p.reference * 0.95) if p is worst else p for p in grips.paired)
    found = results(dataclasses.replace(grips, paired=changed))
    assert found["C2"] == "Fail" and found["C3"] == "Pass"


def test_a_coupon_the_team_has_not_built_cannot_be_checked_against_their_model(grips):
    alone = dataclasses.replace(grips, team_deck=None, team_deck_sha256=None, equivalence=None, paired=(), inner_band=None, excluded=None)
    by_id = {s.item.id: s for s in score(load_rubric(), alone)}
    assert all(by_id[f"B{n}"].result == "Not checked" for n in (1, 2, 3, 4, 5, 7, 8, 9))
    assert by_id["C2"].result == "Not checked" and by_id["C2"].found == Found(team=(NOT_BUILT_BY_TEAM,))
    assert by_id["A2"].result == "Pass" and by_id["D2"].result == "Pass"


def test_a_blank_card_field_fails_the_card_row(grips):
    assert results(dataclasses.replace(grips, card=CARD[:-1] + (("Size", " "),)))["A1"] == "Fail"


def test_the_hand_check_names_the_end_effect_and_is_not_applicable_to_a_short_gauge(grips):
    """D1 compares the gauge centre with free-edge laminate theory. The grips hold the width, so the evidence states the free length in
    widths and N_x/N_y at the centre, and a coupon whose free length is under the rubric's minimum gets N/A, not Fail."""
    item = next(item for item in load_rubric().items if item.id == "D1")
    check = laminate_check(grips.our_deck, grips.solution)
    facts = _Facts(check, "", (0.0,), "mm", "MPa")
    result, found = _laminate(item, grips, facts)
    evidence = "\n".join(found.difference)
    assert result == "Fail" and "Free length 4 widths; N_x at the centre" in evidence and "above free-edge theory" in evidence  # 2 by 8 mesh: 3 %
    short = _Facts(dataclasses.replace(check, free_length_m=check.width_m * 1.1), "", (0.0,), "mm", "MPa")
    result, found = _laminate(item, grips, short)
    assert result == "N/A" and found.difference[0].startswith("The free length is 1.1 widths, under the 3 the hand check needs") and "For the record:" in found.difference
    assert _laminate(item.model_copy(update={"condition": None}), grips, short)[0] == "Fail"  # without the condition the number alone decides


def test_every_item_our_tools_judge_has_a_scorer_and_every_equivalence_item_has_a_place(grips):
    assert {item.id for item in load_rubric().items if item.method == "auto"} <= set(AUTO)  # AUTO keeps E2's scorer while section E is deferred
    for item in grips.equivalence.items:
        homes = [item_id for item_id, pattern in EQUIVALENCE_GROUPS.items() if re.fullmatch(pattern, item.item)]
        assert len(homes) + (item.item in LISTED_ONLY) == 1, item.item


# --------------------------------------------------------------------------------------------- the audit form

def test_the_form_locks_our_results_and_opens_the_auditors_cells(grips, tmp_path):
    rubric = load_rubric()
    scored = score(rubric, grips)
    path = write_form(tmp_path / "audit.xlsx", rubric, scored, [("Coupon", "IM7-UNT-L9"), ("Attempt", "3")], UNITS)
    book = load_workbook(path)
    assert book.sheetnames == [COVER, RUBRIC_SHEET, RECORD]  # no tables given: no Results or Compared values sheet
    cover = book[COVER]
    labels = {cover.cell(r, 1).value: r for r in range(1, cover.max_row + 1)}
    assert cover.cell(labels["Attempt"], 2).value == "3" and cover.cell(2, 1).value == "IM7-UNT-L9, attempt 3"
    assert all(not cover.cell(labels[field], 2).protection.locked for field in AUDITOR_FIELDS) and cover.protection.sheet
    assert str(cover.cell(labels[RULE_FIELD], 2).value).startswith("=IF(") and cover.cell(labels[RULE_FIELD], 2).protection.locked  # fills itself
    record = book[RECORD]
    assert record.sheet_state == "hidden"  # for our tools; the auditor never sees it
    kept = {record.cell(r, 1).value: record.cell(r, 2).value for r in range(1, record.max_row + 1)}
    assert kept["Rubric"] == "version 0.3 (confirmed)" and len(kept["Our results digest"]) == 64 and record.protection.sheet
    sheet = book[RUBRIC_SHEET]
    assert [c.value for c in sheet[1]] == list(RUBRIC_COLUMNS)
    result, observation, cause = (RUBRIC_COLUMNS.index(name) + 1 for name in ("Result", "Observation", "Cause"))
    rows = {sheet.cell(r, 1).value: r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value}
    assert list(rows) == [s.item.id for s in scored]
    banners = [sheet.cell(r, 2).value for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value is None]
    assert banners == ["A. Identity and traceability", "B. Model match: our model is the team's model",
                       "C. Results match: verification against the team's solver", "D. Physics, independent of both solvers"]
    for s in scored:
        cell = sheet.cell(rows[s.item.id], result)
        if s.result is None:
            assert cell.value is None and not cell.protection.locked, s.item.id
        else:
            assert cell.value == s.result and cell.protection.locked, s.item.id
        assert not sheet.cell(rows[s.item.id], observation).protection.locked and not sheet.cell(rows[s.item.id], cause).protection.locked
    lists = {v.formula1 for v in sheet.data_validations.dataValidation}
    assert lists == {'"Pass,Fail,N/A,Not checked"', '"' + ",".join(CAUSES) + '"'}
    assert "\n" in sheet.cell(rows["B5"], RUBRIC_COLUMNS.index("Team") + 1).value  # one fact per line, the team's values beside ours
    assert sheet.cell(rows["B5"], RUBRIC_COLUMNS.index("Team") + 1).value == sheet.cell(rows["B5"], RUBRIC_COLUMNS.index("Ours") + 1).value
    assert sheet.cell(rows["D1"], RUBRIC_COLUMNS.index("Team") + 1).value is None  # the hand check is ours alone


def test_the_form_computes_the_rule_counts_the_rows_and_shows_the_numbers_behind_them(grips, tmp_path):
    """Form v0.3 (ADR-0044 amendment 1): the cover names what is audited and compared and takes the verdict; its rule verdict and section
    counts are formulas over the auditor's entries; the headline numbers and the values our tools computed are on the Results and
    Compared values sheets; how the package was built is on the Record sheet; the rows are coloured and the sheets print."""
    rubric = load_rubric()
    scored = score(rubric, grips)
    tables = form_tables(rubric, grips, scored)
    files = [("The model", "nodes, elements, layup (section B)"), ("  Team's deck", "coupon-2x8-grips-im7-8552.inp"), ("  Our deck", "grips-coupon-abaqus.inp")]
    record = [("Team deck", "coupon-2x8-grips-im7-8552.inp"), ("Team deck SHA-256", "0" * 64), ("Engine", "anisotropy-fea 0.1.0, commit abc1234")]
    path = write_form(tmp_path / "audit.xlsx", rubric, scored, [("Coupon", "IM7-UNT-L9"), ("Test", "unnotched tension"), ("Attempt", "3"),
                                                              ("Package", "2026-10-24_IM7-UNT-L9")], UNITS, tables=tables, files=files, record=record)
    book = load_workbook(path)
    assert book.sheetnames == [COVER, RUBRIC_SHEET, RESULTS_SHEET, VALUES_SHEET, RECORD] and book.properties.title == "Coupon audit: 2026-10-24_IM7-UNT-L9"
    cover = book[COVER]
    labels = {cover.cell(r, 1).value: r for r in range(1, cover.max_row + 1)}
    assert cover.cell(2, 1).value == "IM7-UNT-L9, unnotched tension, attempt 3"
    assert cover.cell(labels["Team's deck"], 1).alignment.indent == 2 and not cover.cell(labels["Team's deck"], 1).font.b  # a row under "The model"
    assert (labels["Package"] < labels[COMPARED] < labels["The model"] < labels["Team's deck"] < labels["Our deck"] < labels[AUDITOR_BLOCK] < labels[RULE_FIELD] < labels[SECTIONS]
            < labels["A. Identity and traceability"] < labels[HOW_TO] < labels[TO_SEND])
    rule = cover.cell(labels[RULE_FIELD], 2).value
    assert rule.startswith('=IF(COUNTIFS(Rubric!$E$2:$E$200,"yes",Rubric!$G$2:$G$200,"Not checked")') and rule.endswith('"Pass with observations","Pass")))')
    assert ',"",IF(' in rule and 'COUNTIF(Rubric!$K$2:$K$200,"?*")' in rule and cover.cell(labels[RULE_FIELD], 2).protection.locked  # blank until complete
    sections = [cover.cell(r, 1).value for r in range(labels[SECTIONS] + 1, labels[SECTIONS] + 5)]
    assert sections == ["A. Identity and traceability", "B. Model match: our model is the team's model",
                        "C. Results match: verification against the team's solver", "D. Physics, independent of both solvers"]
    assert cover.cell(labels["A. Identity and traceability"], 2).value == "the coupon card and the unit system" and cover.max_column == 2
    assert cover.cell(labels[SECTIONS] + 5, 1).value is None  # no section E, and no counts table
    for noise in ("Engine commit", "Team deck SHA-256", "Our results digest", "Run by", "Example finding", "Findings", "ADR", "as filled in", "rule"):
        assert not any(noise in str(cover.cell(r, c).value) for r in range(1, cover.max_row + 1) for c in (1, 2)), noise
    sent = [cover.cell(r, 1).value for r in range(labels[TO_SEND] + 1, cover.max_row + 1) if cover.cell(r, 1).value]
    assert sent[0].startswith("C1 (Total reaction force): the reaction force") and [line[:2] for line in sent] == ["C1", "C2", "C4", "C7"]
    sheet = book[RUBRIC_SHEET]
    result = RUBRIC_COLUMNS.index("Result") + 1
    last = len(scored) + len({s.item.section for s in scored}) + 1  # a banner row per section, under the header
    judged = {sheet.cell(r, 1).value: sheet.cell(r, RUBRIC_COLUMNS.index("Judged by") + 1).value for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value}
    assert judged["C1"] == "auditor" and judged["C2"] == "auditor" and judged["C3"] == "our tools"  # the team's values missing (S12 for C2): the auditor judges
    coloured = {str(r.sqref) for r in sheet.conditional_formatting}
    assert f"G2:G{last}" in coloured and result == 7 and sheet.print_title_rows == "$1:$1" and sheet.page_setup.orientation == "landscape"
    assert sheet.row_dimensions[3].height is not None and all(c.font.name == FONT for c in sheet[3] if c.value is not None)  # row 3: A1, under the A banner
    results = book[RESULTS_SHEET]
    first = {results.cell(r, 1).value: r for r in range(1, results.max_row + 1) if results.cell(r, 1).value}
    glance = dict(tables.glance)
    assert first[GLANCE] < first["S11 in the graded band"] < first["Hand check at the gauge centre"]
    assert glance["S11 in the graded band"].startswith("24 of 24 within 1 % of ") and "; worst " in glance["S11 in the graded band"]
    assert results.cell(first["Reaction at the loaded grip"], 2).value.startswith("ours ") and "not in the team's export" in glance["Reaction at the loaded grip"]
    assert "the free length is 4 widths" in glance["Hand check at the gauge centre"]
    titles = [results.cell(r, 1).value for r in range(1, results.max_row + 1) if results.cell(r, 1).font.b and results.cell(r, 1).font.sz == 11]
    assert [t.split(",")[0].split(":")[0] for t in titles] == [GLANCE, "Stresses in the graded band", "Strains at the gauge centre",
                                                               "Section forces at the gauge centre", "Ply stresses at the gauge centre in material axes (MPa)",
                                                               "Reaction at the loaded grip (N)"]
    checks = [t.check for t in tables.tables]  # every table says how to redo its numbers from the package's files or in Abaqus
    assert all(c.startswith("How to check: ") for c in checks) and len(checks) == 5
    assert "4-results/team-vs-ours.csv" in checks[0] and "LE11, LE22 and LE12 at the integration point of element" in checks[1] and "Ply-1 (0 deg)" in checks[1]
    assert "S11, S22 and S12 at element" in checks[3] and "reaction force RF2 at the top reference point" in checks[4]
    shown = [results.cell(r, 1).value for r in range(1, results.max_row + 1) if str(results.cell(r, 1).value).startswith("How to check: ")]
    assert shown == checks and all(results.cell(r, 1).font.i for r in range(1, results.max_row + 1) if str(results.cell(r, 1).value).startswith("How to check"))
    band = tables.tables[0]
    assert band.columns[:4] == ("Quantity", "Layer (1 = bottom)", "Values compared", "Within tolerance")
    assert [(row[0], row[1], row[2], row[3]) for row in band.rows if row[1] == "all"] == [("S11", "all", 24, 24), ("S22", "all", 24, 24), ("Mises", "all", 24, 24)]
    header = first["Stresses in the graded band, y 10 to 90 mm: the team's Abaqus against our engine"] + 1
    assert results.protection.sheet and results.cell(header, 1).value == "Quantity" and results.cell(header + 1, 6).number_format == "0.00"
    values = book[VALUES_SHEET]
    assert values.max_row == len(grips.paired) + 1 and values.auto_filter.ref == f"A1:K{len(grips.paired) + 1}" and not values.protection.sheet
    assert [c.value for c in values[1]][:6] == ["Team element", "Our element", "x (mm)", "y (mm)", "Layer (1 = bottom)", "Quantity"]
    assert values.cell(2, 7).number_format == "0.00" and values.cell(2, 3).number_format == "0.000"
    kept = {book[RECORD].cell(r, 1).value: book[RECORD].cell(r, 2).value for r in range(1, book[RECORD].max_row + 1)}
    assert kept["Team deck SHA-256"] == "0" * 64 and kept["Engine"] == "anisotropy-fea 0.1.0, commit abc1234" and kept["Rubric"] == "version 0.3 (confirmed)"
    assert book[RECORD].sheet_state == "hidden" and not any("ADR" in str(c.value) for row in results.iter_rows() for c in row if c.value is not None)
    assert read_form  # a form with these sheets reads back: test_a_form_with_the_numbers_behind_it_reads_back


# --------------------------------------------------------------------------------------------- a returned form

def filled(folder, evidence, results=None, verdict="Pass", observations=None, causes=None, changed=None, when="2026-10-24", auditor="A. Engineer",
           tables=False, package=None):
    """The form of the grips package as an auditor returns it. D1 is set to Pass so that a clean form can pass: on this coarse
    2 by 8 mesh the centre element is about 3 % from laminate theory, beyond the rubric's 1 %."""
    rubric = load_rubric()
    scored = tuple(dataclasses.replace(s, result="Pass") if s.item.id == "D1" else s for s in score(rubric, evidence))
    cover = [("Coupon", "IM7-UNT-L9"), ("Attempt", "3")] + ([("Package", package)] if package else [])
    path = write_form(folder / "audit.xlsx", rubric, scored, cover, UNITS, tables=form_tables(rubric, evidence, scored) if tables else None,
                      record=[("Team deck", "coupon-2x8-grips-im7-8552.inp"), ("Team deck SHA-256", "0" * 64)])
    book = load_workbook(path)
    sheet = book[RUBRIC_SHEET]
    rows = {sheet.cell(r, 1).value: r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value}
    column = {name: RUBRIC_COLUMNS.index(name) + 1 for name in RUBRIC_COLUMNS}
    for s in scored:
        if s.result is None:
            sheet.cell(rows[s.item.id], column["Result"]).value = (results or {}).get(s.item.id, "Pass")
    for item_id, value in (changed or {}).items():
        sheet.cell(rows[item_id], column["Result"]).value = value
    for name, entries in (("Observation", observations), ("Cause", causes)):
        for item_id, value in (entries or {}).items():
            sheet.cell(rows[item_id], column[name]).value = value
    cover_sheet = book[COVER]
    labels = {cover_sheet.cell(r, 1).value: r for r in range(1, cover_sheet.max_row + 1)}
    for label, value in (("Auditor", auditor), ("Audit date", when), ("Verdict", verdict)):
        cover_sheet.cell(labels[label], 2).value = value
    book.save(path)
    return path


def test_a_clean_form_reads_back(grips, tmp_path):
    form = read_form(filled(tmp_path, grips, when=datetime(2026, 10, 24, 9, 30)))
    assert form.verdict == "Pass" and form.auditor == "A. Engineer" and form.date == "2026-10-24" and form.counts() == "none"
    assert form.cover["Attempt"] == "3" and form.cover["Team deck SHA-256"] == "0" * 64  # the cover and the Record sheet, read as one
    assert {r.id: r.result for r in form.rows}["B6"] == "Pass"


def test_a_form_with_the_numbers_behind_it_reads_back(grips, tmp_path):
    """The Results and Compared values sheets, and the cover's formulas, do not get in the way of reading a returned form."""
    form = read_form(filled(tmp_path, grips, tables=True))
    assert form.verdict == "Pass" and form.cover[RULE_FIELD] == "Pass"  # the word `filled` wrote where the spreadsheet would save the formula's value
    assert load_workbook(form.path).sheetnames == [COVER, RUBRIC_SHEET, RESULTS_SHEET, VALUES_SHEET, RECORD]


def test_the_verdict_follows_the_rule(grips, tmp_path):
    noted = read_form(filled(tmp_path, grips, verdict="Pass with observations", observations={"B3": "the layup table omits the ply thickness"}))
    assert noted.verdict == "Pass with observations" and noted.counts() == "1 observation"
    failed = read_form(filled(tmp_path, grips, verdict="Fail", results={"B6": "Fail"}, causes={"B6": "our input"}))  # a required row fails
    assert failed.verdict == "Fail" and failed.counts() == "1 fail (B6)" and {r.id: r.cause for r in failed.rows}["B6"] == "our input"


def test_the_rule_on_its_own():
    row = lambda item, required, result, observation="": Row(item, "auditor", required, result, observation, "")  # noqa: E731
    assert verdict_of((row("A1", True, "Pass"), row("C1", False, "Not checked"))) == "Pass"  # an optional row may stay unchecked
    assert verdict_of((row("A1", True, "Pass"), row("E2", False, "Fail"))) == "Pass with observations"  # an optional row fails
    assert verdict_of((row("A1", True, "N/A", "the card names no standard"),)) == "Pass with observations"  # an observation on a passing row
    assert verdict_of((row("A1", True, "Fail"),)) == "Fail"
    with pytest.raises(AuditError, match="required rows A1"):
        verdict_of((row("A1", True, "Not checked"),))


@pytest.mark.parametrize("kwargs, message", [
    (dict(verdict="Pass", observations={"B3": "a missing table"}), "the rule gives 'Pass with observations'"),
    (dict(changed={"C3": "Fail"}, causes={"C3": "engine"}), "a result our tools gave was changed"),  # C3 is our tools' row; C2 waits on S12
    (dict(results={"B6": ""}), "row B6 has no result"),
    (dict(results={"B6": "Not checked"}), "required rows B6 have no result other than Not checked"),
    (dict(verdict="Fail", results={"B6": "Fail"}), "row B6 fails without a cause"),
    (dict(causes={"C3": "gremlins"}), "row C3: the cause 'gremlins' is not one of"),
    (dict(auditor=""), "the cover's Auditor is empty"),
    (dict(when="last Tuesday"), "is not a date"),
    (dict(verdict="Approved"), "the verdict 'Approved' is not one of"),
])
def test_an_incomplete_or_inconsistent_form_is_refused_by_name(grips, tmp_path, kwargs, message):
    with pytest.raises(AuditError, match=re.escape(message)):
        read_form(filled(tmp_path, grips, **kwargs))


def test_a_file_that_is_not_a_form_is_refused(tmp_path):
    (tmp_path / "notes.xlsx").write_text("not a workbook", encoding="utf-8")
    with pytest.raises(AuditError, match="not an audit form"):
        read_form(tmp_path / "notes.xlsx")


# --------------------------------------------------------------------------------------------- filing a returned form

CASE_PAGE = """# IM7-UNT-L9, unnotched tension

**Case.** Category: team test | Status: in progress | Accepted attempt: 03

<!-- generated: result -->
<!-- end generated -->

## Attempt log

| NN | Date | Tried | Result | Outcome | Evidence |
| --- | --- | --- | --- | --- | --- |
| 01 | 2026-10-20 | a | b | c | d |
| 02 | 2026-10-21 | a | b | c | d |
| 03 | 2026-10-22 | a | b | c | d |

## Where to look
"""


@pytest.fixture
def coupon_folders(tmp_path, monkeypatch):
    """A model folder with a case page, and an empty local folder for the coupon, in place of paths.MODELS and paths.LOCAL."""
    models, local = tmp_path / "models", tmp_path / "local"
    (models / "IM7-UNT-L9").mkdir(parents=True)
    (models / "IM7-UNT-L9" / "README.md").write_text(CASE_PAGE, encoding="utf-8")
    (local / "IM7-UNT-L9").mkdir(parents=True)
    monkeypatch.setattr(paths, "MODELS", models)
    monkeypatch.setattr(paths, "LOCAL", local)
    monkeypatch.setattr(case_pages, "result_block", lambda folder: "the result\n")  # the CalculiX record is not what this test is about
    return models / "IM7-UNT-L9", local / "IM7-UNT-L9"


def returned(tmp_path, grips, coupon="IM7-UNT-L9", **kwargs):
    folder = tmp_path / "came-back"
    folder.mkdir(exist_ok=True)
    path = filled(folder, grips, package="2026-10-23_IM7-UNT-L9", **kwargs)  # a Package row, as the coupon document writes one
    book = load_workbook(path)
    cover = book[COVER]
    labels = {cover.cell(r, 1).value: r for r in range(1, cover.max_row + 1)}
    cover.cell(labels["Coupon"], 2).value = coupon
    book.save(path)
    return path


def test_a_returned_form_is_kept_and_listed_on_the_case_page(grips, tmp_path, coupon_folders):
    model, local = coupon_folders
    form, kept, case = case_pages.file_audit("IM7-UNT-L9", returned(tmp_path, grips), "2026-10-24")
    assert kept == local / paths.TEAM_FILES / "2026-10-24" / "2026-10-23_IM7-UNT-L9-audit.xlsx" and kept.is_file()
    assert form.verdict == "Pass" and case.status == "in progress"
    page = (model / "README.md").read_text(encoding="utf-8")
    assert page.index("## Audits") < page.index("## Attempt log") and "| 2026-10-24 | 2026-10-23_IM7-UNT-L9 | 3 | A. Engineer | 2026-10-24 | **Pass** |" in page
    assert case_pages.fill(page, case_pages.page_blocks(model, page), model / "README.md") == page  # the block is current
    assert case_pages.file_audit("IM7-UNT-L9", kept, "2026-10-25")[1] == kept  # a form already kept stays where it is


@pytest.mark.parametrize("kwargs, message", [
    (dict(coupon="IM7-UNC-L9"), "the form is for 'IM7-UNC-L9', not IM7-UNT-L9"),
    (dict(attempt="9"), "attempt '9' is not a row of the case page's attempt log (01 to 03)"),
])
def test_a_form_that_is_not_the_coupons_is_refused(grips, tmp_path, coupon_folders, kwargs, message):
    path = returned(tmp_path, grips, **{k: v for k, v in kwargs.items() if k == "coupon"})
    if "attempt" in kwargs:
        book = load_workbook(path)
        cover = book[COVER]
        row = next(r for r in range(1, cover.max_row + 1) if cover.cell(r, 1).value == "Attempt")
        cover.cell(row, 2).value = kwargs["attempt"]
        book.save(path)
    with pytest.raises(case_pages.CasePageError, match=re.escape(message)):
        case_pages.file_audit("IM7-UNT-L9", path, "2026-10-24")
    assert not (coupon_folders[1] / paths.TEAM_FILES).exists()


def test_a_form_for_another_team_deck_is_refused(grips, tmp_path, coupon_folders):
    verification = coupon_folders[1] / paths.OURS / paths.VERIFICATION
    verification.mkdir(parents=True)
    (verification / paths.MANIFEST).write_text(f'[team]\ndeck = "{GRIPS_FIXTURE.as_posix()}"\n', encoding="utf-8")
    with pytest.raises(case_pages.CasePageError, match="the form's team deck is not coupon-2x8-grips-im7-8552.inp"):
        case_pages.file_audit("IM7-UNT-L9", returned(tmp_path, grips), "2026-10-24")
