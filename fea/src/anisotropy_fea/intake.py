"""The coupon input form (ADR-0046): the Excel template an engineer fills in with the inputs of a coupon test, and its reading into
the two spec files our tools build from.

`write_blank_form` writes the template: one input sheet with a labelled row per input (the team's cells yellow, the fixed words
as lists, the units and what the row means beside each), a "How to fill it in" sheet, a Definitions sheet that explains every term
with an example, a read-only sheet of the library materials the team may pick, and a hidden Record sheet for our tools. The library list and the versions are read at write time, so the form is
generated, never kept as a file in the repository.

`read_form` reads a filled-in form back by its labels, and `read_values` reads the same rows from one JSON object, the values file
the studio sends (ADR-0046 amendment 2); `read_input` takes either. `check_form` turns the cells into typed values and names every
problem at once (ADR-0003); `designation_for` gives the coupon its name by the rules of the coupon tests (one code, one meaning);
`spec_data` builds the spec as a plain dict with our defaults filled in; `render_toml` writes it as the commented TOML the rest of
the tools read; `build_specs` does all of that for the team's build and its CalculiX twin and lets `spec.parse_spec` judge the
exact text that will be written. The form never re-implements a spec rule.

Units on the form are mm, MPa and N; the spec is written in "mm-MPa".
"""

from __future__ import annotations

import json
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from anisotropy_engine.library import LibraryMaterial, library_material, list_datasets

from . import __version__, forms, paths
from .spec import (BENDING_KINDS, ABAQUS_NAME, DEFAULT_NAMES, DESIGNATION, IN_PLANE_KINDS, TEST_CODES, CouponSpec, SpecError, full_angles, load_spec, parse_spec,
                   to_si)
from .units import unit_factors

FORM_VERSION = "0.4"
INPUT, HOWTO, DEFINITIONS, LIBRARY, RECORD = "Coupon", "How to fill it in", "Definitions", "Library materials", "Record"
UNITS = "mm-MPa"
TEST_ORDER = ("unnotched_tension", "compression", "open_hole_tension", "shear_pm45", "three_point_bending", "four_point_bending")  # as the form lists them
assert set(TEST_ORDER) == set(IN_PLANE_KINDS) | set(BENDING_KINDS)
TESTS = {TEST_CODES[kind][1]: kind for kind in TEST_ORDER}  # the test's name on the form -> the spec's kind
YES_NO = ("yes", "no")
CONTROLS = ("displacement", "force")
GRIPS = ("coupling", "nodal")
REGIONS = ("tabs", "line")
PLIES = 16  # rows offered for the ply angles
HOLE_DEFAULTS = {"circumferential": 32, "radial_layers": 6, "ring_ratio": 2.0, "growth": 1.2}  # the public open-hole check's O-grid
ELEMENT_SIZE_DEFAULT = 2.0  # mm, the public checks' row spacing
FPF_CRITERIA = ("max_stress", "tsai_hill")  # what "all" reports when the material has strengths; Tsai-Wu joins when f12* is given
CRITERION_WORDS = ("all", "maximum stress", "tsai-hill", "tsai-wu", "none")  # the Failure criterion row
CRITERION_OF_WORD = {"maximum stress": "max_stress", "tsai-hill": "tsai_hill", "tsai-wu": "tsai_wu"}
WORD_OF_CRITERION = {criterion: word for word, criterion in CRITERION_OF_WORD.items()}
VALUES_SUFFIX = ".json"  # a values file: the form's rows as one JSON object, what the studio sends (ADR-0046 amendment 2)
MATERIAL_CODE = re.compile(r"^[A-Z][A-Z0-9]*$")
LAYUP_CODE = re.compile(r"^L[1-9][0-9]*$")
LIBRARY_HELP = ("pick a material by its id; the sheet Library materials lists each one with its constants and source, and once you pick one "
                "this cell shows its constants")  # shown until a material is picked; an Excel text in a formula holds at most 255 characters
assert len(LIBRARY_HELP) <= 255 and '"' not in LIBRARY_HELP


class IntakeError(ValueError):
    """A form that is not ours, incomplete or inconsistent: every problem, each naming its row."""


# --------------------------------------------------------------------------------------------- the rows of the form

@dataclass(frozen=True)
class Field:
    key: str
    label: str
    kind: str  # "text", "number", "whole", "date", "list", "angle"
    help: str = ""
    required: bool = False
    default: Any = None
    words: tuple[str, ...] = ()  # the list a "list" field offers; the library list is bound to the Library sheet instead
    at_least: float | None = None
    above: float | None = None


BLOCKS: tuple[tuple[str, tuple[Field, ...]], ...] = (
    ("Who fills this in", (
        Field("filled_by", "Filled in by", "text", "your name, so we know whom to ask", required=True),
        Field("date", "Date", "date", "the day you filled it in, as YYYY-MM-DD, e.g. 2026-10-26", required=True),
        Field("notes", "Notes to us", "text", "anything we should know: where a value comes from, which values are estimates, what you want to learn from this coupon"),
    )),
    ("The coupon", (
        Field("test", "Test", "list", "the kind of test: unnotched tension (a plain strip pulled), unnotched compression (a plain strip pushed), "
              "open-hole tension (a strip with a hole at its centre, pulled), in-plane shear (a strip of +45 and -45 plies only, pulled), "
              "three-point bending (a strip on two supports pushed down at midspan) or four-point bending (pushed down at two lines); "
              "see Definitions, Tests", required=True, words=tuple(TESTS)),
        Field("material_code", "Material code", "text", "a short code for the material in the coupon's name, capital letters and digits: X840PW in "
              "X840PW-UNT-L1; leave blank to reuse the code we already use for this material"),
        Field("layup_code", "Layup code", "text", "a short code for the stacking sequence: L1, L2 and so on; leave blank and we reuse the code of this "
              "same stacking sequence or give the next free one"),
        Field("standard", "Standard", "text", "the test standard the specimen follows, e.g. ASTM D3039, or none", default="none"),
        Field("size_from", "Where the size comes from", "text", "where the length, width and tabs come from: a standard, a drawing, or your own design",
              required=True),
        Field("team_model", "Your Abaqus model of this coupon", "text", "the name of your Abaqus model or job for this coupon, or not built yet",
              default="not built yet"),
        Field("test_data", "Test data", "text", "measured results we can compare with (a test report or a file name), or none", default="none"),
    )),
    ("Geometry (mm): the coupon lies with its length along y, the direction it is pulled, pushed or spanned", (
        Field("length", "Length (mm)", "number", "the whole coupon from end to end along the pull, tabs included, e.g. 250; for bending, the span plus the "
              "two overhangs", required=True, above=0),
        Field("width", "Width (mm)", "number", "across the coupon, e.g. 25", required=True, above=0),
        Field("tab_length", "Tab length (mm)", "number", "how much of each end the grip holds, e.g. 56; 0 for no tabs, and then set Grip region to line. "
              "Required for the pulled and pushed tests; leave it blank for bending, which has supports instead of grips", at_least=0),
    )),
    ("Bending (three- and four-point bending only; ignored for the other tests)", (
        Field("span", "Support span (mm)", "number", "the distance between the two support lines, e.g. 48; ASTM D7264 takes 32 times the laminate "
              "thickness. The coupon lies across the supports with its length along y", above=0),
        Field("overhang", "Overhang (mm)", "number", "how far the coupon reaches past each support, so that Length = span + 2 overhang; 0 for a coupon "
              "that ends at its supports", default=0.0, at_least=0),
        Field("load_span", "Loading span (mm)", "number", "four-point bending only: the distance between the two loading noses, centred on the span, "
              "e.g. half the support span", above=0),
    )),
    ("Open hole (open-hole tension only; ignored for the other tests)", (
        Field("hole_diameter", "Hole diameter (mm)", "number", "the diameter of the hole at the centre of the coupon, e.g. 6", above=0),
        Field("hole_circumferential", "Nodes around the hole", "whole", "how many mesh points sit on the edge of the hole; a multiple of 4, at least 8. "
              "More points give a rounder hole and a sharper peak stress", default=HOLE_DEFAULTS["circumferential"], at_least=8),
        Field("hole_radial", "Element rings around the hole", "whole", "how many rings of elements surround the hole before the regular mesh starts; "
              "more rings give more detail near the hole", default=HOLE_DEFAULTS["radial_layers"], at_least=1),
        Field("hole_ring", "Ring ratio", "number", "how far the ringed region reaches from the hole's centre, in hole radii: 2 means a square reaching "
              "2 radii out on each side; it must stay inside the width", default=HOLE_DEFAULTS["ring_ratio"], above=1),
        Field("hole_growth", "Ring growth", "number", "how much thicker each ring is than the one inside it: 1.2 means 20% thicker, so the smallest "
              "elements sit at the edge of the hole, where stress changes fastest; 1 keeps every ring the same", default=HOLE_DEFAULTS["growth"], at_least=1),
    )),
    ("Layup, bottom ply first", tuple(
        Field(f"ply_{i}", f"Ply {i}" + (" (bottom)" if i == 1 else ""), "angle",
              "the fibre angle of the bottom ply in degrees, measured from x, which runs across the width: 0 = fibres across the coupon, "
              "90 = fibres along the pull, 45 and -45 diagonal. Then one ply per row upward; leave the rest blank" if i == 1
              else "the next ply up" if i == 2 else "")
        for i in range(1, PLIES + 1)) + (
        Field("symmetric", "Symmetric", "list", "yes: the plies listed are the lower half and are mirrored, so 0, 45, -45, 90 becomes 8 plies; "
              "no: the plies listed are the whole stack", required=True, words=YES_NO),
        Field("ply_thickness", "Ply thickness (mm)", "number", "the cured thickness of one ply, e.g. 0.127; leave blank only when the library "
              "material gives one (its Ply thickness column on the sheet Library materials)", above=0),
    )),
    ("Material: pick one from our library, or give your own constants", (
        Field("library", "Library material", "list", LIBRARY_HELP),
        Field("own_name", "Own material: name", "text", "the material's name as in your Abaqus model, no spaces, e.g. X840-Z60"),
        Field("own_E1", "Own material: E1 (MPa)", "number", "stiffness along the fibres (Young's modulus in direction 1)", above=0),
        Field("own_E2", "Own material: E2 (MPa)", "number", "stiffness across the fibres, in the plane of the ply (direction 2)", above=0),
        Field("own_G12", "Own material: G12 (MPa)", "number", "in-plane shear stiffness (shear modulus in the 1-2 plane)", above=0),
        Field("own_nu12", "Own material: nu12", "number", "Poisson's ratio: how much the ply narrows across the fibres when stretched along them, "
              "e.g. 0.3; no unit"),
        Field("own_F1t", "Own material: F1t (MPa)", "number", "strength along the fibres in tension. Give all five strengths, or none: without them "
              "we predict no failure", above=0),
        Field("own_F1c", "Own material: F1c (MPa)", "number", "strength along the fibres in compression, as a positive number", above=0),
        Field("own_F2t", "Own material: F2t (MPa)", "number", "strength across the fibres in tension", above=0),
        Field("own_F2c", "Own material: F2c (MPa)", "number", "strength across the fibres in compression, as a positive number", above=0),
        Field("own_F6", "Own material: F6 (MPa)", "number", "in-plane shear strength", above=0),
        Field("own_source", "Own material: source", "text", "where these numbers come from: the supplier's datasheet and its date, or a test report. "
              "They stay with this coupon and never go into our library"),
        Field("G13", "G13 (MPa)", "number", "out-of-plane shear stiffness in the plane of the fibres and the thickness, as in your Abaqus material "
              "(Lamina). No measured value? For a tape use G13 = G12 and say so in the notes; see Definitions, Material", required=True, above=0),
        Field("G23", "G23 (MPa)", "number", "out-of-plane shear stiffness across the fibres and the thickness, as in your Abaqus material (Lamina). "
              "No measured value? For a tape use G23 = E2 / 3 and say so in the notes; see Definitions, Material", required=True, above=0),
    )),
    ("Failure criterion: which first-ply failure criterion we report", (
        Field("criteria", "Failure criterion", "list", "all: every criterion the material allows (maximum stress and Tsai-Hill, and Tsai-Wu when "
              "f12* is given below); or one of them; or none. A criterion needs the material's strengths; see Definitions, First-ply failure",
              default="all", words=CRITERION_WORDS),
        Field("f12_star", "Tsai-Wu interaction f12*", "number", "the normalised interaction term of the Tsai-Wu criterion, a number between -1 and 1, "
              "neither included; required for Tsai-Wu; leave it blank otherwise, and Tsai-Wu is not reported"),
    )),
    ("Loading: the bottom grip is held, the top grip moves; in bending, the supports hold and the nose moves", (
        Field("control", "Loading type", "list", "displacement: the top grip moves by the value, in mm; force: a total force, in N, pulls or pushes "
              "the top grip. Bending takes displacement: the loading nose travels by the value", required=True, words=CONTROLS),
        Field("value", "Value (mm or N)", "number", "the amount as a positive number, e.g. 1 mm or 10000 N; for compression we push, for the other "
              "tests we pull; in bending the nose pushes down by this travel", required=True, above=0),
    )),
    ("Mesh: an element size, or your Abaqus mesh", (
        Field("element_size", "Element size (mm)", "number", "the size of the small pieces (elements) the coupon is divided into for the calculation, "
              "e.g. 2. Smaller is finer: more detail near the grips and a hole, but slower. We put a row of nodes exactly on each tab line",
              default=ELEMENT_SIZE_DEFAULT, above=0),
        Field("mesh_nx", "Your Abaqus mesh: elements across the width", "whole", "only if you already meshed this coupon in Abaqus: the number of "
              "elements across the width. Give this and the next two instead of an element size, and our mesh matches yours node for node", at_least=1),
        Field("mesh_partitions", "Your Abaqus mesh: partition lines (mm)", "text", "where your partition lines cross the length, measured from the "
              "bottom end, comma-separated, e.g. 56, 110, 140, 194"),
        Field("mesh_seeds", "Your Abaqus mesh: rows per segment", "text", "how many rows of elements lie in each segment between those lines, from the "
              "bottom, so one more number than there are lines, e.g. 28, 27, 10, 27, 28"),
    )),
    ("Grips (ignored for bending, which rests on supports)", (
        Field("grips", "Grip kind", "list", "coupling: the gripped nodes at each end are tied to one reference point, held at the bottom and moved at "
              "the top, as in your Abaqus models; nodal: the gripped nodes themselves are held or moved", default="coupling", words=GRIPS),
        Field("region", "Grip region", "list", "tabs: every node on the tab is held; line: only the row of nodes at each end (use it when Tab length "
              "is 0)", default="tabs", words=REGIONS),
    )),
    ("Names in your Abaqus model (optional: give them and our Abaqus file uses them)", (
        Field("name_part", "Abaqus name: part", "text", "the part's name; leave all five blank and we use ours"),
        Field("name_instance", "Abaqus name: instance", "text", "the instance's name in the assembly, e.g. Coupon-1"),
        Field("name_material", "Abaqus name: material", "text", "the material's name under Property, Material"),
        Field("name_layup", "Abaqus name: layup", "text", "the composite layup's name"),
        Field("name_section", "Abaqus name: section set", "text", "the set the layup is assigned to"),
    )),
)
FIELDS: dict[str, Field] = {f.key: f for _, fields in BLOCKS for f in fields}
KEY_OF_LABEL: dict[str, str] = {f.label: f.key for f in FIELDS.values()}
OWN_KEYS = ("own_name", "own_E1", "own_E2", "own_G12", "own_nu12", "own_F1t", "own_F1c", "own_F2t", "own_F2c", "own_F6", "own_source")
STRENGTH_KEYS = ("own_F1t", "own_F1c", "own_F2t", "own_F2c", "own_F6")
NAME_KEYS = {"name_part": "part", "name_instance": "instance", "name_material": "material", "name_layup": "layup", "name_section": "section_set"}

# Which rows of the form a spec rule concerns, when `parse_spec` refuses the rendered spec.
ROWS_OF_KEY = {
    "geometry.length": "length", "geometry.width": "width", "geometry.tab_length": "tab_length", "geometry.hole_diameter": "hole_diameter",
    "geometry.span": "span", "geometry.overhang": "overhang", "geometry.load_span": "load_span",
    "mesh.hole": "hole_circumferential", "mesh.element_size": "element_size", "mesh.nx": "mesh_nx", "mesh.y_partitions": "mesh_partitions",
    "mesh.y_seeds": "mesh_seeds", "mesh.ny": "element_size", "mesh.mesher": "test", "layup.angles": "ply_1", "layup.ply_thickness": "ply_thickness",
    "material.library": "library", "material.E1": "own_E1", "material.E2": "own_E2", "material.G12": "own_G12", "material.nu12": "own_nu12",
    "material.F1t": "own_F1t", "material.F1c": "own_F1c", "material.F2t": "own_F2t", "material.F2c": "own_F2c", "material.F6": "own_F6",
    "material.transverse": "G13", "grips.region": "region", "grips.kind": "grips", "loading.value": "value", "loading.control": "control",
    "names.part": "name_part", "names.instance": "name_instance", "names.material": "name_material", "names.layup": "name_layup",
    "names.section_set": "name_section", "coupon.designation": "material_code", "output.fpf_criteria": "criteria", "material.f12_star": "f12_star",
}

HOW_TO = (
    "What this form is for. You describe one coupon test; from it we build the model in our engine, solve it, and send back a package: "
    "the document, our Abaqus deck of the same coupon, and the audit form on which you compare your Abaqus results with ours.",
    "Fill in the yellow cells only. Rows marked 'required' must be given; the others have a default or may stay blank. Units are mm, MPa and N. "
    "Every row says what it means beside it; the sheet Definitions explains each term with an example.",
    "Directions: the coupon lies with its length along y, the direction it is pulled or pushed, and its width along x. Ply angles are measured "
    "from x, as in your Abaqus models: 0 = fibres across the coupon, 90 = fibres along the pull.",
    "Layup: one ply per row, bottom ply first, angles in degrees. With Symmetric = yes, the plies you list are the lower half and are mirrored.",
    "Material: either pick one from the sheet Library materials, or give your own constants with their source. Not both. G13 and G23 are needed "
    "in both cases, for the shell section of the Abaqus model; take them from your Abaqus material, or estimate them as the sheet Definitions "
    "says and write in the notes that they are estimates. If the library material has no strength in a column, we predict no failure; if it has "
    "no ply thickness, give one.",
    "Failure criterion: all reports every criterion the material's strengths allow, maximum stress and Tsai-Hill, and Tsai-Wu too when you give "
    "its interaction term f12*; or pick one criterion, or none.",
    "Mesh: give an element size and we divide the coupon so that a row of nodes sits on each tab line; or, when you already have an Abaqus model, "
    "give its mesh (elements across the width, the partition lines from the bottom, and the rows per segment) so that ours matches it node for node.",
    "Open hole: the hole rows are read only for open-hole tension. Bending: the span, overhang and loading span rows are read only for the two "
    "bending tests, which rest on two support lines instead of grips. Loading: give the amount; for compression we apply it as a push, and in bending "
    "the nose pushes down by it.",
    "Where we read results: at the centre of the coupon and just past the bottom tab; on an open-hole coupon, beside and above the hole and halfway "
    "to the tab. You do not enter these points: we place them from the geometry.",
    "Names: if you give the names of your Abaqus part, instance, material, layup and section set, our deck opens in Abaqus/CAE with them.",
    "Send the file back as it is. If something is missing or inconsistent, we send it back with every problem named.",
)

# The sheet Definitions: per group, (term, what it means, an example or a tip). Plain words; the questions an engineer asks first.
GLOSSARY: tuple[tuple[str, tuple[tuple[str, str, str], ...]], ...] = (
    ("The coupon and its directions", (
        ("Coupon", "A test specimen: a flat strip of the laminate, held at both ends by grips and pulled or pushed along its length.",
         "A 250 x 25 mm strip with 56 mm tabs."),
        ("x and y", "Lay the coupon flat with its length running up the page. x runs across the width from the left edge (x = 0); y runs along "
         "the length from the bottom end (y = 0). The coupon is pulled or pushed along y.",
         "The centre of a 250 x 25 mm coupon is at x = 12.5, y = 125."),
        ("Tab, tab line", "A tab is the end region a grip holds. The tab line is where the tab ends and the free part of the coupon begins.",
         "With 56 mm tabs on a 250 mm coupon the tab lines are at y = 56 and y = 194."),
        ("Gauge section", "The free part between the tabs, where the coupon should fail and where results are read.", ""),
    )),
    ("Tests", (
        ("Unnotched tension", "A plain strip pulled along its length. It gives the laminate's tensile stiffness and strength.", "ASTM D3039."),
        ("Unnotched compression", "A plain strip pushed along its length. It gives the compressive stiffness and strength. Our model is flat and "
         "in-plane, so it does not buckle.", "ASTM D6641."),
        ("Open-hole tension", "A strip with a hole drilled through its centre, pulled. The hole concentrates the stress at its edge.",
         "ASTM D5766: 36 mm wide with a 6 mm hole."),
        ("In-plane shear", "A strip made of +45 and -45 plies only, pulled. The plies are loaded in shear, which gives the shear stiffness and "
         "strength.", "ASTM D3518. Every ply +45 or -45, as many of each."),
        ("Three-point bending", "A strip resting on two support lines, pushed down at midspan by a loading nose that moves by the value you give. It "
         "gives the flexural stiffness and strength; the outer plies carry most of the stress, the bottom ones in tension.",
         "ASTM D7264 procedure A: 13 mm wide, span 32 times the thickness."),
        ("Four-point bending", "A strip on two support lines, pushed down by two loading noses centred on the span. Between the noses the moment "
         "is constant, so the failure is not under a nose.", "ASTM D7264 procedure B: the loading span is half the support span."),
        ("Support span, overhang, loading span", "The span is the distance between the support lines; the overhang is what reaches past each "
         "support (Length = span + 2 overhang); the loading span is the distance between the two noses of four-point bending.",
         "Span 48, overhang 6: a 60 mm coupon."),
    )),
    ("Layup and plies", (
        ("Ply", "One layer of the laminate: fibres in resin, all at one angle.", ""),
        ("Ply angle", "The direction of a ply's fibres, in degrees, measured from x (across the width), positive counterclockwise when you look "
         "down on the coupon. 0 = fibres across the coupon, 90 = fibres along the pull, 45 and -45 diagonal. This is how your Abaqus models are "
         "set up: the coupon drawn along y, the layup in the global axes.",
         "If you think of 0 as the direction of the pull, add 90 to every angle: 0, 45, -45, 90 becomes 90, -45, 45, 0."),
        ("Bottom ply first", "List the plies from the bottom face up, in the order you would lay them on the tool.", ""),
        ("Symmetric", "The stack is mirrored about its middle. With Symmetric = yes, list only the lower half.",
         "0, 45, -45, 90 with yes is 8 plies: 0, 45, -45, 90, 90, -45, 45, 0."),
        ("Ply thickness", "The cured thickness of one ply.", "0.127 mm (0.005 in)."),
    )),
    ("Material", (
        ("Library material", "A material whose constants we hold with their source and page. Pick its id; the sheet Library materials lists each "
         "one with its constants, and the row on the form shows the one you picked.", "im7-8552: IM7/8552 carbon/epoxy tape, NCAMP."),
        ("Empty cell on the sheet Library materials", "The source does not give that value. With no strength we cannot predict when a ply fails, "
         "so we leave failure out; with no ply thickness, give one on the form.", ""),
        ("Own material", "Your own constants, with where they come from. They stay with this coupon and never go into our library.", ""),
        ("Directions 1, 2, 3", "The ply's own directions: 1 along the fibres, 2 across the fibres in the plane of the ply, 3 through the thickness.",
         ""),
        ("E1, E2", "Stiffness (Young's modulus) along the fibres (E1) and across them (E2).", "IM7/8552 in our library: E1 = 162,096 MPa, "
         "E2 = 8,963 MPa."),
        ("G12", "In-plane shear stiffness: how the ply resists shear in its own plane.", "IM7/8552: G12 = 4,688 MPa."),
        ("nu12", "Poisson's ratio: how much the ply narrows across the fibres when it is stretched along them. No unit.", "IM7/8552: 0.316."),
        ("Strengths F1t, F1c, F2t, F2c, F6", "The stress at which a ply fails: along the fibres in tension (F1t) and compression (F1c), across the "
         "fibres in tension (F2t) and compression (F2c), and in in-plane shear (F6). Compression as a positive number. Give all five or none.", ""),
        ("Failure criterion", "The rule that says when a ply has failed, from its stresses and strengths. Maximum stress: a ply fails when one "
         "stress reaches its strength. Tsai-Hill and Tsai-Wu: the stresses are combined, so no single one governs. All reports every criterion "
         "the material's strengths allow.", "all, our default; or one of maximum stress, tsai-hill, tsai-wu; or none."),
        ("f12*", "The interaction term of the Tsai-Wu criterion, normalised, between -1 and 1 (neither included). Tsai-Wu needs it; the other "
         "criteria do not.", "Leave it blank and Tsai-Wu is not reported."),
        ("G13, G23", "Out-of-plane shear stiffness: G13 in the plane that holds the fibres and the thickness, G23 across the fibres and the "
         "thickness. Abaqus needs them for the shell section (its Lamina material asks for them). Our own solver does not use them, so they do "
         "not change our results for a thin coupon pulled or pushed.",
         "No measured value: for a tape the usual estimate is G13 = G12 and G23 = E2 / 3 (the ply taken as the same in every direction across "
         "the fibres, with nu23 = 0.5; MIL-HDBK-17-3F, section 10.2.4.1). For a woven fabric there is no reliable estimate: ask the supplier. "
         "Either way, write in the notes which values are estimates."),
        ("Source", "Where your own constants come from.", "The supplier's datasheet and its date, or a test report and its number."),
    )),
    ("Loading and grips", (
        ("Displacement", "The top grip moves by a set distance; the bottom grip is held. We report the force it takes.", "1 mm."),
        ("Force", "A set total force pulls or pushes the top grip; the bottom grip is held. We report how far it moves.", "10000 N."),
        ("Coupling grip", "The gripped nodes at each end are tied to one reference point, which moves like a rigid grip: held at the bottom, moved "
         "at the top. This is how your Abaqus models grip the coupon.", ""),
        ("Nodal grip", "The gripped nodes are held or moved directly, without a reference point.", ""),
        ("Grip region", "tabs: every node on the tab is held. line: only the row of nodes at each end.", "Use line when Tab length is 0."),
    )),
    ("Mesh", (
        ("Mesh, element, node", "For the calculation the coupon is divided into small four-sided pieces called elements; their corners are the "
         "nodes. The stress is worked out in every element. The whole set of elements is the mesh.",
         "A 250 x 20 mm coupon with 2 mm elements has 10 across and 125 along: 1,250 elements."),
        ("Element size", "About how long an element's side is. We divide the width, and each part of the length, into whole numbers of elements "
         "of about this size, with a row of nodes exactly on each tab line.", "2 mm, our default; 1 mm for a finer check."),
        ("Finer mesh", "Smaller elements. They show more detail where the stress changes quickly (at the end of a tab, around a hole) and take "
         "longer to solve. In the middle of a plain coupon the stress is the same everywhere, so a finer mesh changes nothing there.", ""),
        ("Your Abaqus mesh", "If you have meshed this coupon in Abaqus, give its seeds instead of an element size: elements across the width, the "
         "partition lines from the bottom, and the rows in each segment. Our mesh is then the same as yours node for node, and the results can "
         "be compared element by element.", "10 across; lines at 56, 110, 140, 194; rows 28, 27, 10, 27, 28."),
        ("Mesh around a hole", "Nodes around the hole: points on the hole's edge. Element rings: rings of elements around the hole. Ring ratio: how "
         "far the ringed region reaches, in hole radii. Ring growth: how much thicker each ring is than the one inside it.",
         "Our defaults: 32 nodes, 6 rings, ratio 2, growth 1.2."),
    )),
    ("What we send back", (
        ("Reaction force", "The force the grip needs to make the movement you gave.", ""),
        ("First-ply failure", "The load at which the first ply anywhere in the coupon reaches its strength, by the criterion you chose on the "
         "Failure criterion row; given as a fraction of your load, with the place and the ply.", "0.93: the first ply fails at 93% of a 1 mm pull."),
        ("Where we read results", "At the centre of the coupon and just past the bottom tab; on an open-hole coupon, beside and above the hole and "
         "halfway to the tab. We place these points from the geometry; you do not enter them.", ""),
        ("Coupon name", "Material code, test code and layup code. Test codes: UNT unnotched tension, UNC unnotched compression, OHT open-hole "
         "tension, IPS in-plane shear.", "X840PW-UNT-L1: X840 plain weave, unnotched tension, layup 1."),
    )),
)


# --------------------------------------------------------------------------------------------- the library sheet

LIBRARY_COLUMNS = ("Id", "Name", "Dataset", "E1 (MPa)", "E2 (MPa)", "G12 (MPa)", "nu12", "F1t (MPa)", "F1c (MPa)", "F2t (MPa)", "F2c (MPa)",
                   "F6 (MPa)", "Ply thickness (mm)", "Source", "Rights")


def library_records() -> list[LibraryMaterial]:
    """Every library material the form may offer: those whose primary dataset gives the elastic constants, by id."""
    ids = sorted({d["material_id"] for d in list_datasets() if d["primary"] and d["elastic_inputs"]})
    records = []
    for material_id in ids:
        try:
            records.append(library_material(material_id))
        except ValueError:  # a record the engine cannot resolve to one value per constant (a ranged thickness): a spec could not use it either
            continue
    return records


def library_row(lib: LibraryMaterial) -> tuple:
    mpa = lambda pa: None if pa is None else pa / 1e6  # noqa: E731
    s = lib.lamina.strengths
    return (lib.material_id, lib.name, lib.dataset_id, mpa(lib.lamina.E1), mpa(lib.lamina.E2), mpa(lib.lamina.G12), lib.lamina.nu12,
            *(mpa(getattr(s, role)) if s is not None else None for role in ("F1t", "F1c", "F2t", "F2c", "F6")),
            None if lib.ply_thickness_m is None else lib.ply_thickness_m * 1e3, lib.citation, lib.rights)


# --------------------------------------------------------------------------------------------- the blank form

def write_blank_form(path: Path | str) -> Path:
    """The coupon input form, blank: the input sheet, the instructions, the library materials and the hidden record."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    path = Path(path)
    records = library_records()
    book = Workbook()
    book.properties.creator = "Anisotropy AI"
    book.properties.title = "Coupon input form"
    sheet = book.active
    sheet.title = INPUT
    sheet["A1"] = "Coupon input form"
    sheet["A1"].font = forms.font(14, bold=True)
    sheet.cell(2, 1, "The inputs our engine needs to build and solve your coupon. Fill in the yellow cells and send the file back.").font = forms.font(11, muted=True)
    widths = {1: 44, 2: 30, 3: 76}
    row = 4
    library_cell = None
    for title, fields in BLOCKS:
        for col in (1, 2, 3):
            forms.fill(sheet.cell(row, col), forms.BANNER_FILL)
        sheet.cell(row, 1, title).font = forms.font(11, bold=True)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        row += 1
        for f in fields:
            sheet.cell(row, 1, f.label).font = forms.font(bold=f.required)
            sheet.cell(row, 1).alignment = forms.wrap()  # the label beside the first line of its explanation
            cell = sheet.cell(row, 2, f.default)
            cell.alignment = forms.wrap()
            forms.open_cell(cell)
            if f.kind == "date":
                cell.number_format = "yyyy-mm-dd"
            help_text = ("required" + ("; " if f.help else "") if f.required else "") + f.help
            sheet.cell(row, 3, help_text or None).font = forms.font(muted=True)
            sheet.cell(row, 3).alignment = forms.wrap()
            if f.kind == "list" and f.words:
                forms.listing(sheet, f.words, [cell.coordinate])
            elif f.key == "library":
                library_cell = cell
            elif f.kind in ("number", "whole", "angle"):
                forms.number_rule(sheet, [cell.coordinate], whole=f.kind == "whole", at_least=f.at_least, above=f.above)
            row += 1
        row += 1
    last = row - 1
    if library_cell is not None and records:
        n = len(records) + 1
        ids = f"'{LIBRARY}'!$A$2:$A${n}"
        forms.listing(sheet, ids, [library_cell.coordinate])
        column = lambda letter: f"'{LIBRARY}'!${letter}$2:${letter}${n}"  # noqa: E731
        pick = lambda letter: f"INDEX({column(letter)},MATCH({library_cell.coordinate},{ids},0))"  # noqa: E731
        shown = (f'=IFERROR({pick("B")}&": E1 "&{pick("D")}&", E2 "&{pick("E")}&", G12 "&{pick("F")}&" MPa, nu12 "&{pick("G")}'
                 f'&IF({pick("M")}="","",", ply "&{pick("M")}&" mm"),"{LIBRARY_HELP}")')
        shown_cell = sheet.cell(library_cell.row, 3, shown)
        shown_cell.font = forms.font(muted=True)
        shown_cell.alignment = forms.wrap()
        sheet.row_dimensions[library_cell.row].height = 13.0 * forms.lines_of(LIBRARY_HELP, widths[3]) + 4  # fit_rows skips formulas
    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width
    forms.fit_rows(sheet, {3: widths[3]}, 4, last)
    sheet.freeze_panes = "A4"
    forms.protect(sheet)
    forms.printable(sheet, landscape=False)

    sheet = book.create_sheet(HOWTO)
    sheet["A1"] = HOWTO
    sheet["A1"].font = forms.font(14, bold=True)
    row = 3
    for text in HOW_TO:
        forms.paragraph(sheet, row, text, 2, 120)
        row += 1
    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 80
    forms.protect(sheet)
    forms.printable(sheet, landscape=False)

    write_definitions(book.create_sheet(DEFINITIONS))

    sheet = book.create_sheet(LIBRARY)
    sheet.append(list(LIBRARY_COLUMNS))
    for cell in sheet[1]:
        cell.font = forms.font(bold=True)
        forms.fill(cell, forms.HEADER_FILL)
        cell.alignment = forms.wrap()
    sheet.row_dimensions[1].height = 30
    for lib in records:
        sheet.append(list(library_row(lib)))
    for j, width in enumerate((16, 28, 34, 10, 10, 10, 8, 10, 10, 10, 10, 10, 12, 60, 18), start=1):
        sheet.column_dimensions[get_column_letter(j)].width = width
    for cells in sheet.iter_rows(min_row=2, min_col=4, max_col=13):
        for cell in cells:
            if isinstance(cell.value, float):
                cell.number_format = "0.000" if cell.column in (7, 13) else "0"
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(LIBRARY_COLUMNS))}{len(records) + 1}"
    forms.protect(sheet)
    forms.printable(sheet, landscape=True, title_rows="1:1")

    sheet = book.create_sheet(RECORD)
    for r, (label, value) in enumerate((("Form version", FORM_VERSION), ("Made on", date.today().isoformat()),
                                        ("Engine", f"anisotropy-fea {__version__}"), ("Library materials listed", len(records))), start=1):
        sheet.cell(r, 1, label).font = forms.font(bold=True)
        sheet.cell(r, 2, value)
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 40
    forms.protect(sheet)
    sheet.sheet_state = "hidden"

    forms.one_font(book)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def write_definitions(sheet) -> None:
    """The sheet Definitions: a banner per group of GLOSSARY, then one row per term with what it means and an example."""
    from openpyxl.utils import get_column_letter

    widths = {1: 30, 2: 74, 3: 56}
    sheet["A1"] = DEFINITIONS
    sheet["A1"].font = forms.font(14, bold=True)
    sheet.cell(2, 1, "What each term on the form means, with an example. The row on the sheet Coupon says the same in short.").font = forms.font(11, muted=True)
    for col, heading in enumerate(("Term", "What it means", "Example or tip"), start=1):
        cell = sheet.cell(4, col, heading)
        cell.font = forms.font(bold=True)
        forms.fill(cell, forms.HEADER_FILL)
    row = 5
    for title, terms in GLOSSARY:
        for col in (1, 2, 3):
            forms.fill(sheet.cell(row, col), forms.BANNER_FILL)
        sheet.cell(row, 1, title).font = forms.font(11, bold=True)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        row += 1
        for term, meaning, example in terms:
            sheet.cell(row, 1, term).font = forms.font(bold=True)
            sheet.cell(row, 2, meaning)
            sheet.cell(row, 3, example or None).font = forms.font(muted=True)
            for col in (1, 2, 3):
                sheet.cell(row, col).alignment = forms.wrap()
            row += 1
    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width
    forms.fit_rows(sheet, widths, 5, row - 1)
    sheet.freeze_panes = "A5"
    forms.protect(sheet)
    forms.printable(sheet, landscape=True, title_rows="4:4")


def fill_form(blank: Path | str, values: dict[str, Any], out: Path | str) -> Path:
    """A blank form filled in by field key (an `angles` list spreads over the ply rows): for tests and demonstrations."""
    from openpyxl import load_workbook

    values = dict(values)
    for i, angle in enumerate(values.pop("angles", []), start=1):
        values[f"ply_{i}"] = angle
    book = load_workbook(blank)
    sheet = book[INPUT]
    rows = {KEY_OF_LABEL[forms.text_of(sheet.cell(r, 1).value)]: r for r in range(1, sheet.max_row + 1)
            if forms.text_of(sheet.cell(r, 1).value) in KEY_OF_LABEL}
    for key, value in values.items():
        if key not in rows:
            raise KeyError(f"the form has no row for {key!r}")
        sheet.cell(rows[key], 2).value = value
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)
    return out


# --------------------------------------------------------------------------------------------- a filled-in form

@dataclass(frozen=True)
class FormInput:
    """A filled-in form as read: the cell of every known row by field key, as typed."""

    path: Path
    version: str
    cells: dict[str, Any]

    def label(self, key: str) -> str:
        return FIELDS[key].label


def read_form(path: Path | str) -> FormInput:
    """The cells of a filled-in form, by field key; refused when the file is not a coupon input form."""
    from openpyxl import load_workbook

    path = Path(path)
    try:
        book = load_workbook(path, data_only=True)
    except Exception as exc:  # openpyxl raises several kinds for a file that is not a workbook
        raise IntakeError(f"{path.name}: not a coupon input form ({exc})") from None
    if RECORD not in book.sheetnames or INPUT not in book.sheetnames:
        raise IntakeError(f"{path.name}: not a coupon input form (no {INPUT} and {RECORD} sheets)")
    version = forms.labelled(book[RECORD]).get("Form version", "")
    if not version:
        raise IntakeError(f"{path.name}: not a coupon input form (no form version on its record)")
    sheet = book[INPUT]
    cells: dict[str, Any] = {}
    for r in range(1, sheet.max_row + 1):
        key = KEY_OF_LABEL.get(forms.text_of(sheet.cell(r, 1).value))
        if key is not None:
            value = sheet.cell(r, 2).value
            cells[key] = None if isinstance(value, str) and not value.strip() else value
    missing = [FIELDS[key].label for key in FIELDS if key not in cells]
    if len(missing) > len(FIELDS) // 2:
        raise IntakeError(f"{path.name}: not a coupon input form (its rows are not ours)")
    if missing:
        raise IntakeError(f"{path.name}: rows missing from the form: {', '.join(missing)}")
    return FormInput(path, version, cells)


def read_values(path: Path | str) -> FormInput:
    """The form's rows read from a values file: one JSON object whose keys are the rows' field keys, as the studio sends them.

    `form_version` is required and must be the current form's; `angles` may stand in for the ply rows; a mesh row may hold a list
    where the form holds comma-separated text; a yes-or-no row may hold a boolean; an absent key is a blank cell; an unknown key is
    refused by name (ADR-0003).
    """
    path = Path(path)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise IntakeError(f"{path.name}: not a values file ({exc})") from None
    if not isinstance(data, dict):
        raise IntakeError(f"{path.name}: a values file is one JSON object of the form's rows by field key")
    data = dict(data)
    version = data.pop("form_version", None)
    if version is None:
        raise IntakeError(f"{path.name}: form_version is required; the current form is {FORM_VERSION}")
    if str(version) != FORM_VERSION:
        raise IntakeError(f"{path.name}: form_version {version!r} is not the current form's, {FORM_VERSION}; the rows follow the current definition")
    angles = data.pop("angles", None)
    if angles is not None:
        if not isinstance(angles, list):
            raise IntakeError(f"{path.name}: angles must be a list of ply angles, bottom ply first")
        if len(angles) > PLIES:
            raise IntakeError(f"{path.name}: angles lists {len(angles)} plies; the form holds {PLIES} ply rows")
        if any(f"ply_{i}" in data for i in range(1, PLIES + 1)):
            raise IntakeError(f"{path.name}: give angles or the ply rows (ply_1, ply_2, ...), not both")
        for i, angle in enumerate(angles, start=1):
            data[f"ply_{i}"] = angle
    unknown = sorted(key for key in data if key not in FIELDS)
    if unknown:
        raise IntakeError(f"{path.name}: unknown keys {', '.join(unknown)}; the keys are the form's field keys")
    cells: dict[str, Any] = {}
    for key, f in FIELDS.items():
        value = data.get(key)
        if isinstance(value, list):
            value = ", ".join(str(v) for v in value)
        elif isinstance(value, bool) and f.words == YES_NO:
            value = "yes" if value else "no"
        elif isinstance(value, dict):
            raise IntakeError(f"{path.name}: {key} holds an object; a row holds one value")
        cells[key] = None if value is None or (isinstance(value, str) and not value.strip()) else value
    return FormInput(path, str(version), cells)


def read_input(path: Path | str) -> FormInput:
    """A filled-in form (.xlsx) or a values file (.json), by its suffix."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        return read_form(path)
    if path.suffix.lower() == VALUES_SUFFIX:
        return read_values(path)
    raise IntakeError(f"{path.name}: a coupon arrives as a filled-in form (.xlsx) or a values file (.json)")


def values_of_spec(spec: CouponSpec) -> dict[str, Any]:
    """The form's rows, by field key, that `spec_data` would turn back into this spec: how the studio reopens a coupon to change it.

    Works on a spec as parsed (its own units) or loaded (SI). The rows the spec does not record (who filled it in, when, the notes)
    are blank. A spec the form cannot state, in other units or asking for a set of criteria the row does not offer, is refused.
    """
    if spec.units != UNITS:
        raise IntakeError(f"the form states a coupon in {UNITS}; the spec is in {spec.units}")
    if not (spec.material.transverse.plane_stress and spec.material.transverse.transverse_isotropy):
        raise IntakeError("the form fixes plane_stress = true and transverse_isotropy = true for the out-of-plane constants; the spec states them otherwise")
    if spec.mesh.element_order != 1:
        raise IntakeError(f"the form fixes element_order = 1; the spec asks for order {spec.mesh.element_order}")
    f = unit_factors(spec.units) if spec.si else {kind: 1.0 for kind in unit_factors(spec.units)}
    L, S, F = f["length"], f["stress"], f["force"]
    back = lambda v, k: None if v is None else v / k  # noqa: E731
    g, m, mat, t = spec.geometry, spec.mesh, spec.material, spec.material.transverse
    values: dict[str, Any] = {"form_version": FORM_VERSION, "filled_by": None, "date": None, "notes": None, "test": TEST_CODES[spec.kind][1]}
    if spec.coupon is not None:
        match = DESIGNATION.match(spec.coupon.designation)
        values.update(material_code=match["material"], layup_code=match["layup"], standard=spec.coupon.standard, size_from=spec.coupon.size_from,
                      team_model=spec.coupon.team_model, test_data=spec.coupon.test_data)
    values.update(length=back(g.length, L), width=back(g.width, L), tab_length=back(g.tab_length, L) if g.tab_length is not None else None,
                  hole_diameter=back(g.hole_diameter, L), span=back(g.span, L), overhang=back(g.overhang, L), load_span=back(g.load_span, L))
    if m.hole is not None:
        values.update(hole_circumferential=m.hole.circumferential, hole_radial=m.hole.radial_layers, hole_ring=m.hole.ring_ratio, hole_growth=m.hole.growth)
    values.update(angles=[int(a) if float(a).is_integer() else a for a in spec.layup.angles], symmetric="yes" if spec.layup.symmetric else "no",
                  ply_thickness=None if spec.layup.ply_thickness == "library" else back(spec.layup.ply_thickness, L))
    if mat.library is not None:
        values["library"] = mat.library
    else:
        values.update(own_name=spec.names.material if spec.names is not None else None, own_E1=back(mat.E1, S), own_E2=back(mat.E2, S),
                      own_G12=back(mat.G12, S), own_nu12=mat.nu12, own_F1t=back(mat.F1t, S), own_F1c=back(mat.F1c, S), own_F2t=back(mat.F2t, S),
                      own_F2c=back(mat.F2c, S), own_F6=back(mat.F6, S), own_source=mat.source)
    values.update(G13=back(t.G13, S), G23=back(t.G23, S), f12_star=mat.f12_star)
    strengths = mat.library is not None and library_material(mat.library, mat.dataset).lamina.strengths is not None or mat.F1t is not None
    asked = list(spec.output.fpf_criteria)
    every = [*FPF_CRITERIA, *(["tsai_wu"] if mat.f12_star is not None else [])] if strengths else []
    if asked == every:
        values["criteria"] = "all"
    elif not asked:
        values["criteria"] = "none"
    elif len(asked) == 1:
        values["criteria"] = WORD_OF_CRITERION[asked[0]]
    else:
        raise IntakeError(f"the form's Failure criterion row offers all, one criterion or none; the spec asks for {', '.join(asked)}")
    values.update(control=spec.loading.control, value=abs(back(spec.loading.value, L if spec.loading.control == "displacement" else F)))
    if m.y_seeds is not None:
        values.update(mesh_nx=m.nx, mesh_partitions=[back(p, L) for p in m.y_partitions], mesh_seeds=list(m.y_seeds))
    elif m.element_size is not None:
        values["element_size"] = back(m.element_size, L)
    elif g.tab_length:  # nx and ny with tabs: the same mesh seeded by segment, the tab lines being node rows (validate_spec)
        tab_rows = round(m.ny * g.tab_length / g.length)
        values.update(mesh_nx=m.nx, mesh_partitions=[back(g.tab_length, L), back(g.length - g.tab_length, L)],
                      mesh_seeds=[tab_rows, m.ny - 2 * tab_rows, tab_rows])
    else:  # nx and ny without tabs: an element size that mesh_data rounds back to the same counts, or nothing the form can state
        size = g.length / m.ny
        if max(1, round(g.width / size)) != m.nx:
            raise IntakeError(f"the form states a mesh by element size or by seeds per segment; {m.nx} by {m.ny} elements on this coupon is neither")
        values["element_size"] = back(size, L)
    if spec.grips is not None:
        values.update(grips=spec.grips.kind, region=spec.grips.region)
    if spec.names is not None:
        names = spec.names.model_dump()
        given = {key: names[key] for key in ("part", "instance", "layup", "section_set") if names[key] != DEFAULT_NAMES[key]}
        if mat.library is not None and names["material"] != mat.library.upper():
            given["material"] = names["material"]
        for key, spec_key in NAME_KEYS.items():
            if spec_key in given:
                values[key] = given[spec_key]
    return {key: value for key, value in values.items() if value is not None or key in ("filled_by", "date", "notes")}


def spec_from_values(values: Values, designation: Designation, source_name: str) -> CouponSpec:
    """The spec the team's build would be, in SI, judged by `parse_spec`; nothing is written."""
    main, _ = build_specs(values, designation, source_name)
    return to_si(parse_spec(main, source_name))


@dataclass(frozen=True)
class Values:
    """What the form says, typed and checked: the inputs `spec_data` builds the spec from."""

    filled_by: str
    day: str
    notes: str
    kind: str
    material_code: str | None
    layup_code: str | None
    standard: str
    size_from: str
    team_model: str
    test_data: str
    length: float
    width: float
    tab_length: float
    hole_diameter: float | None
    hole: dict[str, float | int]
    angles: list[float]
    symmetric: bool
    ply_thickness: float | None
    library: str | None
    own: dict[str, Any] | None  # name, E1, E2, G12, nu12, strengths (dict or None), source
    G13: float
    G23: float
    control: str
    value: float  # an amount; the sign follows the test
    element_size: float | None
    mesh_nx: int | None
    partitions: list[float] | None
    seeds: list[int] | None
    grips: str
    region: str
    names: dict[str, str] | None
    criteria: tuple[str, ...] | None = None  # the criteria asked for by name; None = "all", every one the material allows
    f12_star: float | None = None
    span: float | None = None  # bending kinds: the distance between the support lines
    overhang: float | None = None  # bending kinds: what reaches past each support
    load_span: float | None = None  # four-point bending: the distance between the loading noses

    @property
    def test_code(self) -> str:
        return TEST_CODES[self.kind][0]


def check_form(form: FormInput) -> Values:
    """The typed values of a filled-in form; refused with every problem named, each by its row."""
    cells, problems = form.cells, []

    def text(key: str) -> str:
        value = forms.text_of(cells.get(key))
        if not value and FIELDS[key].required:
            problems.append(f"{form.label(key)}: required")
        return value

    def number(key: str, whole: bool = False) -> float | int | None:
        value = cells.get(key)
        f = FIELDS[key]
        if value is None:
            if f.required:
                problems.append(f"{f.label}: required")
            return None
        try:
            number_ = float(str(value).replace(",", ".")) if isinstance(value, str) else float(value)
        except (TypeError, ValueError):
            problems.append(f"{f.label}: a number is needed, not {forms.text_of(value)!r}")
            return None
        if whole and number_ != int(number_):
            problems.append(f"{f.label}: a whole number is needed, not {number_:g}")
            return None
        if f.above is not None and not number_ > f.above:
            problems.append(f"{f.label}: must be greater than {f.above:g}, not {number_:g}")
        if f.at_least is not None and not number_ >= f.at_least:
            problems.append(f"{f.label}: must be at least {f.at_least:g}, not {number_:g}")
        return int(number_) if whole else number_

    def word(key: str) -> str:
        f = FIELDS[key]
        value = forms.text_of(cells.get(key)).lower()
        if not value:
            if f.required:
                problems.append(f"{f.label}: required; one of {', '.join(f.words)}")
            return f.default or ""
        if value not in f.words:
            problems.append(f"{f.label}: {forms.text_of(cells.get(key))!r} is not one of {', '.join(f.words)}")
        return value

    def numbers(key: str, whole: bool = False) -> list | None:
        raw = forms.text_of(cells.get(key))
        if not raw:
            return None
        try:
            parts = [float(p.strip()) for p in raw.replace(";", ",").split(",") if p.strip()]
            if whole and any(p != int(p) for p in parts):
                raise ValueError
            return [int(p) for p in parts] if whole else parts
        except ValueError:
            problems.append(f"{form.label(key)}: comma-separated {'whole ' if whole else ''}numbers are needed, not {raw!r}")
            return None

    filled_by, notes = text("filled_by"), text("notes")
    raw_day = cells.get("date")
    if isinstance(raw_day, (datetime, date)):
        day = forms.text_of(raw_day)
    else:
        day = forms.text_of(raw_day)
        if not day:
            problems.append("Date: required")
        else:
            try:
                day = date.fromisoformat(day).isoformat()
            except ValueError:
                problems.append(f"Date: {day!r} is not a date (write YYYY-MM-DD)")
    test_name = forms.text_of(cells.get("test"))
    kind = TESTS.get(test_name.lower() if test_name else "", "")
    if not kind:
        problems.append(f"Test: {test_name!r} is not one of {', '.join(TESTS)}" if test_name else f"Test: required; one of {', '.join(TESTS)}")
    material_code = text("material_code").upper() or None
    if material_code and not MATERIAL_CODE.match(material_code):
        problems.append(f"Material code: {material_code!r} is not capital letters and digits starting with a letter")
    layup_code = text("layup_code").upper() or None
    if layup_code and not LAYUP_CODE.match(layup_code):
        problems.append(f"Layup code: {layup_code!r} is not L followed by a number, such as L1")
    standard, size_from = text("standard") or "none", text("size_from")
    team_model, test_data = text("team_model") or "not built yet", text("test_data") or "none"

    length, width, tab_length = number("length"), number("width"), number("tab_length")
    bending = kind in BENDING_KINDS
    if tab_length is None and kind and not bending:
        problems.append("Tab length (mm): required for the pulled and pushed tests (0 for no tabs)")
    if bending and tab_length:
        problems.append("Tab length (mm): a bending coupon has supports, not tabs; leave it blank or 0")
    span, overhang, load_span = number("span"), number("overhang"), number("load_span")
    if bending:
        if span is None:
            problems.append("Support span (mm): required for bending")
        overhang = 0.0 if overhang is None else overhang
        if span is not None and length is not None and abs(length - (span + 2 * overhang)) > 1e-9 * length:
            problems.append(f"Length (mm): must equal the support span plus twice the overhang, {span + 2 * overhang:g}, not {length:g}")
        if kind == "four_point_bending":
            if load_span is None:
                problems.append("Loading span (mm): required for four-point bending")
            elif span is not None and not load_span < span:
                problems.append(f"Loading span (mm): must be smaller than the support span {span:g}, not {load_span:g}")
        elif load_span is not None:
            problems.append("Loading span (mm): belongs to four-point bending only")
    elif kind and (span is not None or load_span is not None or overhang):  # the overhang row carries a default of 0 on every form
        problems.append("Support span, Overhang and Loading span: belong to the bending tests only")
    hole_diameter = number("hole_diameter")
    if kind == "open_hole_tension" and hole_diameter is None:
        problems.append("Hole diameter (mm): required for open-hole tension")
    hole = {"circumferential": number("hole_circumferential", whole=True), "radial_layers": number("hole_radial", whole=True),
            "ring_ratio": number("hole_ring"), "growth": number("hole_growth")}
    hole = {k: HOLE_DEFAULTS[k] if v is None else v for k, v in hole.items()}

    angles: list[float] = []
    ended = False
    for i in range(1, PLIES + 1):
        value = cells.get(f"ply_{i}")
        if value is None:
            ended = True
            continue
        if ended:
            problems.append(f"Ply {i}: a blank ply row above it; list the plies without gaps")
            continue
        try:
            angles.append(float(value))
        except (TypeError, ValueError):
            problems.append(f"Ply {i}: an angle in degrees is needed, not {forms.text_of(value)!r}")
    if not angles:
        problems.append("Ply 1 (bottom): at least one ply is needed")
    symmetric = word("symmetric") == "yes"
    ply_thickness = number("ply_thickness")

    library = forms.text_of(cells.get("library")) or None
    own_given = {key: cells.get(key) for key in OWN_KEYS if cells.get(key) is not None}
    own = None
    if library and own_given:
        problems.append("Library material and Own material: give one of the two, not both")
    elif not library and not own_given:
        problems.append("Library material or Own material: one of the two is needed")
    elif own_given:
        own = {"name": text("own_name"), "E1": number("own_E1"), "E2": number("own_E2"), "G12": number("own_G12"), "nu12": number("own_nu12"),
               "source": text("own_source")}
        for key in ("own_name", "own_E1", "own_E2", "own_G12", "own_nu12", "own_source"):
            if cells.get(key) is None:
                problems.append(f"{form.label(key)}: required with your own constants")
        if own["name"] and not ABAQUS_NAME.match(own["name"]):
            problems.append(f"Own material: name: {own['name']!r} must start with a letter and hold only letters, digits, - and _")
        strengths = {key[4:]: number(key) for key in STRENGTH_KEYS}
        given = [k for k, v in strengths.items() if v is not None]
        if given and len(given) != len(STRENGTH_KEYS):
            problems.append(f"Own material strengths: give all five (F1t, F1c, F2t, F2c, F6) or none; given {', '.join(given)}")
        own["strengths"] = strengths if len(given) == len(STRENGTH_KEYS) else None
    G13, G23 = number("G13"), number("G23")

    strengths_known: bool | None = None  # whether the material has strengths, when the form says enough to know
    if library and not own_given:
        try:
            strengths_known = library_material(library).lamina.strengths is not None
        except ValueError as exc:
            problems.append(f"{form.label('library')}: {exc}")
    elif own is not None:
        strengths_known = own["strengths"] is not None
    criterion_word, f12_star = word("criteria"), number("f12_star")
    criteria: tuple[str, ...] | None = None
    if criterion_word == "none":
        criteria = ()
    elif criterion_word in CRITERION_OF_WORD:
        criteria = (CRITERION_OF_WORD[criterion_word],)
    if f12_star is not None and not -1.0 < f12_star < 1.0:
        problems.append(f"{form.label('f12_star')}: must lie strictly between -1 and 1, not {f12_star:g}")
    if criteria and strengths_known is False:
        problems.append(f"{form.label('criteria')}: the material has no strengths, so no criterion can be reported; choose none, pick a library "
                        "material with strengths, or give your own")
    if criteria and "tsai_wu" in criteria and f12_star is None:
        problems.append(f"{form.label('f12_star')}: required for Tsai-Wu")

    control = word("control")
    value = number("value")
    if bending and control == "force":
        problems.append("Loading type: bending takes displacement (the travel of the loading nose); force control is not built")

    element_size = number("element_size")
    mesh_nx = number("mesh_nx", whole=True)
    partitions, seeds = numbers("mesh_partitions"), numbers("mesh_seeds", whole=True)
    abaqus_mesh = [v for v in (mesh_nx, partitions, seeds) if v is not None]
    size_given = cells.get("element_size") is not None
    if abaqus_mesh and size_given:
        problems.append("Element size and Your Abaqus mesh: give one of the two, not both (clear the element size to use your mesh)")
    elif abaqus_mesh and len(abaqus_mesh) != 3:
        problems.append("Your Abaqus mesh: give all three rows (elements across the width, partition lines, rows per segment)")
    elif not abaqus_mesh and not size_given:
        problems.append("Element size (mm): required unless you give your Abaqus mesh")
    if partitions is not None and seeds is not None and len(seeds) != len(partitions) + 1:
        problems.append(f"Your Abaqus mesh: rows per segment: {len(seeds)} counts for {len(partitions)} partition lines; one count per segment, "
                        f"one more than the lines")
    grips, region = word("grips"), word("region")

    names = {spec_key: forms.text_of(cells.get(key)) for key, spec_key in NAME_KEYS.items() if cells.get(key) is not None}
    for key, spec_key in NAME_KEYS.items():
        if spec_key in names and not ABAQUS_NAME.match(names[spec_key]):
            problems.append(f"{form.label(key)}: {names[spec_key]!r} must start with a letter and hold only letters, digits, - and _")
    if problems:
        raise IntakeError(f"{form.path.name}: " + "; ".join(problems))
    return Values(filled_by, day, notes, kind, material_code, layup_code, standard, size_from, team_model, test_data, length, width, tab_length or 0.0,
                  hole_diameter, hole, angles, symmetric, ply_thickness, library, own, G13, G23, control, value, element_size if size_given else None,
                  mesh_nx, partitions, seeds, grips, region, names or None, criteria, f12_star,
                  span=span if bending else None, overhang=overhang if bending else None, load_span=load_span if bending else None)


# --------------------------------------------------------------------------------------------- the coupon's name

@dataclass(frozen=True)
class Designation:
    material_code: str
    test_code: str
    layup_code: str

    @property
    def name(self) -> str:
        return f"{self.material_code}-{self.test_code}-{self.layup_code}"


def material_meaning(values: Values) -> tuple:
    """What a material code stands for: a library record (and its primary dataset), or one set of own constants, in SI."""
    if values.library:
        return ("library", values.library, library_material(values.library).dataset_id)
    own = values.own
    return ("own", *(round(own[k] * 1e6, 3) for k in ("E1", "E2", "G12")), round(own["nu12"], 6))


def spec_material_meaning(spec) -> tuple:
    mat = spec.material
    if mat.library is not None:
        return ("library", mat.library, mat.dataset or library_material(mat.library).dataset_id)
    return ("own", round(mat.E1, 3), round(mat.E2, 3), round(mat.G12, 3), round(mat.nu12, 6))


def existing_coupons(root: Path) -> dict[str, Any]:
    """Every coupon under `root` with a spec, by name (the rule of paths.model_dirs on another root), loaded in SI."""
    root = Path(root)
    if not root.is_dir():
        return {}
    return {d.name: load_spec(d / "ours" / f"{d.name}.toml") for d in sorted(root.iterdir()) if d.is_dir() and (d / "ours" / f"{d.name}.toml").is_file()}


def designation_for(values: Values, root: Path) -> Designation:
    """The coupon's name, <material code>-<test code>-<layup code>, by the rules every coupon on this machine follows: a code
    means one thing and one thing has one code. A blank code is reused from the coupons that exist, or, for a layup, proposed."""
    problems = []
    coupons = existing_coupons(root)
    materials: dict[str, tuple] = {}
    layups: dict[str, tuple] = {}
    for spec in coupons.values():
        match = DESIGNATION.match(spec.coupon.designation) if spec.coupon else None
        if match is None:
            continue
        materials.setdefault(match.group("material"), spec_material_meaning(spec))
        layups.setdefault(match.group("layup"), tuple(full_angles(spec)))
    meaning = material_meaning(values)
    code = values.material_code
    if code is None:
        code = next((c for c, m in materials.items() if m == meaning), None)
        if code is None:
            problems.append("Material code: this material has no code yet; give one (capital letters and digits, e.g. IM7)")
    else:
        if code in materials and materials[code] != meaning:
            problems.append(f"Material code: {code} already names another material on this machine; give another code")
        other = next((c for c, m in materials.items() if m == meaning and c != code), None)
        if other is not None:
            problems.append(f"Material code: this material already has the code {other}; use it")
    stacking = tuple(float(a) for a in (values.angles + values.angles[::-1] if values.symmetric else values.angles))
    layup = values.layup_code
    if layup is None:
        layup = next((c for c, angles in layups.items() if angles == stacking), None)
        if layup is None:
            n = 1
            while f"L{n}" in layups:
                n += 1
            layup = f"L{n}"
    else:
        if layup in layups and layups[layup] != stacking:
            problems.append(f"Layup code: {layup} already names another stacking sequence on this machine; give another code")
        other = next((c for c, angles in layups.items() if angles == stacking and c != layup), None)
        if other is not None:
            problems.append(f"Layup code: this stacking sequence already has the code {other}; use it")
    if problems:
        raise IntakeError("; ".join(problems))
    return Designation(code, values.test_code, layup)


# --------------------------------------------------------------------------------------------- the spec

def mesh_data(values: Values) -> tuple[dict, dict[str, str]]:
    """The [mesh] table: the team's Abaqus seeds, or seeds from an element size with a node row on each tab line."""
    comments: dict[str, str] = {}
    if values.kind == "open_hole_tension":
        size = values.element_size if values.element_size is not None else ELEMENT_SIZE_DEFAULT
        data = {"mesher": "gmsh", "element_order": 1, "element_size": size, "hole": dict(values.hole)}
        comments["mesh.element_size"] = "the blocks around the hole; the square and the rings follow [mesh.hole]"
        return data, comments
    data: dict = {"mesher": "structured", "element_order": 1}
    if values.seeds is not None:
        data.update(nx=values.mesh_nx, y_partitions=values.partitions, y_seeds=values.seeds)
        comments["mesh.y_partitions"] = "the team's Abaqus mesh, seeded by segment as they seeded it"
        return data, comments
    size, L, w, tab = values.element_size, values.length, values.width, values.tab_length
    rows = lambda length: max(1, round(length / size))  # noqa: E731
    data["nx"] = rows(w)
    if values.kind in BENDING_KINDS:  # a node row on each support line and each nose line
        lines = sorted({y for y in bending_lines_of(values) if 0.0 < y < L})
        edges = [0.0, *lines, L]
        data.update(y_partitions=lines, y_seeds=[rows(b - a) for a, b in zip(edges, edges[1:])])
        comments["mesh.y_seeds"] = f"about {size:g} mm rows, with a node row on each support line and each loading-nose line"
        return data, comments
    if tab > 0:
        data.update(y_partitions=[tab, L - tab], y_seeds=[rows(tab), rows(L - 2 * tab), rows(tab)])
        comments["mesh.y_seeds"] = f"about {size:g} mm rows, with a node row on each tab line"
    else:
        data["ny"] = rows(L)
        comments["mesh.ny"] = f"about {size:g} mm rows"
    return data, comments


def bending_lines_of(values: Values) -> list[float]:
    """The support lines and the loading-nose line(s) of a bending coupon, as y from the bottom end (what spec.node_lines will compute)."""
    L = values.length
    lines = [values.overhang, L - values.overhang]
    if values.kind == "three_point_bending":
        lines.append(L / 2)
    else:
        lines += [L / 2 - values.load_span / 2, L / 2 + values.load_span / 2]
    return lines


def probes_data(values: Values) -> dict:
    """Where the results are read: the gauge centre, a point near the grip, and beside a hole when there is one; on a bending coupon the
    centre element under the nose and a quarter-span node and element."""
    L, w, tab = values.length, values.width, values.tab_length
    x, y = w / 2, L / 2
    if values.kind in BENDING_KINDS:
        quarter = (values.overhang + y) / 2
        return {"nodes": {"quarter": [x, quarter]}, "elements": {"centre": [x, y], "quarter": [x, quarter]}}
    if values.kind == "open_hole_tension":
        r = values.hole_diameter / 2
        far_y = (tab + y) / 2
        return {"nodes": {"centre_top": [x, y + r], "hole_side": [x + r, y], "gauge": [x, far_y]},
                "elements": {"hole_side": [x + r + 0.2, y], "hole_top": [x, y + r + 0.2], "far": [x, far_y]}}
    near = min(tab + 1.0, y) if tab > 0 else min(1.0, y)
    return {"nodes": {"centre": [x, y]}, "elements": {"centre": [x, y], "near_grip": [x, near]}}


def gauge_probes(kind: str) -> tuple[tuple[str, ...], tuple[str, ...]]:
    """The node and element probes a CalculiX record keeps: those in the gauge. Beside a tab line or a hole the bilinear kernel element
    and S8R differ in the small components, so those probes are left out of the record, as the coupons' records were made (ADR-0049)."""
    if kind in BENDING_KINDS:  # the element beside the nose line differs on its secondary components (the record reading says how much)
        return (("quarter",), ("quarter",))
    return (("gauge",), ("far",)) if kind == "open_hole_tension" else (("centre",), ("centre",))


def material_data(values: Values) -> tuple[dict, dict[str, str], str, bool]:
    """The [material] table, its comments, the material's name for the title, and whether it has strengths."""
    comments = {"material.transverse.plane_stress": "nu13 = nu23 = 0 for CalculiX, the plane stress Abaqus LAMINA and our engine assume",
                "material.transverse.transverse_isotropy": "E3 = E2: a stated choice, used only by CalculiX",
                "material.transverse.G13": "as given on the form", "material.transverse.G23": "as given on the form"}
    transverse = {"G13": values.G13, "G23": values.G23, "plane_stress": True, "transverse_isotropy": True}
    if values.library:
        lib = library_material(values.library)
        data = {"library": lib.material_id, "dataset": lib.dataset_id, "transverse": transverse}
        comments["material.library"] = lib.citation.replace(" · ", "; ")
        return data, comments, lib.name, lib.lamina.strengths is not None
    own = values.own
    data = {k: own[k] for k in ("E1", "E2", "G12", "nu12")}
    if own["strengths"] is not None:
        data.update(own["strengths"])
    data["source"] = own["source"]
    data["transverse"] = transverse
    comments["material.source"] = "the team's own constants, as given on the form; they stay in this folder"
    return data, comments, own["name"], own["strengths"] is not None


def spec_data(values: Values, designation: Designation, grips_kind: str, form_name: str) -> tuple[dict, dict[str, str], tuple[str, ...]]:
    """The spec as a plain dict in the order the file is written, its comments by dotted key, and its header lines."""
    mesh, comments = mesh_data(values)
    material, material_comments, material_name, strengths = material_data(values)
    comments.update(material_comments)
    kind = values.kind
    bending = kind in BENDING_KINDS
    test_name = TEST_CODES[kind][1]
    sign = -1.0 if kind == "compression" or bending else 1.0  # a push, and the nose moving down onto the coupon, are negative
    layup = "[" + "/".join(f"{a:g}" for a in values.angles) + "]" + ("s" if values.symmetric else "")
    if bending:
        grip_words = f"{values.span:g} mm span" + (f", {values.load_span:g} mm loading span" if kind == "four_point_bending" else "")
        load_words = f"nose{'s' if kind == 'four_point_bending' else ''} pushed down {values.value:g} mm"
    else:
        grip_words = (f"{'coupled' if grips_kind == 'coupling' else 'nodal'} grips on {values.tab_length:g} mm tabs" if values.region == "tabs"
                      else f"{'coupled' if grips_kind == 'coupling' else 'nodal'} grips on the end lines")
        load_words = (f"{values.value:g} mm {'push' if kind == 'compression' else 'pull'}" if values.control == "displacement"
                      else f"{values.value:g} N {'compression' if kind == 'compression' else 'tension'}")
    hole_words = f", {values.hole_diameter:g} mm hole in {values.width:g} mm" if kind == "open_hole_tension" else ""
    title = f"{designation.name}, {test_name}: {material_name} {layup}{hole_words}, {grip_words}, {load_words}"
    if bending:
        geometry: dict = {"length": values.length, "width": values.width, "span": values.span, "overhang": values.overhang}
        if kind == "four_point_bending":
            geometry["load_span"] = values.load_span
    else:
        geometry = {"length": values.length, "width": values.width, "tab_length": values.tab_length}
    if kind == "open_hole_tension":
        geometry["hole_diameter"] = values.hole_diameter
    data: dict = {
        "title": title, "kind": kind, "units": UNITS,
        "coupon": {"designation": designation.name, "standard": values.standard, "size_from": values.size_from, "team_model": values.team_model,
                   "test_data": values.test_data},
        "geometry": geometry, "mesh": mesh,
        "layup": {"angles": [int(a) if float(a).is_integer() else a for a in values.angles], "symmetric": values.symmetric,
                  "ply_thickness": values.ply_thickness if values.ply_thickness is not None else "library"},
        "material": material,
        **({} if bending else {"grips": {"kind": grips_kind, "region": values.region}}),
        "loading": {"control": values.control, "value": sign * values.value},
        "probes": probes_data(values),
    }
    if values.names or values.own:  # the team's own material always carries the name they gave it into the deck
        deck_material = values.own["name"] if values.own else values.library.upper()
        given = {**DEFAULT_NAMES, "material": deck_material, **(values.names or {})}
        data["names"] = {key: given[key] for key in ("part", "instance", "material", "layup", "section_set")}
    if values.f12_star is not None:
        material["f12_star"] = values.f12_star
        comments["material.f12_star"] = "the Tsai-Wu interaction term, as given on the form"
    if values.criteria is None:  # "all": every criterion the material allows
        criteria = [*FPF_CRITERIA, *(["tsai_wu"] if values.f12_star is not None else [])] if strengths else []
    else:
        criteria = list(values.criteria)
    data["output"] = {"fpf_criteria": criteria,
                      "abaqus": {"field_preselect": True, "node_print": True, "el_print": True},
                      "calculix": {"node_print": True, "el_print": True, "frd": True, "el_print_global": True}}
    comments["layup.angles"] = "bottom ply first; the lower half when symmetric"
    comments["layup.ply_thickness"] = "from the library record" if values.ply_thickness is None else "as given on the form"
    if not bending:
        comments["grips.kind"] = ("a reference point at each end drives the whole tab, as in the team's decks" if grips_kind == "coupling"
                                  else "the grip nodes held directly: the CalculiX twin, which reports no reaction at a coupled reference node")
    comments["loading.value"] = ("the nose moves down onto the coupon: negative along z" if bending
                                 else "a push: negative along the coupon" if kind == "compression" else "along the coupon")
    if bending:
        comments["geometry.span"] = "the distance between the support lines, as given on the form; w is held along each line"
    comments["output.fpf_criteria"] = (("every criterion the material allows" if values.criteria is None else "as chosen on the form") if criteria
                                       else ("no strengths given, so no failure criterion" if not strengths else "none chosen on the form"))
    header = (f"{designation.name}, {test_name}: built from the coupon input form {form_name}, filled in by {values.filled_by} on {values.day}.",
              *([f"Notes from the team: {values.notes}"] if values.notes else []),
              "Units mm-MPa. Our defaults (element order, plane stress, probes, output requests) are commented where they stand.")
    return data, comments, header


INLINE = ("nodes", "elements", "abaqus", "calculix")  # nested tables written on one line, as every spec in this project writes them


def toml_value(value) -> str:
    if isinstance(value, str):
        return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (list, tuple)):
        return "[" + ", ".join(toml_value(v) for v in value) + "]"
    if isinstance(value, dict):
        return "{ " + ", ".join(f"{k} = {toml_value(v)}" for k, v in value.items()) + " }"
    if isinstance(value, float):
        text = f"{value:.9g}"
        return text if "." in text or "e" in text else text + ".0"
    return str(value)


def render_toml(data: dict, comments: dict[str, str] | None = None, header: tuple[str, ...] = ()) -> str:
    """The spec dict as the commented TOML our tools read: scalars first, then one table per key, sub-tables for nested tables,
    and the usual inline tables for probes and output requests. `comments` are keyed by dotted path."""
    comments = comments or {}

    def line(path: str, key: str, value) -> str:
        text = f"{key} = {toml_value(value)}"
        return f"{text}  # {comments[path]}" if path in comments else text

    out = [f"# {h}" for h in header]
    for key, value in data.items():
        if not isinstance(value, dict):
            out.append(line(key, key, value))

    def table(path: str, body: dict) -> None:
        out.append("")
        out.append(f"[{path}]")
        for key, value in body.items():
            if isinstance(value, dict) and key not in INLINE:
                continue
            out.append(line(f"{path}.{key}", key, value))
        for key, value in body.items():
            if isinstance(value, dict) and key not in INLINE:
                table(f"{path}.{key}", value)

    for key, value in data.items():
        if isinstance(value, dict):
            table(key, value)
    return "\n".join(out) + "\n"


def form_rows_of(message: str) -> list[str]:
    """The form rows a spec message concerns, by the table.key names it mentions."""
    rows = []
    for key in re.findall(r"\b(?:geometry|mesh|layup|material|grips|loading|probes|names|coupon|output)\.[A-Za-z0-9_]+", message):
        field_key = ROWS_OF_KEY.get(key) or ROWS_OF_KEY.get(key.rsplit(".", 1)[0])
        if field_key and FIELDS[field_key].label not in rows:
            rows.append(FIELDS[field_key].label)
    return rows


def build_specs(values: Values, designation: Designation, form_name: str) -> tuple[str, str]:
    """The team's build and its CalculiX twin as TOML texts, each judged by `parse_spec` before it is returned."""
    texts = []
    for grips_kind, stem in ((values.grips, designation.name), ("nodal", f"{designation.name}-nodal")):
        data, comments, header = spec_data(values, designation, grips_kind, form_name)
        text = render_toml(data, comments, header)
        try:
            parse_spec(text, f"{stem}.toml")
        except SpecError as exc:
            rows = form_rows_of(str(exc))
            raise IntakeError(f"{form_name}: the spec refuses the inputs: {exc}" + (f" (form rows: {', '.join(rows)})" if rows else "")) from None
        texts.append(text)
    return texts[0], texts[1]


# --------------------------------------------------------------------------------------------- filing a form

@dataclass(frozen=True)
class Intake:
    """What reading a filled-in form did: the coupon's name and folder, the two spec texts, the files written, and what to run next."""

    designation: Designation
    folder: Path
    main: str
    twin: str
    written: tuple[Path, ...]
    commands: tuple[str, ...]


def next_commands(folder: Path, name: str, kind: str) -> tuple[str, ...]:
    """What to run from fea/ once the specs exist, with the coupon's paths filled in; the attempt folder is the one new_attempt.py prints.
    `tools/run_coupon.py` runs the same steps in one go."""
    import os

    rel = Path(os.path.relpath(folder, paths.FEA)).as_posix()
    if rel.startswith("../../"):  # a root outside the repository (a scratch folder): its full path reads better than a chain of ..
        rel = Path(folder).resolve().as_posix()
    ours = f"{rel}/{paths.OURS}"
    run = f"{ours}/{paths.RUNS}/attempt-NN_<date>/run"
    nodes, elements = gauge_probes(kind)
    return (
        f"uv run python tools/regenerate_decks.py {name} {name}-nodal --ccx",
        f"uv run python tools/new_attempt.py {name}   # prints the folder attempt-NN_<date>",
        f"uv run python examples/build_model.py {ours}/{name}.toml --out {run}/kernel",
        f"uv run python examples/run_calculix.py {ours}/{name}-nodal.toml --out {run}/calculix",
        f'uv run python tools/record_from_dat.py {ours}/{name}-nodal.toml {run}/calculix/{name}-nodal-ccx.dat --version 2.22 --run-by "..." '
        f'--machine "..." --tolerance 0.03 --floor 0.02 --node-probes {" ".join(nodes)} --element-probes {" ".join(elements)} --notes "..." '
        f'--out {ours}/{paths.VERIFICATION}/{name}-calculix.toml',
        f"uv run --project fea python docs/models/coupon.py {name}   # from the repository root",
    )


def material_words(values: Values) -> str:
    """The material as the READMEs name it, with its rights class, as the workspace rules ask."""
    if values.library:
        lib = library_material(values.library)
        return f"{lib.name} (library record {lib.material_id}, dataset {lib.dataset_id}; rights class {lib.rights})"
    own = values.own
    return (f"the team's own constants, named {own['name']} (source: {own['source']}; rights class user-supplied: private, kept in this folder and "
            "never published)")


def size_words(values: Values) -> str:
    size = f"{values.length:g} mm long, {values.width:g} mm wide"
    if values.tab_length:
        size += f", {values.tab_length:g} mm tabs at each end"
    if values.kind in BENDING_KINDS:
        size += f", {values.span:g} mm span" + (f", {values.load_span:g} mm loading span" if values.kind == "four_point_bending" else "")
    if values.kind == "open_hole_tension":
        size += f", central hole of {values.hole_diameter:g} mm"
    return size


def team_readme(values: Values, designation: Designation, form_name: str, day: str) -> str:
    """The README of the dated team folder: who made the file, when, what it is, the material and its rights class."""
    lines = [
        f"# {designation.name}: the coupon input form, received {day}",
        "",
        f"Filled in by {values.filled_by} on {values.day}; read on {day} by `fea/tools/coupon_from_form.py` into `ours/{designation.name}.toml` "
        f"and `ours/{designation.name}-nodal.toml`.",
        "",
        f"Material: {material_words(values)}. The team's Abaqus model of this coupon: {values.team_model}. Test data: {values.test_data}.",
    ]
    if values.notes:
        lines += ["", f"Notes from the team: {values.notes}"]
    what = ("the coupon's values as the studio sent them: the form's rows as one JSON object" if form_name.lower().endswith(VALUES_SUFFIX)
            else "the coupon input form as received")
    lines += ["", "| File | What it is |", "| --- | --- |", f"| `{form_name}` | {what} (form version {FORM_VERSION}) |", ""]
    return "\n".join(lines)


def case_page(values: Values, designation: Designation, form_name: str, day: str, commands: tuple[str, ...]) -> str:
    """A case page to start from: the case line, the problem, an empty result block, the attempt log's first row, where to look, the
    rebuild commands. The status and the accepted attempt are written by hand from here on."""
    name, test = designation.name, TEST_CODES[values.kind][1]
    layup = "[" + "/".join(f"{a:g}" for a in values.angles) + "]" + ("s" if values.symmetric else "")
    material = material_words(values).split(" (")[0]
    where = [
        (f"`team/{day}/`", "the coupon input form as received", "the team"),
        (f"`ours/{name}.toml`", f"the spec of the team's build; its deck `ours/decks/{name}-abaqus.inp` is the one to open in Abaqus/CAE",
         "`fea/tools/coupon_from_form.py`"),
        (f"`ours/{name}-nodal.toml`", "the CalculiX twin: the same grips as direct constraints on the tab nodes" if values.kind not in BENDING_KINDS
         else "the CalculiX twin: the same model (a bending coupon has no couplings to replace)", "`fea/tools/coupon_from_form.py`"),
        ("`ours/decks/`", "the decks written from the two specs; never edited by hand", "`fea/tools/regenerate_decks.py`"),
        ("`ours/verification/`", "the CalculiX record and its reading; the team comparison once the team has built the coupon", "us"),
        ("`ours/runs/`", "`attempt-NN_<date>/run/`, NN being the row of the attempt log", "`fea/tools/new_attempt.py`, then our tools"),
        (f"`documents/<date>_{name}/`", "the packages that go out, one per delivery", "`docs/models/coupon.py`"),
    ]
    parts = [
        f"# {name}, {test} (`{name}`)",
        "",
        "**Case.** Category: team test | Status: in progress | Accepted attempt: 01",
        "",
        f"The team's {test} coupon {name}: {size_words(values)}; layup {layup}; material {material}. Stated on the coupon input form "
        f"`{form_name}` filled in by {values.filled_by} on {values.day}; the specs are written from it and nothing has run yet.",
        "",
        f"**Coupon.** Standard: {values.standard}. The size comes from {values.size_from}. The team's model: {values.team_model}. Test data: "
        f"{values.test_data}.",
        "",
        "**Status.** In progress. Verified means our engine agrees with another solver on this model; validated means it agrees with test data.",
        "",
        "## Result",
        "",
        "<!-- generated: result -->",
        "<!-- end generated -->",
        "",
        "## Attempt log",
        "",
        "Every try, failed ones included, oldest first; this table is the record of each one. A try that ran something keeps its runs in "
        "`ours/runs/attempt-NN_<date>/run/`, NN being its row here; start that folder with `tools/new_attempt.py`.",
        "",
        "| # | Date | What was tried | Result | Outcome | Evidence |",
        "| --- | --- | --- | --- | --- | --- |",
        f"| 01 | {day} | The team's input form `{form_name}` read into the specs `{name}.toml` and `{name}-nodal.toml`; nothing run yet | -- | "
        f"to run | `team/{day}/` |",
        "",
        "## Where to look",
        "",
        "Three parts beside this page: what the team sent, what we built, and what goes out.",
        "",
        "| Where | What it holds | Made by |",
        "| --- | --- | --- |",
        *(f"| {a} | {b} | {c} |" for a, b, c in where),
        "",
        "## Rebuild",
        "",
        "From `fea/`:",
        "",
        "```sh",
        *commands[:-1],
        "```",
        "",
        f"From the repository root, the coupon document and its package: `{commands[-1].split('   #')[0]}`.",
        "",
    ]
    return "\n".join(parts)


def convert(form_path: Path | str, *, root: Path | str | None = None, day: str | None = None, replace: bool = False, dry_run: bool = False) -> Intake:
    """Read a filled-in form, or a values file, and start its coupon under `root` (default .local/fea): the two specs in ours/, the
    input as received in team/<day>/ with a README, and a case page when there is none. Specs that exist are refused unless
    `replace`, which rewrites the two specs and files the input it read them from, so that every attempt has its input; the README
    and the case page are never rewritten. With `dry_run` nothing is written."""
    form_path = Path(form_path)
    root = Path(root) if root is not None else paths.LOCAL
    day = day or date.today().isoformat()
    date.fromisoformat(day)
    form = read_input(form_path)
    values = check_form(form)
    designation = designation_for(values, root)
    main, twin = build_specs(values, designation, form_path.name)
    name = designation.name
    folder = root / name
    ours = folder / paths.OURS
    specs = {ours / f"{name}.toml": main, ours / f"{name}-nodal.toml": twin}
    existing = [p.name for p in specs if p.exists()]
    if existing and not replace:
        raise IntakeError(f"{name}: {', '.join(existing)} already exist under {folder}; read the form again with --replace to rewrite the two specs")
    commands = next_commands(folder, name, values.kind)
    if dry_run:
        return Intake(designation, folder, main, twin, (), commands)
    written = []
    ours.mkdir(parents=True, exist_ok=True)
    for path, text in specs.items():
        path.write_text(text, encoding="utf-8", newline="\n")
        written.append(path)
    received = folder / paths.TEAM_FILES / day
    copy = received / form_path.name
    if copy.resolve() != form_path.resolve():
        received.mkdir(parents=True, exist_ok=True)
        copy.write_bytes(form_path.read_bytes())
        written.append(copy)
    if not existing:
        readme = received / "README.md"
        if not readme.exists():
            received.mkdir(parents=True, exist_ok=True)
            readme.write_text(team_readme(values, designation, form_path.name, day), encoding="utf-8", newline="\n")
            written.append(readme)
        page = folder / "README.md"
        if not page.exists():
            page.write_text(case_page(values, designation, form_path.name, day, commands), encoding="utf-8", newline="\n")
            written.append(page)
    return Intake(designation, folder, main, twin, tuple(written), commands)
