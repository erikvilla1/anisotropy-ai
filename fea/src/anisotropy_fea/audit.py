"""The coupon audit (ADR-0043): one rubric that every audit of a coupon package is scored against.

The rubric (`paths.RUBRIC`) lists each check once: what is checked, what passes, who judges it (our tools or the auditor),
the tolerance, whether a Pass verdict needs it, the tests it applies to and what the team must send for it. This module reads
it and refuses one that is incomplete or inconsistent (ADR-0003): a missing field, an item numbered out of order or under an
unknown section, a test code `spec.TEST_CODES` does not know, or a function in `source` that does not exist.

`score` gives each item of a package its result from what our tools found (`Evidence`): a result for the items our tools
judge, and the evidence the auditor reads for the rest, one fact per line. `write_form` writes the audit form, the Excel workbook
the auditor fills in: a cover that says what is audited and which files are compared and takes the verdict, the rubric with the
computed results locked and the auditor's Result, Observation and Cause open, the words of each answer offered as a list, the
verdict the rule gives as a formula that follows the auditor's entries, from `form_tables` the numbers behind our rows as values
on a Results sheet and every compared value on a Compared values sheet, and last a Record sheet of how the package was built
(ADR-0044 and its amendment 1). A row that fails is the finding: there is no findings sheet (ADR-0043 amendment 3). Evidence is
the claim the auditor verifies, one fact per line; where the team's export lacks a value, the row states ours and is the auditor's
(ADR-0043 amendment 4).

The words an audit form uses are fixed here, so the form, the document and the reader of a returned form agree on them.
"""

from __future__ import annotations

import hashlib
import importlib
import re
import tomllib
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Callable, Literal

import numpy as np
from pydantic import BaseModel, ConfigDict, Field, PositiveFloat, ValidationError, model_validator

from . import paths
from .equivalence import DeckEquivalence, EquivalenceItem
from .mesh import KernelError
from .model import Deck, NodeKey, all_boundaries, expanded_layers, resolve_nodes
from .cases import element_name
from .report import LaminateCheck, gripped_groups, laminate_check
from .results import MembraneSolution
from .side_by_side import PairedValue, agreement, ply_values
from .forms import FONT, fit_rows, labelled as _labelled, lines_of, listing, one_font, printable, protect, text_of as _text  # noqa: F401  the house style of the forms
from .spec import TEST_CODES
from .units import from_si

RESULTS = ("Pass", "Fail", "N/A", "Not checked")  # an item's result on the form
VERDICTS = ("Pass", "Pass with observations", "Fail")  # the audit's verdict
CAUSES = ("engine", "our input", "their model", "documentation", "tolerance")  # what a failed row traces to, and so who fixes it
ITEM_ID = re.compile(r"^(?P<section>[A-Z])(?P<number>[1-9][0-9]*)$")


class RubricError(ValueError):
    """The rubric is incomplete or inconsistent."""


class Record(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")


class Tolerance(Record):
    """Exactly one of: a fraction (of the value or largest magnitude the acceptance names), or a length in metres."""

    relative: PositiveFloat | None = None
    length_m: PositiveFloat | None = None

    @model_validator(mode="after")
    def one_kind(self) -> Tolerance:
        if (self.relative is None) == (self.length_m is None):
            raise ValueError("give exactly one of relative or length_m")
        return self


class Section(Record):
    id: str
    title: str = Field(min_length=1)
    meaning: str = ""  # what the section checks, in the words the form's cover gives the team


class Condition(Record):
    """What a coupon must satisfy for an item to apply to it; otherwise the item is N/A, with the reason stated."""

    min_free_length_widths: PositiveFloat  # the free length between the grips, in widths of the coupon


class Item(Record):
    id: str
    section: str
    check: str = Field(min_length=1)
    acceptance: str = Field(min_length=1)
    method: Literal["auto", "auditor"]
    required: bool
    source: tuple[str, ...] = ()  # "module.function" or "module.Class.method" in anisotropy_fea
    tolerance: Tolerance | None = None
    applies_to: tuple[str, ...] = ()  # test codes; empty: every test
    condition: Condition | None = None  # when set, a coupon that does not meet it gets N/A on this item
    needs: str | None = Field(default=None, min_length=1)


class Rubric(Record):
    version: str = Field(min_length=1)
    status: Literal["draft", "confirmed"]
    sections: tuple[Section, ...] = Field(alias="section", min_length=1)
    items: tuple[Item, ...] = Field(alias="item", min_length=1)


def parse_rubric(text: str, source_path: str | None = None) -> Rubric:
    """The rubric in `text`, checked; `source_path` only names it in errors."""
    where = source_path or "rubric"
    try:
        rubric = Rubric(**tomllib.loads(text))
    except (tomllib.TOMLDecodeError, ValidationError) as exc:
        raise RubricError(f"{where}: {exc}") from None
    validate_rubric(rubric, where)
    return rubric


def load_rubric(path: Path | str = paths.RUBRIC) -> Rubric:
    path = Path(path)
    return parse_rubric(path.read_text(encoding="utf-8"), str(path))


def validate_rubric(rubric: Rubric, where: str = "rubric") -> None:
    """Sections named once; items numbered 1, 2, ... within their section, in file order; known test codes and functions."""

    def need(condition: bool, message: str) -> None:
        if not condition:
            raise RubricError(f"{where}: {message}")

    section_ids = [s.id for s in rubric.sections]
    need(all(re.fullmatch(r"[A-Z]", s) for s in section_ids), f"section ids are single capital letters, not {section_ids}")
    need(len(set(section_ids)) == len(section_ids), f"a section id is used twice: {section_ids}")
    codes = {code for code, _ in TEST_CODES.values()}
    numbers: dict[str, int] = {}
    for item in rubric.items:
        match = ITEM_ID.match(item.id)
        need(match is not None, f"item id {item.id!r}: write the section letter and a number, such as A1")
        need(match.group("section") == item.section, f"item {item.id} sits under section {item.section!r}; its id names {match.group('section')}")
        need(item.section in section_ids, f"item {item.id}: section {item.section!r} is not one of {section_ids}")
        expected = numbers.get(item.section, 0) + 1
        need(int(match.group("number")) == expected, f"item {item.id}: section {item.section} is numbered from 1 in file order; expected {item.section}{expected}")
        numbers[item.section] = expected
        unknown = sorted(set(item.applies_to) - codes)
        need(not unknown, f"item {item.id}: test codes {unknown} are not in spec.TEST_CODES ({sorted(codes)})")
        need(len(set(item.applies_to)) == len(item.applies_to), f"item {item.id}: a test code is listed twice")
        for name in item.source:
            need(_resolves(name), f"item {item.id}: source {name!r} is not a function of anisotropy_fea (write module.function or module.Class.method)")
    empty = [s for s in section_ids if s not in numbers]
    need(not empty, f"sections {empty} have no items")


def _resolves(name: str) -> bool:
    """`module.function` or `module.Class.method` (a property counts), the module one of anisotropy_fea's."""
    module, *attributes = name.split(".")
    if not attributes:
        return False
    try:
        found = importlib.import_module(f"anisotropy_fea.{module}")
    except ImportError:
        return False
    for attribute in attributes:
        found = getattr(found, attribute, None)
    return callable(found) or isinstance(found, property)


def items_for(rubric: Rubric, test_code: str) -> tuple[Item, ...]:
    """The items an audit of a coupon with this test code scores, in rubric order."""
    if test_code not in {code for code, _ in TEST_CODES.values()}:
        raise RubricError(f"test code {test_code!r} is not in spec.TEST_CODES")
    return tuple(item for item in rubric.items if not item.applies_to or test_code in item.applies_to)


# --------------------------------------------------------------------------------------------- scoring

@dataclass(frozen=True)
class Evidence:
    """What our tools found for one coupon package. Lengths and stresses in the stated `units`, forces in N.

    The team's side is empty when the team has not built the coupon: the model and results match cannot then be checked.
    """

    test_code: str
    units: str  # "mm-MPa" or "m-Pa"
    card: tuple[tuple[str, str], ...]  # spec.coupon_card
    our_deck: Deck  # SI
    solution: MembraneSolution  # our kernel on our deck
    material: str | None = None  # the library record the spec names, in words
    team_deck: Deck | None = None  # SI
    team_deck_sha256: str | None = None  # compare.sha256_of; on the form's Record sheet
    equivalence: DeckEquivalence | None = None
    paired: tuple[PairedValue, ...] = ()  # the team's reported values beside ours
    inner_band: tuple[float, float] | None = None  # the graded band: element centroids with lo <= y <= hi
    free_band: tuple[float, float] | None = None  # the free length between the grips
    excluded: str | None = None  # how the regions left out of the grading were set, in words
    team_reaction: float | None = None  # N, at the loaded grip
    fpf: tuple[str, float, int, float] | None = None  # our first-ply failure: criterion, factor of the applied loading, ply (1 = bottom), its angle
    rebuilt: DeckEquivalence | str | None = None  # a deck the auditor built from our guide against ours, or why the two could not be compared
    guide: tuple[int, int] | None = None  # the rebuild guide's open items and ASSUMED tags
    team_reports: tuple[str, ...] = ()  # the file names of the team's field reports, for the "How to check" lines


@dataclass(frozen=True)
class Found:
    """What our tools found for one row, in the three columns of the form: the team's values, ours, and the difference between
    them. Each is a tuple of lines; the lines of the three columns are written one per row where they correspond."""

    team: tuple[str, ...] = ()
    ours: tuple[str, ...] = ()
    difference: tuple[str, ...] = ()

    @property
    def text(self) -> str:
        """The three columns as one text, each line prefixed by its column, for a reader without the sheet."""
        return "\n".join([*(f"Team: {t}" for t in self.team), *(f"Ours: {o}" for o in self.ours), *(f"Difference: {d}" for d in self.difference)])


@dataclass(frozen=True)
class Scored:
    item: Item
    result: str | None  # one of RESULTS for an item our tools judge; None where the auditor judges (an auditor item, or one whose team value is missing)
    found: Found


@dataclass(frozen=True)
class _Facts:
    """Computed once per package for several items."""

    laminate: LaminateCheck | None
    laminate_reason: str  # why there is no laminate check, when there is none
    angles: tuple[float, ...]
    length_unit: str
    stress_unit: str
    reaction_N: float | None = None  # our reaction along the coupon at the top (loaded) grip
    centre: tuple[int, str] | None = None  # the element nearest the coupon's centre: mesh index and name


# Which of deck_equivalence's items (by the whole of their name) answer each B item.
EQUIVALENCE_GROUPS = {
    "B1": r"Nodes|Node positions",
    "B2": r"Elements|Element connectivity",
    "B3": r"Layup, bottom first",
    "B4": r"Section|Ply thickness|Integration points per ply|Shell orientation",
    "B5": r"Material (E1|E2|nu12|G12|G13|G23)( \(.+\))?",
    "B7": r"Couplings|Coupling(,| at) .+",
    "B8": r"Boundary conditions, node by node|Concentrated loads",
    "B9": r"Steps",
}
LISTED_ONLY = ("Part / instance names", "Material names", "Section set / layup names", "Strengths (*Fail Stress)")  # shown, not judged
NOT_BUILT_BY_TEAM = "The team has not built this coupon, so there is no model of theirs to compare"
TEAM_BLANK = "left blank, for the team to fill in"  # the Team column of a row whose value the team has not exported
UNIT_WORDS = {"mm-MPa": ("mm and MPa", "mm, MPa and N"), "m-Pa": ("m and Pa", "m, Pa and N")}  # a deck's units; the audit's
DOF_NAMES = {1: "U1", 2: "U2", 3: "U3", 4: "UR1", 5: "UR2", 6: "UR3"}  # Abaqus/CAE's names for the degrees of freedom


def equivalence_group(equivalence: DeckEquivalence, item_id: str) -> list[EquivalenceItem]:
    return [i for i in equivalence.items if re.fullmatch(EQUIVALENCE_GROUPS[item_id], i.item)]


def _pct(fraction: float) -> str:
    return f"{100 * fraction:.3g} %"


def _layup_text(angles: tuple[float, ...], full: bool = False) -> str:
    """Whether the layup is balanced and symmetric; with `full`, the layup itself first."""
    symmetric = list(angles) == list(reversed(angles))
    off_axis = sorted({abs(a) for a in angles if a % 90 != 0})
    unbalanced = [a for a in off_axis if angles.count(a) != angles.count(-a)]
    balance = ("balanced" if not unbalanced else "unbalanced (" + "; ".join(f"{angles.count(a)} plies at +{a:g}, {angles.count(-a)} at -{a:g}"
                                                                             for a in unbalanced) + ")")
    text = f"{balance}, {'symmetric' if symmetric else 'not symmetric'}"
    return f"[{'/'.join(f'{a:g}' for a in angles)}], {len(angles)} plies, {text}" if full else text


def _deck_angles(deck: Deck) -> tuple[float, ...]:
    return tuple(layer.angle_deg for layer in expanded_layers(next(iter(deck.parts.values())).sections[0]))


def _node_where(deck: Deck, key: NodeKey, units: str, length_unit: str) -> str:
    """A held node in words: a reference point by its position, an instance node by its label and position."""
    instance, label = key
    if instance is None:
        x, y = (from_si(float(v), "length", units) for v in deck.assembly_nodes.xyz[deck.assembly_nodes.index_of(label)][:2])
        return f"reference point at ({x:g}, {y:g}) {length_unit}"
    part = deck.parts[deck.instances[instance].part]
    rows = np.flatnonzero(part.nodes.labels == label)
    if rows.size == 0:
        return f"node {label}"
    x, y = (from_si(float(v), "length", units) for v in part.nodes.xyz[int(rows[0])][:2])
    return f"node {label} at ({x:g}, {y:g}) {length_unit}"


def _dof_text(held: dict[int, float], units: str, length_unit: str) -> str:
    """'U2 = 10 mm; U1 = U3 = UR1 = UR2 = UR3 = 0': the prescribed dofs grouped by value, the moved ones first."""
    by_value: dict[str, list[str]] = {}
    for dof, value in sorted(held.items()):
        shown = f"{from_si(value, 'length', units):g} {length_unit}" if dof <= 3 and value != 0 else f"{value:g}"
        by_value.setdefault(shown, []).append(DOF_NAMES.get(dof, f"dof {dof}"))
    return "; ".join(sorted((f"{' = '.join(names)} = {shown}" for shown, names in by_value.items()), key=lambda text: text.endswith("= 0")))


def _boundary_lines(deck: Deck, units: str, length_unit: str) -> list[str]:
    """B8's column for one deck: a line per group of nodes held the same way, then the concentrated loads."""
    held: dict[NodeKey, dict[int, float]] = {}
    sets: dict[NodeKey, set[str]] = {}
    for boundary in all_boundaries(deck):
        for key in resolve_nodes(deck, boundary.target, boundary.line):
            for dof in range(boundary.dof_first, boundary.dof_last + 1):
                held.setdefault(key, {})[dof] = boundary.value
            if isinstance(boundary.target, str):
                sets.setdefault(key, set()).add(boundary.target)
    groups: dict[tuple, list[NodeKey]] = {}
    for key, dofs in held.items():
        groups.setdefault(tuple(sorted(dofs.items())), []).append(key)
    lines = []
    for signature, keys in groups.items():
        names = sorted({name for key in keys for name in sets.get(key, ())})
        where = _node_where(deck, keys[0], units, length_unit) if len(keys) == 1 else f"{len(keys)} nodes"
        via = f" (set{'s' if len(names) > 1 else ''} {', '.join(names)})" if names else ""
        lines.append(f"{where[0].upper() + where[1:]}{via}: {_dof_text(dict(signature), units, length_unit)}")
    loads: dict[tuple, int] = {}
    for step in deck.steps:
        for load in step.cloads:
            keys = resolve_nodes(deck, load.target, load.line)
            where = _node_where(deck, keys[0], units, length_unit) if len(keys) == 1 else f"{len(keys)} nodes"
            via = f" (set {load.target})" if isinstance(load.target, str) else ""
            loads[(load.dof, load.value, where + via)] = loads.get((load.dof, load.value, where + via), 0) + 1
    if not loads:
        lines.append("Concentrated loads: none")
    for (dof, value, where), count in loads.items():
        lines.append(f"Concentrated load {value:g} N on {DOF_NAMES.get(dof, f'dof {dof}')} at {where}" + (f", {count} times" if count > 1 else ""))
    return lines


def _short(item: str) -> str:
    """The item's name as the Difference column refers to it: 'Material E1 (MPa)' is 'E1'."""
    name = item.split(" (")[0]
    return name[9:] if name.startswith("Material ") else name


def _model_match(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
    if ev.equivalence is None:
        return "Not checked", Found(team=(NOT_BUILT_BY_TEAM,))
    group = equivalence_group(ev.equivalence, item.id)
    if not group:
        return "N/A", Found(difference=("Neither deck has one",) if item.id == "B7" else ("deck_equivalence reports no such item",))
    team, ours, difference = [], [], []
    for i in group:
        if i.item == "Node positions":
            offset = from_si(ev.equivalence.largest_node_offset, "length", ev.units)
            difference.append("Node positions: the same, node for node" if i.same
                              else f"Node positions: differ; the largest offset between matched nodes is {offset:.3g} {facts.length_unit}")
        elif i.item == "Element connectivity":
            difference.append("Element connectivity: the same, element for element" if i.same else f"Element connectivity: differs; ours {i.ours}")
        else:
            label = i.item if item.id != "B5" else _short(i.item) + (" (" + i.item.split(" (")[1] if " (" in i.item else "")
            team.append(f"{label}: {i.team}")
            ours.append(f"{label}: {i.ours}")
            difference.append(f"{_short(i.item)}: {'the same' if i.same else 'differs'}")
    differ = [i for i in group if i.same is False]
    result = "Fail" if differ else "Pass"
    if item.id == "B3":
        team_angles = _deck_angles(ev.team_deck)
        team.append(f"Balance: {_layup_text(team_angles)}")
        ours.append(f"Balance: {_layup_text(facts.angles)}")
        difference.append("Balance: the same" if _layup_text(team_angles) == _layup_text(facts.angles) else "Balance: differs")
    if item.id == "B5" and differ:
        difference.append("Largest relative difference: " + _largest_material_difference(ev.team_deck, ev.our_deck))
    if item.id == "B8":  # the founder's request: the team's boundary conditions and loads and ours, in full, side by side
        team = _boundary_lines(ev.team_deck, ev.units, facts.length_unit)
        ours = _boundary_lines(ev.our_deck, ev.units, facts.length_unit)
    return result, Found(tuple(team), tuple(ours), tuple(difference))


def _largest_material_difference(team: Deck, ours: Deck) -> str:
    """The constant whose ply values differ most, relative to ours: the size a six-digit table hides."""
    plies = lambda deck: expanded_layers(next(iter(deck.parts.values())).sections[0])  # noqa: E731
    worst = (0.0, "")
    for name in ("E1", "E2", "nu12", "G12", "G13", "G23"):
        for a, b in zip(plies(team), plies(ours)):
            x, y = getattr(team.materials[a.material].elastic, name), getattr(ours.materials[b.material].elastic, name)
            if x is not None and y is not None and y != 0:
                worst = max(worst, (abs(x - y) / abs(y), name))
    return f"{worst[0]:.2g} ({worst[1]})"


def _card(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
    if not ev.card:
        return "Fail", Found(ours=("The spec has no [coupon] table",))
    blank = [label for label, value in ev.card if not value.strip()]
    lines = [f"{label}: {value}" for label, value in ev.card if value.strip()]
    return ("Fail" if blank else "Pass"), Found(ours=tuple(lines), difference=("Blank: " + ", ".join(blank),) if blank else ())


def _units(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
    words = lambda units: UNIT_WORDS.get(units, (str(units), str(units)))  # noqa: E731
    team = (words(ev.team_deck.units)[0],) if ev.team_deck is not None else ()
    ours = (words(ev.our_deck.units)[0], f"This audit: {words(ev.units)[1]}")
    decks = [d for d in (ev.team_deck, ev.our_deck) if d is not None]
    ok = all(d.units == ev.units for d in decks)
    return ("Pass" if ok else "Fail"), Found(team, ours, ("The same",) if ok else ("Differ",))


def _our_largest_in_band(quantity: str, ev: Evidence, facts: _Facts) -> str | None:
    """Our largest |value| of a quantity over the graded band, by element and ply, where the team reported none."""
    if ev.inner_band is None:
        return None
    try:
        values = ply_values(ev.solution, ev.units)[quantity]
    except KernelError:
        return None
    mesh = ev.solution.mesh
    centres = from_si(mesh.xy[mesh.connectivity[:, :4]].mean(axis=1), "length", ev.units)
    lo, hi = ev.inner_band
    inside = np.flatnonzero((centres[:, 1] >= lo) & (centres[:, 1] <= hi))
    if inside.size == 0:
        return None
    e, k = np.unravel_index(np.argmax(np.abs(values[inside])), values[inside].shape)
    e = int(inside[e])
    return (f"{quantity}: largest {abs(values[e, k]):.6g} {facts.stress_unit} in the band, at our element {mesh.element_keys[e][1]} "
            f"({centres[e, 0]:g}, {centres[e, 1]:g}), ply {int(k) + 1}")


def _stresses(quantities: tuple[str, ...]) -> Callable[[Item, Evidence, _Facts], tuple[str | None, Found]]:
    def scorer(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
        if ev.equivalence is None:
            return "Not checked", Found(team=(NOT_BUILT_BY_TEAM,))
        if ev.inner_band is None:
            return "Not checked", Found(difference=("No graded band is set for this coupon",))
        reported = [q for q in quantities if any(p.quantity == q for p in ev.paired)]
        if not reported:
            return "Not checked", Found(team=(f"No field report of {', '.join(quantities)} in the export",))
        lo, hi = ev.inner_band
        layers = sorted({p.layer for p in ev.paired if p.quantity in reported})
        team = [f"Graded band y {lo:g} to {hi:g} {facts.length_unit}; layers {', '.join(map(str, layers))}"]
        ours = [f"The same elements and layers, from our deck"]
        difference = [f"Tolerance {_pct(item.tolerance.relative)} of the team's largest value in the band"]
        failed = False
        for q in reported:
            a = agreement(list(ev.paired), q, "inner", ev.inner_band, item.tolerance.relative)
            if a.count == 0:
                team.append(f"{q}: no reported element in the band")
                ours.append(f"{q}: --")
                difference.append(f"{q}: not compared")
                failed = True
                continue
            failed |= a.within != a.count
            w = a.worst
            in_band = [p for p in ev.paired if p.quantity == q and lo <= p.y <= hi]
            ours_largest = max(in_band, key=lambda p: abs(p.ours))
            team.append(f"{q}: {a.count} values; largest {a.largest:.6g} {facts.stress_unit}")
            ours.append(f"{q}: {a.count} values; largest {abs(ours_largest.ours):.6g} {facts.stress_unit}")
            difference.append(f"{q}: {a.within} of {a.count} within {_pct(a.fraction)}; worst {_pct(a.worst_fraction)} at team element "
                              f"{w.reference_element}, layer {w.layer}")
        missing = [q for q in quantities if q not in reported]
        for q in missing:
            team.append(f"{q}: not in the reports, {TEAM_BLANK}")
            ours.append(_our_largest_in_band(q, ev, facts) or f"{q}: --")
            difference.append(f"{q}: not compared; the Result waits until the team has verified {q} against our value")
        # a reported quantity outside the tolerance fails the row; a quantity the team has not exported leaves the row to the auditor
        return ("Fail" if failed else None if missing else "Pass"), Found(tuple(team), tuple(ours), tuple(difference))
    return scorer


def _reaction(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
    ours = (f"{abs(facts.reaction_N):.6g} N along the coupon at the top grip",) if facts.reaction_N is not None else ("Not computed: no gripped group found",)
    if ev.team_reaction is None:
        return None, Found(team=(TEAM_BLANK,), ours=ours)
    team = (f"{abs(ev.team_reaction):.6g} N",)
    if facts.reaction_N is None:
        return "Not checked", Found(team, ours)
    difference = abs(abs(facts.reaction_N) - abs(ev.team_reaction)) / abs(ev.team_reaction)
    return (("Pass" if difference <= item.tolerance.relative else "Fail"),
            Found(team, ours, (f"{_pct(difference)} apart; tolerance {_pct(item.tolerance.relative)}",)))


def _node_text(mesh, index: int) -> str:
    instance, label = mesh.node_keys[index]
    return f"node {label}" if instance is not None else f"reference node {label}"


def _displacement(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
    """C4: our displacements and the strain at the gauge centre, for the team to set its own beside."""
    mesh, u = ev.solution.mesh, from_si(ev.solution.displacement, "length", ev.units)
    i1, i2 = int(np.argmax(np.abs(u[:, 0]))), int(np.argmax(np.abs(u[:, 1])))
    lines = [f"Largest displacement: U1 {u[i1, 0]:.4g} {facts.length_unit} at {_node_text(mesh, i1)}; U2 {u[i2, 1]:.4g} {facts.length_unit} "
             f"at {_node_text(mesh, i2)}"]
    if facts.centre is not None:
        e, name = facts.centre
        eps = ev.solution.strain[e]
        lines.append(f"Strain at the gauge centre (element {name}): eps_x {eps[0]:.4g}, eps_y {eps[1]:.4g}, gamma_xy {eps[2]:.4g}")
    return None, Found(team=(TEAM_BLANK,), ours=tuple(lines))


def _fpf(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
    if ev.fpf is None:
        return None, Found(team=(TEAM_BLANK,), ours=("Not computed: the spec names no failure criterion",))
    criterion, factor, ply, angle = ev.fpf
    return None, Found(team=(TEAM_BLANK,), ours=(f"The first ply to fail is ply {ply} ({angle:g} deg) by {criterion.replace('_', ' ')}, at {factor:.3g} of "
                                                "the applied loading",))


def laminate_distance(check: LaminateCheck) -> tuple[float, float]:
    """How far the kernel's centre element is from free-edge laminate theory: the largest strain difference as a share of the largest
    laminate-theory strain, and the same for the ply stresses in material axes."""
    strain = float(np.max(np.abs(check.kernel_strain - check.strain)) / np.max(np.abs(check.strain)))
    stress = float(np.max(np.abs(check.kernel_ply_stress_12 - check.ply_stress_12)) / np.max(np.abs(check.ply_stress_12)))
    return strain, stress


def end_effect(check: LaminateCheck) -> tuple[float, float, float]:
    """Three measures of the grips' hold on the width: the free length in widths; the lateral section force at the centre as a share
    of the longitudinal one, zero with free lateral edges and more where the hold still reaches the centre; and the reaction's excess
    over free-edge laminate theory, the stiffening of the held ends, which under displacement control shifts strain from the ends to
    the free length even where N_x at the centre has faded."""
    n_x, n_y = check.kernel_section_force[0], check.kernel_section_force[1]
    restraint = abs(n_x / n_y) if n_y != 0.0 else float("nan")
    excess = check.kernel_reaction_N / check.reaction_N - 1.0 if check.reaction_N != 0.0 else float("nan")
    return check.free_length_m / check.width_m, restraint, excess


def _laminate(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
    check = facts.laminate
    if check is None:
        return "Not checked", Found(difference=("No laminate-theory check for this model: " + facts.laminate_reason,))
    strain, stress = laminate_distance(check)
    widths, restraint, excess = end_effect(check)
    eps = check.kernel_strain
    ours = (f"Centre element {check.centre_element}: eps_x {eps[0]:.4g}, eps_y {eps[1]:.4g}, gamma_xy {eps[2]:.4g}",
            "Ply stresses in material axes: on the Results sheet")
    lines = [f"Strain: within {_pct(strain)} of classical laminate theory with free lateral edges",
             f"Ply stresses: within {_pct(stress)} of laminate theory",
             f"Free length {widths:.2g} widths; N_x at the centre {_pct(restraint)} of N_y (zero with free edges: the hold the grips keep on the width)",
             f"Reaction {_pct(abs(excess))} {'above' if excess >= 0 else 'below'} free-edge theory: the held ends stiffen the coupon"]
    if item.condition is not None and widths < item.condition.min_free_length_widths:
        return "N/A", Found(ours=ours, difference=(f"The free length is {widths:.2g} widths, under the {item.condition.min_free_length_widths:g} the hand "
                                                   "check needs: the grips' hold on the width reaches the gauge centre", "For the record:", *lines))
    return ("Pass" if max(strain, stress) <= item.tolerance.relative else "Fail"), Found(ours=ours, difference=(*lines, f"Tolerance {_pct(item.tolerance.relative)}"))


def _equilibrium(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
    deck = ev.our_deck
    applied = np.zeros(2)
    for step in deck.steps:
        for load in step.cloads:
            if load.dof in (1, 2):
                applied[load.dof - 1] += load.value * len(resolve_nodes(deck, load.target, load.line))
    residual = ev.solution.total_reaction + applied
    scale = abs(facts.reaction_N) if facts.reaction_N is not None else float(np.abs(ev.solution.reaction).sum(axis=0).max()) / 2
    if scale == 0.0:
        return "Not checked", Found(difference=("No reaction to compare the sum with",))
    share = float(np.max(np.abs(residual))) / scale
    return (("Pass" if share <= item.tolerance.relative else "Fail"),
            Found(ours=(f"Reactions sum to ({ev.solution.total_reaction[0]:.3g}, {ev.solution.total_reaction[1]:.3g}) N",
                        f"Applied loads ({applied[0]:.3g}, {applied[1]:.3g}) N"),
                  difference=(f"Residual {share:.2g} of the {scale:.6g} N reaction; tolerance {item.tolerance.relative:g}",)))


def _not_built(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
    return "Not checked", Found(difference=("This check is not built yet",))


def _rebuilt(item: Item, ev: Evidence, facts: _Facts) -> tuple[str | None, Found]:
    if ev.rebuilt is None:
        return "Not checked", Found(team=("No deck rebuilt from the guide has come back",))
    if isinstance(ev.rebuilt, str):
        return "Fail", Found(team=("The deck rebuilt from the guide",), difference=(f"Cannot be matched to ours: {ev.rebuilt}",))
    differ = ev.rebuilt.differences
    if not differ:
        return "Pass", Found(team=("The deck rebuilt from the guide",), ours=("Our deck",), difference=("Matches item by item",))
    return "Fail", Found(team=("The deck rebuilt from the guide",), ours=("Our deck",),
                         difference=("Differs:", *(f"{i.item}: rebuilt {i.team}; ours {i.ours}" for i in differ)))


AUTO: dict[str, Callable[[Item, Evidence, _Facts], tuple[str | None, Found]]] = {
    "A1": _card, "A2": _units,
    **{item_id: _model_match for item_id in EQUIVALENCE_GROUPS},
    "C1": _reaction, "C2": _stresses(("S11", "S22", "S12")), "C3": _stresses(("Mises",)), "C4": _displacement,
    "C7": _fpf, "D1": _laminate, "D2": _equilibrium, "D3": _not_built, "D4": _not_built, "E2": _rebuilt,
}


def _listed(ev: Evidence, name: str) -> str | None:
    """The team's text of an equivalence item that is listed, not judged."""
    return next((i.team for i in ev.equivalence.items if i.item == name), None) if ev.equivalence else None


def _auditor_evidence(item: Item, ev: Evidence, facts: _Facts) -> Found:
    """What our tools can show the auditor for an item the auditor judges."""
    if item.id == "B6":
        strengths = next((i.ours for i in ev.equivalence.items if i.item.startswith("Strengths")), None) if ev.equivalence else None
        team = tuple(f"{label}: {text}" for label, text in (("Material names", _listed(ev, "Material names")),
                                                           ("Strengths", _listed(ev, "Strengths (*Fail Stress)"))) if text)
        ours = (f"Material: {ev.material}" if ev.material else "Material: explicit constants in the spec",
                *((f"Strengths: {strengths}",) if strengths else ()))
        return Found(team, ours)
    if item.id == "C5":
        if ev.equivalence is None:
            return Found(team=(NOT_BUILT_BY_TEAM,))
        return Found(ours=tuple((ev.excluded or "The comparison names no region left out").split("\n")))
    if item.id == "C6":
        if not ev.paired:
            return Found(team=(NOT_BUILT_BY_TEAM if ev.equivalence is None else "No field report in the export",))
        team, ours = [], []
        for q in dict.fromkeys(p.quantity for p in ev.paired):
            members = [p for p in ev.paired if p.quantity == q]
            t, o = max(members, key=lambda p: abs(p.reference)), max(members, key=lambda p: abs(p.ours))
            team.append(f"{q}: largest at element {t.reference_element} ({t.x:g}, {t.y:g})")
            ours.append(f"{q}: largest at element {o.reference_element} ({o.x:g}, {o.y:g})")
        return Found(tuple(team), tuple(ours), (f"Element numbers are the team's; (x, y) in {facts.length_unit}",))
    if item.id == "D5":
        centre = ()
        if facts.centre is not None:
            e, name = facts.centre
            eps = ev.solution.strain[e]
            centre = (f"At the centre element {name}: gamma_xy = {eps[2]:.4g}, eps_y = {eps[1]:.4g}",)
        return Found(ours=(f"Layup {_layup_text(facts.angles, full=True)}", *centre, "Our engine's model is a membrane, so it shows no bending"))
    if item.id == "E3":
        return Found() if ev.guide is None else Found(ours=(f"The guide lists {ev.guide[0]} open item(s) at its end and tags {ev.guide[1]} value(s) ASSUMED",))
    if item.id == "E4":
        card = dict(ev.card)
        return Found(ours=(f"Test data on the coupon card: {card.get('Test data', 'not stated')}",))
    return Found()


def _facts(evidence: Evidence) -> _Facts:
    deck, solution = evidence.our_deck, evidence.solution
    try:
        check, reason = laminate_check(deck, solution), ""
    except KernelError as exc:
        check, reason = None, str(exc)
    try:
        _, top = gripped_groups(deck, solution)
        reaction: float | None = float(solution.reaction[list(top)].sum(axis=0)[1])
    except KernelError:
        reaction = None
    mesh = solution.mesh
    instance_rows = [k for k in range(mesh.node_count) if mesh.node_keys[k][0] is not None]
    middle = 0.5 * (mesh.xy[instance_rows].min(axis=0) + mesh.xy[instance_rows].max(axis=0))
    centres = mesh.xy[mesh.connectivity[:, :4]].mean(axis=1)
    e = int(np.argmin(np.linalg.norm(centres - middle, axis=1)))
    length_unit, stress_unit = {"mm-MPa": ("mm", "MPa"), "m-Pa": ("m", "Pa")}[evidence.units]
    return _Facts(check, reason, _deck_angles(deck), length_unit, stress_unit, reaction, (e, element_name(mesh.element_keys[e])))


def score(rubric: Rubric, evidence: Evidence) -> tuple[Scored, ...]:
    """Every item that applies to the package's test, in rubric order: a result where our tools judge, evidence for the auditor. An item
    our tools would judge but whose team value is not in the export states our value and is the auditor's (result None)."""
    facts = _facts(evidence)
    scored = []
    for item in items_for(rubric, evidence.test_code):
        if item.method == "auditor":
            scored.append(Scored(item, None, _auditor_evidence(item, evidence, facts)))
            continue
        scorer = AUTO.get(item.id)
        result, found = scorer(item, evidence, facts) if scorer else ("Not checked", Found(difference=("No scorer is built for this item yet",)))
        scored.append(Scored(item, result, found))
    return tuple(scored)


def tolerance_text(item: Item, units: str) -> str:
    if item.tolerance is None:
        return ""
    if item.tolerance.relative is not None:
        relative = item.tolerance.relative
        return f"{relative:g} relative" if relative < 1e-4 else f"{100 * relative:g} %"
    unit = {"mm-MPa": "mm", "m-Pa": "m"}[units]
    return f"{from_si(item.tolerance.length_m, 'length', units):g} {unit}"


# --------------------------------------------------------------------------------------------- the audit form

COVER, RUBRIC_SHEET, RECORD = "Cover", "Rubric", "Record"  # read back from a returned form: the first two sheets, in this order, and Record by name
RESULTS_SHEET, VALUES_SHEET = "Results", "Compared values"  # the numbers behind our rows, for the reader; not read back
AUDITOR_FIELDS = ("Auditor", "Audit date")  # the cover's cells the auditor fills in
RULE_FIELD = "Verdict"  # the cover's formula: the verdict the rule gives for the results and observations as they stand; blank until every required row has one
DIGEST = "Our results digest"  # the Record row that lets a returned form show our tools' results were not changed
RUBRIC_COLUMNS = ("ID", "Check", "What passes", "Judged by", "Required", "Tolerance", "Result", "Team", "Ours", "Difference", "Observation", "Cause")
FOUND_COLUMNS = ("Team", "Ours", "Difference")  # what our tools found: the team's values, ours, and the difference between them
AUDITOR_COLUMNS = ("Observation", "Cause")  # open on every row; Result is open on the rows the auditor judges
RUBRIC_ROWS = 200  # how far down the rubric sheet the cover's formula looks: more rows than any rubric has
COMPARED = "What is compared"  # the cover's block: the values the audit sets side by side, and the files each side comes from
AUDITOR_BLOCK = "Filled in by the auditor"
SECTIONS = "The sections"  # the cover's list: each section and what it means
HOW_TO = "How to fill it in"
GLANCE = "At a glance"  # the Results sheet's first block: the headline numbers
TO_SEND = "For the next delivery, the team's export should also include"  # the cover's list of what the rows waiting on the team need


@dataclass(frozen=True)
class Table:
    """One table of our numbers on the Results sheet: a title, a header row, rows of values, a note under them and the number format
    of each column (None for General; a shorter tuple leaves the rest General)."""

    title: str
    columns: tuple[str, ...]
    rows: tuple[tuple, ...]
    note: str = ""
    formats: tuple[str | None, ...] = ()
    check: str = ""  # "How to check": the file and the operation that reproduce the table's numbers, for the team


@dataclass(frozen=True)
class FormTables:
    """The numbers behind the results, computed by our tools and written to the form as values (ADR-0001): the headline lines at a
    glance, the tables of the Results sheet, every compared value for the Compared values sheet, and what the team's next export
    should include for the rows that wait on it."""

    glance: tuple[tuple[str, str], ...]
    tables: tuple[Table, ...]
    values: Table | None
    to_send: tuple[str, ...]


def form_tables(rubric: Rubric, evidence: Evidence, scored: tuple[Scored, ...]) -> FormTables:
    """The form's numbers, from the evidence the rows were scored from: the band agreement per quantity and layer, the gauge centre
    against laminate theory, the reaction, and every compared value."""
    ev = evidence
    length_unit, stress_unit = {"mm-MPa": ("mm", "MPa"), "m-Pa": ("m", "Pa")}[ev.units]
    items = {item.id: item for item in rubric.items}
    glance: list[tuple[str, str]] = []
    tables: list[Table] = []
    values: Table | None = None
    stress, length, percent = "0.00", "0.000", "0.000"
    deck_name = Path(ev.our_deck.source_path).name if ev.our_deck.source_path else "our deck"
    reports = (f"the team's own reports ({', '.join(ev.team_reports)}) give" if len(ev.team_reports) > 1 else
               f"the team's own report ({ev.team_reports[0]}) gives" if ev.team_reports else "the team's own field reports give")

    if ev.paired and ev.inner_band is not None:
        lo, hi = ev.inner_band
        quantities = list(dict.fromkeys(p.quantity for p in ev.paired))
        in_band = {q: [p for p in ev.paired if p.quantity == q and lo <= p.y <= hi] for q in quantities}
        largest_of = {q: max((abs(p.reference) for p in in_band[q]), default=0.0) for q in quantities}
        rows: list[tuple] = []
        for q in quantities:
            tolerance = items["C3" if q == "Mises" else "C2"].tolerance.relative
            largest = largest_of[q]
            for layer in [*sorted({p.layer for p in in_band[q]}), None]:
                group = in_band[q] if layer is None else [p for p in in_band[q] if p.layer == layer]
                if not group:
                    continue
                worst = max(group, key=lambda p: abs(p.difference))
                within = sum(abs(p.difference) <= tolerance * largest for p in group)
                share = abs(worst.difference) / largest if largest else None
                rows.append((q, "all" if layer is None else layer, len(group), within, 100 * tolerance, largest, abs(worst.difference),
                             None if share is None else 100 * share, worst.reference_element, worst.x, worst.y))
                if layer is None:
                    glance.append((f"{q} in the graded band", f"{within} of {len(group)} within {_pct(tolerance)} of {largest:.6g} {stress_unit}; worst "
                                   f"{_pct(share) if share is not None else 'n/a'} at team element {worst.reference_element}, layer {worst.layer}"))
        tables.append(Table(f"Stresses in the graded band, y {lo:g} to {hi:g} {length_unit}: the team's Abaqus against our engine",
                            ("Quantity", "Layer (1 = bottom)", "Values compared", "Within tolerance", "Tolerance (% of the largest team value)",
                             f"Largest team value ({stress_unit})", f"Worst difference ({stress_unit})", "Worst difference (% of the largest)",
                             "Team element of the worst", f"x ({length_unit})", f"y ({length_unit})"), tuple(rows),
                            "A value is within tolerance when |ours - team| is at most the tolerance times the largest team value of that quantity in the "
                            "band (rubric rows C2 and C3). The same numbers are in 4-results/team-vs-ours.csv.",
                            (None, None, None, None, "0.0", stress, stress, percent, None, length, length),
                            f"How to check: in 4-results/team-vs-ours.csv keep the rows of one quantity with y between {lo:g} and {hi:g} {length_unit}; the "
                            f"largest |team| value, the largest |ours - team| and its element are the numbers above. {reports[0].upper() + reports[1:]} "
                            "the same largest value over the rows with Y in that range."))
        values = Table("Every compared value: the team's reported value beside ours at the same element and layer",
                       ("Team element", "Our element", f"x ({length_unit})", f"y ({length_unit})", "Layer (1 = bottom)", "Quantity", f"Team ({stress_unit})",
                        f"Ours ({stress_unit})", f"Ours - team ({stress_unit})", "Share of the largest team value in the graded band", "In the graded band"),
                       tuple((p.reference_element, p.our_element, p.x, p.y, p.layer, p.quantity, p.reference, p.ours, p.difference,
                              abs(p.difference) / largest_of[p.quantity] if largest_of[p.quantity] else None, "yes" if lo <= p.y <= hi else "no")
                             for p in ev.paired),
                       formats=(None, None, length, length, None, None, stress, stress, stress, "0.0000", None))

    try:
        check: LaminateCheck | None = laminate_check(ev.our_deck, ev.solution)
    except KernelError:
        check = None
    hand_check = any(s.item.id == "D1" for s in scored)  # the gauge centre of an open-hole coupon is at the hole: no laminate theory there
    top_set = None
    try:
        _, top = gripped_groups(ev.our_deck, ev.solution)
        reference = tuple(ev.solution.mesh.node_keys[i] for i in top if ev.solution.mesh.node_keys[i][0] is None)
        top_set = next((name for name, keys in ev.our_deck.nsets.items() if reference and tuple(keys) == reference), None)
    except KernelError:
        pass
    reaction_check = (f"How to check: in Abaqus on our deck {deck_name}, the reaction force RF2 at the top reference point"
                      + (f" (set {top_set})" if top_set else "") + " at the end of the step")
    if check is not None and not hand_check:
        tables.append(Table("Reaction at the loaded grip (N)", ("Our engine", "Team"),
                            ((check.kernel_reaction_N, ev.team_reaction if ev.team_reaction is not None else "not exported"),),
                            "Rubric rows C1 and D2. The laminate-theory hand check (row D1) does not apply to this test.", ("0.0", "0.0"),
                            reaction_check + "."))
        glance.append(("Reaction at the loaded grip", f"ours {check.kernel_reaction_N:.6g} N; " + (
            f"team {ev.team_reaction:.6g} N, {_pct(abs(check.kernel_reaction_N - ev.team_reaction) / abs(ev.team_reaction))} apart"
            if ev.team_reaction is not None else "not in the team's export")))
    if check is not None and hand_check:
        widths, restraint, excess = end_effect(check)
        strain, stress_distance = laminate_distance(check)
        per_length = 1.0 / from_si(1.0, "length", ev.units)  # N/m to N per stated length unit
        strain_scale = float(np.max(np.abs(check.strain)))
        force_scale = float(np.max(np.abs(check.section_force)))
        condition = items["D1"].condition if "D1" in items else None
        if condition is None:
            applies_note = applies_short = ""
        elif widths >= condition.min_free_length_widths:
            applies_note, applies_short = f" Rubric row D1 applies from {condition.min_free_length_widths:g} widths.", ""
        else:
            applies_note = f" Rubric row D1 applies from {condition.min_free_length_widths:g} widths, so for this coupon it is N/A."
            applies_short = f"; row D1 is N/A under {condition.min_free_length_widths:g} widths"
        label = check.centre_element.rsplit(".", 1)[-1]  # the element's label in the deck, as Abaqus/CAE asks for it
        width = from_si(check.width_m, "length", ev.units)
        if check.control == "displacement":
            applied = (f"eps_y = {check.applied / check.free_length_m:.4g} (the {from_si(check.applied, 'length', ev.units):.4g} {length_unit} pull over "
                       f"the {from_si(check.free_length_m, 'length', ev.units):.4g} {length_unit} free length)")
        else:
            applied = f"N_y = {check.applied / check.width_m * per_length:.6g} N/{length_unit} (the {check.applied:.6g} N load over the {width:g} {length_unit} width)"
        ply0 = next((f"Ply-{k + 1}" for k, ply in enumerate(check.plies) if ply.angle_deg % 180 == 0), None)
        where = (f"section point of {ply0} (0 deg), whose axes are the coupon's" if ply0
                 else "any ply, remembering that Abaqus reports them in the ply's own axes")
        team_element = next((p.reference_element for p in ev.paired if str(p.our_element) == label), None)
        tables.append(Table(f"Strains at the gauge centre, element {check.centre_element}: laminate theory with free lateral edges against our engine",
                            ("Strain", "Laminate theory", "Our engine", "Difference (% of the largest theory strain)"),
                            tuple((name, float(theory), float(ours), 100 * abs(float(ours) - float(theory)) / strain_scale)
                                  for name, theory, ours in zip(("eps_x", "eps_y", "gamma_xy"), check.strain, check.kernel_strain)),
                            f"Strains are dimensionless; gamma_xy is the engineering shear strain. The largest difference is {_pct(strain)} of the largest "
                            f"theory strain (rubric row D1).{applies_note}", (None, "0.000000", "0.000000", percent),
                            f"How to check: laminate theory with any laminate calculator, for the layup and material on the coupon card, at {applied} and "
                            f"N_x = N_xy = 0. Our engine, in Abaqus on our deck {deck_name}: XY Data from ODB field output, LE11, LE22 and LE12 at the "
                            f"integration point of element {label}, {where}; LE12 is the engineering shear strain."))
        tables.append(Table(f"Section forces at the gauge centre, element {check.centre_element} (N/{length_unit}): laminate theory against our engine",
                            ("Section force", "Laminate theory", "Our engine", "Difference (% of the largest theory force)"),
                            tuple((name, float(theory) * per_length, float(ours) * per_length, 100 * abs(float(ours) - float(theory)) / force_scale)
                                  for name, theory, ours in zip(("N_x", "N_y", "N_xy"), check.section_force, check.kernel_section_force)),
                            f"Section forces are per unit width. The free length between the grips is {from_si(check.free_length_m, 'length', ev.units):.4g} "
                            f"{length_unit}, {widths:.2g} widths. N_x at the centre is {_pct(restraint)} of N_y: the hold the grips keep on the width there, "
                            f"zero with free edges. The reaction is {_pct(abs(excess))} {'above' if excess >= 0 else 'below'} free-edge theory: the held ends "
                            "stiffen the coupon, and under displacement control the free length then carries a little more strain than the pull over the "
                            "free length.", (None, length, length, percent),
                            "How to check: laminate theory is the laminate stiffness A of the card's layup and material times the strains above; our "
                            f"engine's N_y times the width ({width:g} {length_unit}) is our reaction below."))
        ply_rows = []
        for k, ply in enumerate(check.plies):
            theory = [float(v) for v in from_si(check.ply_stress_12[k], "stress", ev.units)]
            ours = [float(v) for v in from_si(check.kernel_ply_stress_12[k], "stress", ev.units)]
            ply_rows.append((k + 1, ply.angle_deg, theory[0], ours[0], theory[1], ours[1], theory[2], ours[2]))
        tables.append(Table(f"Ply stresses at the gauge centre in material axes ({stress_unit}): laminate theory against our engine",
                            ("Ply (1 = bottom)", "Angle (deg)", "S11 theory", "S11 ours", "S22 theory", "S22 ours", "S12 theory", "S12 ours"), tuple(ply_rows),
                            f"The largest difference is {_pct(stress_distance)} of the largest theory value (rubric row D1).",
                            (None, "0", stress, stress, stress, stress, stress, stress),
                            f"How to check: in Abaqus on our deck {deck_name}, S11, S22 and S12 at element {label}, one section point per ply, in the ply's "
                            "material axes." + (f" The team's own reports hold S11 and S22 at the same element, their element {team_element}, at layers 1 "
                                                "and 10, to set beside the 'ours' columns." if team_element is not None else "")))
        if ev.team_reaction is not None:
            team = f"team {ev.team_reaction:.6g} N, {_pct(abs(check.kernel_reaction_N - ev.team_reaction) / abs(ev.team_reaction))} apart"
        else:
            team = "not in the team's export"
        tables.append(Table("Reaction at the loaded grip (N)", ("Laminate theory", "Our engine", "Team"),
                            ((check.reaction_N, check.kernel_reaction_N, ev.team_reaction if ev.team_reaction is not None else "not exported"),),
                            "Laminate theory is the free-edge laminate under the same pull or load (rubric rows C1 and D2).", ("0.0", "0.0", "0.0"),
                            reaction_check + f"; laminate theory is N_y times the width ({width:g} {length_unit})."))
        glance.append(("Reaction at the loaded grip", f"ours {check.kernel_reaction_N:.6g} N; {team}"))
        glance.append(("Hand check at the gauge centre", f"strain within {_pct(strain)} and ply stresses within {_pct(stress_distance)} of laminate theory "
                       f"with free edges; the free length is {widths:.2g} widths{applies_short}"))

    to_send = tuple(f"{s.item.id} ({s.item.check}): {s.item.needs}" for s in scored if s.result in (None, "Not checked") and s.item.needs)
    return FormTables(tuple(glance), tuple(tables), values, to_send)


def write_form(path: Path | str, rubric: Rubric, scored: tuple[Scored, ...], cover: list[tuple[str, str]], units: str,
               tables: FormTables | None = None, files: list[tuple[str, str]] | tuple = (), record: list[tuple[str, str]] | tuple = ()) -> Path:
    """The audit form, version 0.3. The cover says what is audited (`cover`: the coupon, its test, the attempt, the package) and what is
    compared (`files`: the values, and the files each side comes from), takes the auditor's name and date, shows the verdict the rule
    gives (a formula that fills itself from the Results), lists the sections with their meaning, and gives the instructions. The
    Rubric sheet has one row per check under a banner per section: our results locked, the auditor's Result, Observation and Cause
    open, and what our tools found in three columns, the team's values, ours and the difference, one fact per line. With `tables`,
    a Results sheet of the numbers behind our rows (the headline lines first) and a Compared values sheet of every value compared.
    The Record sheet, last and hidden, holds how the package was built (`record`: the team deck and its digest, the engine's commit,
    who built it and when), the rubric version and the digest of our results; it is for our tools, and the auditor never sees it.
    Lists offer the fixed words, results are coloured, and every sheet prints on letter paper. The sheets are protected without a
    password, so a cell is not changed by accident (this is not a security measure)."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter

    path = Path(path)
    bold, title = Font(name=FONT, bold=True, size=10), Font(name=FONT, bold=True, size=14)
    subtitle_font, heading = Font(name=FONT, size=11, color="595959"), Font(name=FONT, bold=True, size=11)
    wrap = Alignment(wrap_text=True, vertical="top")
    open_cell = PatternFill("solid", fgColor="FFF4CC")  # the auditor's cells
    header_fill = PatternFill("solid", fgColor="DDE6F0")
    banner_fill = PatternFill("solid", fgColor="EDEDED")  # a section's row on the Rubric sheet
    unlocked = Protection(locked=False)
    shades = {"Pass": ("C6EFCE", "006100"), "Pass with observations": ("FFEB9C", "9C5700"), "Fail": ("FFC7CE", "9C0006"),
              "N/A": ("EDEDED", "595959"), "Not checked": ("EDEDED", "595959")}
    length_unit, stress_unit = {"mm-MPa": ("mm", "MPa"), "m-Pa": ("m", "Pa")}[units]
    letter = {name: get_column_letter(i + 1) for i, name in enumerate(RUBRIC_COLUMNS)}
    cover_widths = {1: 26, 2: 104}

    def shade(sheet, cells: str, words) -> None:
        """A result or a verdict shows its colour wherever it is written, by the auditor too."""
        for word in words:
            fill, ink = shades[word]
            sheet.conditional_formatting.add(cells, CellIsRule(operator="equal", formula=[f'"{word}"'], fill=PatternFill("solid", fgColor=fill),
                                                               font=Font(name=FONT, color=ink)))

    def rubric_range(name: str) -> str:
        return f"{RUBRIC_SHEET}!${letter[name]}$2:${letter[name]}${RUBRIC_ROWS}"

    def rule_formula() -> str:
        """verdict_of, as a formula: Fail if a required row fails; Pass with observations if any other row fails or an observation is
        written; Pass otherwise; blank while a required row has no result other than Not checked."""
        required, result, observation = rubric_range("Required"), rubric_range("Result"), rubric_range("Observation")
        missing = f'COUNTIFS({required},"yes",{result},"Not checked")+COUNTIFS({required},"yes",{result},"")'
        required_fail = f'COUNTIFS({required},"yes",{result},"Fail")'
        any_fail, observations = f'COUNTIF({result},"Fail")', f'COUNTIF({observation},"?*")'
        return f'=IF({missing}>0,"",IF({required_fail}>0,"Fail",IF(OR({any_fail}>0,{observations}>0),"Pass with observations","Pass")))'

    book = Workbook()
    book.properties.creator = "Anisotropy AI"
    facts = dict(cover)
    book.properties.title = f"Coupon audit: {facts.get('Package', path.stem)}"
    sheet = book.active
    sheet.title = COVER

    def pair(row: int, label: str, value, label_font=bold):
        """A label in the first column and its value in the second; a label that starts with spaces is a row under the one above, indented."""
        head = sheet.cell(row, 1, label.strip())
        head.font = Font(name=FONT, size=10) if label.startswith(" ") else label_font
        if label.startswith(" "):
            head.alignment = Alignment(indent=2, vertical="top")
        cell = sheet.cell(row, 2, value)
        cell.alignment = wrap
        return cell

    def paragraph(row: int, text: str) -> None:
        """One text across both columns."""
        cell = sheet.cell(row, 1, text)
        cell.alignment = wrap
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        sheet.row_dimensions[row].height = 13.0 * lines_of(text, cover_widths[1] + cover_widths[2]) + 4

    sheet["A1"] = "Anisotropy AI coupon audit"
    sheet["A1"].font = title
    subtitle = ", ".join(t for t in (facts.get("Coupon"), facts.get("Test"), f"attempt {facts['Attempt']}" if facts.get("Attempt") else None) if t)
    sheet.cell(2, 1, subtitle).font = subtitle_font
    row = 4
    for label, value in cover:
        pair(row, label, value)
        row += 1
    if files:
        row += 1
        sheet.cell(row, 1, COMPARED).font = heading
        row += 1
        for label, value in files:
            pair(row, label, value)
            row += 1
    row += 1
    sheet.cell(row, 1, AUDITOR_BLOCK).font = heading
    row += 1
    for label in AUDITOR_FIELDS:
        cell = pair(row, label, None)
        cell.fill, cell.protection = open_cell, unlocked
        row += 1
    rule = pair(row, RULE_FIELD, rule_formula())
    shade(sheet, rule.coordinate, VERDICTS)
    row += 2

    sheet.cell(row, 1, SECTIONS).font = heading
    row += 1
    for section in rubric.sections:
        if any(s.item.section == section.id for s in scored):
            pair(row, f"{section.id}. {section.title}", section.meaning or None)
            row += 1
    row += 1

    sheet.cell(row, 1, HOW_TO).font = heading
    row += 1
    numbers = ("The Results sheet holds the numbers behind our tools' rows"
               + (", and Compared values every value compared, with a filter; both are also in the package's 4-results folder."
                  if tables is not None and tables.values is not None else "; our ply stresses are also in the package's 4-results folder.")
               if tables is not None and tables.tables else "Our ply stresses are in the package's 4-results folder.")
    instructions = [
        "On the Rubric sheet each row is one check: what passes, the team's values, ours, and the difference between them, one fact per line. The "
        "rows our tools judged are filled in and locked. Give a Result in every yellow cell: compare the team's value with ours and pick Pass, Fail, "
        "N/A or Not checked from the list. Where the Team column says 'left blank', enter the team's value in Observation first.",
        "To pass, every row marked Required must be Pass or N/A; the Verdict above then fills itself. For a row that fails, pick its Cause: engine, "
        "our input, their model, documentation or tolerance. Write an Observation where you have one.",
        numbers,
        "Put the filled-in form back in the folder you received the package in.",
    ]
    for line in instructions:
        paragraph(row, line)
        row += 1
    if tables is not None and tables.to_send:
        row += 1
        sheet.cell(row, 1, TO_SEND).font = heading
        row += 1
        for line in tables.to_send:
            paragraph(row, line)
            row += 1
    for j, width in cover_widths.items():
        sheet.column_dimensions[get_column_letter(j)].width = width
    fit_rows(sheet, {2: cover_widths[2]}, 4, row)
    protect(sheet)
    printable(sheet, landscape=False)

    sheet = book.create_sheet(RUBRIC_SHEET)
    sheet.append(list(RUBRIC_COLUMNS))
    for cell in sheet[1]:
        cell.font, cell.fill, cell.alignment = bold, header_fill, wrap
    result_col, observation_col, cause_col = (RUBRIC_COLUMNS.index(name) + 1 for name in ("Result", "Observation", "Cause"))
    auditor_rows, item_rows = [], []
    r = 2
    for section in rubric.sections:
        rows = [s for s in scored if s.item.section == section.id]
        if not rows:
            continue
        for col in range(1, len(RUBRIC_COLUMNS) + 1):
            sheet.cell(r, col).fill = banner_fill
        banner = sheet.cell(r, 2, f"{section.id}. {section.title}")
        banner.font, banner.alignment = bold, wrap
        sheet.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(RUBRIC_COLUMNS))
        r += 1
        for s in rows:
            item = s.item
            values = [item.id, item.check, item.acceptance, "our tools" if s.result is not None else "auditor", "yes" if item.required else "no",
                      tolerance_text(item, units) or None, s.result, "\n".join(s.found.team) or None, "\n".join(s.found.ours) or None,
                      "\n".join(s.found.difference) or None, None, None]
            for col, value in enumerate(values, start=1):
                sheet.cell(r, col, value).alignment = wrap
            sheet.cell(r, 1).font = bold
            for col in (observation_col, cause_col) + ((result_col,) if s.result is None else ()):
                cell = sheet.cell(r, col)
                cell.fill, cell.protection = open_cell, unlocked
            item_rows.append(r)
            if s.result is None:
                auditor_rows.append(r)
            r += 1
    last = r - 1
    if auditor_rows:
        listing(sheet, RESULTS, [f"{letter['Result']}{r}" for r in auditor_rows])
    listing(sheet, CAUSES, [f"{letter['Cause']}{r}" for r in item_rows])
    shade(sheet, f"{letter['Result']}2:{letter['Result']}{last}", RESULTS)
    widths = (6, 24, 34, 10, 9, 10, 12, 40, 40, 34, 28, 12)
    for name, width in zip(RUBRIC_COLUMNS, widths):
        sheet.column_dimensions[letter[name]].width = width
    fit_rows(sheet, {i + 1: w for i, w in enumerate(widths)}, 2, last)
    sheet.freeze_panes = "C2"
    protect(sheet)
    printable(sheet, landscape=True, title_rows="1:1")

    if tables is not None and (tables.tables or tables.glance):
        sheet = book.create_sheet(RESULTS_SHEET)
        sheet["A1"] = "The numbers behind our tools' rows"
        sheet["A1"].font = title
        sheet["A2"] = (f"Computed by our engine while the package was built, from the same evidence as the rubric's rows. Lengths in {length_unit}, "
                       f"stresses in {stress_unit}, forces in N. These are values, not formulas.")
        sheet["A2"].alignment = wrap
        sheet.merge_cells("A2:H2")
        sheet.row_dimensions[2].height = 30
        row, columns = 4, 8
        if tables.glance:
            sheet.cell(row, 1, GLANCE).font = heading
            row += 1
            for label, value in tables.glance:
                sheet.cell(row, 1, label).font = bold
                sheet.cell(row, 2, value).alignment = wrap
                sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
                row += 1
            row += 1
        for t in tables.tables:
            sheet.cell(row, 1, t.title).font = heading
            row += 1
            for j, name in enumerate(t.columns, start=1):
                cell = sheet.cell(row, j, name)
                cell.font, cell.fill, cell.alignment = bold, header_fill, wrap
            sheet.row_dimensions[row].height = 44
            row += 1
            for values_row in t.rows:
                for j, value in enumerate(values_row, start=1):
                    cell = sheet.cell(row, j, value)
                    number_format = t.formats[j - 1] if j - 1 < len(t.formats) else None
                    if number_format and isinstance(value, (int, float)) and not isinstance(value, bool):
                        cell.number_format = number_format
                row += 1
            for text, font in ((t.note, None), (t.check, Font(name=FONT, size=10, italic=True, color="595959"))):
                if not text:
                    continue
                cell = sheet.cell(row, 1, text)
                cell.alignment = wrap
                if font is not None:
                    cell.font = font
                sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(4, len(t.columns)))
                sheet.row_dimensions[row].height = 14.0 * max(1, -(-len(text) // 120)) + 4
                row += 1
            row += 1
            columns = max(columns, len(t.columns))
        sheet.column_dimensions["A"].width = 30
        for j in range(2, columns + 1):
            sheet.column_dimensions[get_column_letter(j)].width = 16
        protect(sheet)
        printable(sheet, landscape=True)

    if tables is not None and tables.values is not None:
        sheet = book.create_sheet(VALUES_SHEET)
        sheet.append(list(tables.values.columns))
        for cell in sheet[1]:
            cell.font, cell.fill, cell.alignment = bold, header_fill, wrap
        sheet.row_dimensions[1].height = 44
        for values_row in tables.values.rows:
            sheet.append(list(values_row))
        for j, number_format in enumerate(tables.values.formats, start=1):
            if number_format:
                for cells in sheet.iter_rows(min_row=2, min_col=j, max_col=j):
                    cells[0].number_format = number_format
        for j in range(1, len(tables.values.columns) + 1):
            sheet.column_dimensions[get_column_letter(j)].width = 14
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(tables.values.columns))}{len(tables.values.rows) + 1}"
        printable(sheet, landscape=True, title_rows="1:1")  # left open, so the filter and sorting work; the digest covers the results, not this sheet

    sheet = book.create_sheet(RECORD)
    sheet["A1"] = "Record"
    sheet["A1"].font = title
    sheet["A2"] = "For our tools: how this package was built. This sheet is hidden on purpose; nothing here is filled in."
    sheet["A2"].alignment = wrap
    sheet.merge_cells("A2:B2")
    sheet.row_dimensions[2].height = 30
    row = 4
    digest = results_digest([(s.item.id, s.result) for s in scored if s.result is not None])
    for label, value in [*record, ("Rubric", f"version {rubric.version} ({rubric.status})"), (DIGEST, digest)]:
        sheet.cell(row, 1, label).font = bold
        sheet.cell(row, 2, value).alignment = wrap
        row += 1
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 90
    fit_rows(sheet, {2: 90}, 4, row)
    protect(sheet)
    printable(sheet, landscape=False)
    sheet.sheet_state = "hidden"

    one_font(book)

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def results_digest(results: list[tuple[str, str]]) -> str:
    """SHA-256 of our tools' results, item by item, so a returned form shows they were not changed (the sheet lock has no password)."""
    return hashlib.sha256(";".join(f"{item}={result}" for item, result in results).encode("utf-8")).hexdigest()


# --------------------------------------------------------------------------------------------- a returned form

class AuditError(ValueError):
    """A returned audit form that is incomplete, inconsistent, or not the package's."""


@dataclass(frozen=True)
class Row:
    id: str
    judged_by: str  # "our tools" or "auditor"
    required: bool
    result: str
    observation: str  # the auditor's, free text
    cause: str  # one of CAUSES when the row fails; what the auditor traced the failure to, and so who fixes it


@dataclass(frozen=True)
class ReturnedForm:
    path: Path
    cover: dict[str, str]  # label -> value, as the form's Cover and Record sheets state them (the cover's word wins)
    rows: tuple[Row, ...]
    auditor: str
    date: str  # ISO
    verdict: str  # the rule's, from the rows; the cover's formula shows the same word once the spreadsheet has recalculated it

    def counts(self) -> str:
        """The failed rows and the observations, in words."""
        failed = [r.id for r in self.rows if r.result == "Fail"]
        observed = sum(bool(r.observation) for r in self.rows)
        parts = [f"{len(failed)} fail ({', '.join(failed)})"] if failed else []
        if observed:
            parts.append(f"{observed} observation" + ("s" if observed != 1 else ""))
        return "; ".join(parts) or "none"


def verdict_of(rows: tuple[Row, ...]) -> str:
    """The rule's verdict: Fail if a required row fails; Pass with observations if any other row fails or an observation is written;
    Pass otherwise. Refused while a required row has no result other than Not checked."""
    missing = [r.id for r in rows if r.required and r.result in ("", "Not checked")]
    if missing:
        raise AuditError(f"no verdict yet: the required rows {', '.join(missing)} have no result other than Not checked")
    if any(r.required and r.result == "Fail" for r in rows):
        return "Fail"
    if any(r.result == "Fail" or r.observation for r in rows):
        return "Pass with observations"
    return "Pass"


def read_form(path: Path | str) -> ReturnedForm:
    """A filled-in audit form, checked: every problem is named at once (ADR-0003)."""
    from openpyxl import load_workbook

    path = Path(path)
    try:
        book = load_workbook(path, data_only=True)
    except Exception as exc:  # openpyxl raises several kinds for a file that is not a workbook
        raise AuditError(f"{path.name}: not an audit form ({exc})") from None
    problems: list[str] = []
    if book.sheetnames[:2] != [COVER, RUBRIC_SHEET] or RECORD not in book.sheetnames:
        raise AuditError(f"{path.name}: the sheets are {book.sheetnames}, not {COVER} and {RUBRIC_SHEET} first with a {RECORD} sheet")

    cover = {**_labelled(book[RECORD]), **_labelled(book[COVER])}
    for label in AUDITOR_FIELDS:
        if not cover.get(label):
            problems.append(f"the cover's {label} is empty")
    verdict = cover.get(RULE_FIELD, "")  # the formula's value as the spreadsheet saved it; empty when the file was never recalculated
    if verdict and verdict not in VERDICTS:
        problems.append(f"the verdict {verdict!r} is not one of {list(VERDICTS)}")
    try:
        audit_date = date.fromisoformat(cover.get("Audit date", "")).isoformat()
    except ValueError:
        audit_date = cover.get("Audit date", "")
        if audit_date:
            problems.append(f"the audit date {audit_date!r} is not a date (write YYYY-MM-DD)")

    sheet = book[RUBRIC_SHEET]
    header = [_text(c.value) for c in sheet[1]]
    if header != list(RUBRIC_COLUMNS):
        raise AuditError(f"{path.name}: the rubric sheet's columns are {header}, not {list(RUBRIC_COLUMNS)}")
    col = {name: RUBRIC_COLUMNS.index(name) + 1 for name in RUBRIC_COLUMNS}
    rows = []
    for r in range(2, sheet.max_row + 1):
        item = _text(sheet.cell(r, col["ID"]).value)
        if not ITEM_ID.match(item):  # a section's banner, or an empty row
            continue
        result = _text(sheet.cell(r, col["Result"]).value)
        if result not in RESULTS:
            problems.append(f"row {item}: the result {result!r} is not one of {list(RESULTS)}" if result else f"row {item} has no result")
        cause = _text(sheet.cell(r, col["Cause"]).value)
        if cause and cause not in CAUSES:
            problems.append(f"row {item}: the cause {cause!r} is not one of {list(CAUSES)}")
        elif result == "Fail" and not cause:
            problems.append(f"row {item} fails without a cause (one of {list(CAUSES)})")
        rows.append(Row(item, _text(sheet.cell(r, col["Judged by"]).value), _text(sheet.cell(r, col["Required"]).value) == "yes", result,
                        _text(sheet.cell(r, col["Observation"]).value), cause))
    digest = results_digest([(row.id, row.result) for row in rows if row.judged_by == "our tools"])
    if cover.get(DIGEST) != digest:
        problems.append("a result our tools gave was changed: the rows judged by our tools no longer match the digest on the Record sheet")

    if not problems:
        try:
            rule = verdict_of(tuple(rows))
        except AuditError as exc:
            problems.append(str(exc))
        else:
            if verdict and verdict != rule:
                problems.append(f"the verdict on the form is {verdict!r}; the rule gives {rule!r} for these results and observations")
            verdict = rule
    if problems:
        raise AuditError(f"{path.name}: " + "; ".join(problems))
    return ReturnedForm(path, cover, tuple(rows), cover["Auditor"], audit_date, verdict)


def our_results(path: Path | str) -> dict[str, str]:
    """The results our tools gave on a form, by item, filled in or not: what a later build of the same coupon is compared with."""
    from openpyxl import load_workbook

    sheet = load_workbook(path, data_only=True)[RUBRIC_SHEET]
    col = {name: RUBRIC_COLUMNS.index(name) + 1 for name in RUBRIC_COLUMNS}
    return {_text(sheet.cell(r, col["ID"]).value): _text(sheet.cell(r, col["Result"]).value) for r in range(2, sheet.max_row + 1)
            if _text(sheet.cell(r, col["Judged by"]).value) == "our tools"}


NOT_REGRESSIONS: tuple[str, ...] = ()  # rows about how a package was built rather than what the engine computes; none since rubric 0.3


def regressions(before: dict[str, str], now: dict[str, str]) -> list[str]:
    """The items our tools passed before and do not pass now; an item the rubric no longer has is not one."""
    return [item for item, result in before.items() if result == "Pass" and item in now and now[item] != "Pass" and item not in NOT_REGRESSIONS]
