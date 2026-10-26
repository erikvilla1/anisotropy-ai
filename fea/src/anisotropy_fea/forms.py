"""The house style of the Excel forms the team fills in: the audit form (`audit.write_form`) and the coupon input form
(`intake.write_blank_form`). One font, yellow cells for what the team enters, lists for the fixed words, sheets protected
without a password (against accidents, not a security measure), row heights fitted to the text, letter paper, and a way to
read a filled-in sheet back as label and value. Nothing here knows what a form is about.
"""

from __future__ import annotations

from datetime import date, datetime
from typing import Sequence

from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

FONT = "Arial"
OPEN_FILL = "FFF4CC"  # the team's cells
HEADER_FILL = "DDE6F0"  # a table's header row
BANNER_FILL = "EDEDED"  # a block's banner row
MUTED = "595959"


def font(size: int = 10, bold: bool = False, italic: bool = False, muted: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=MUTED if muted else None)


def wrap(indent: int = 0) -> Alignment:
    return Alignment(wrap_text=True, vertical="top", indent=indent)


def open_cell(cell) -> None:
    """A cell the team fills in: yellow and unlocked."""
    cell.fill = PatternFill("solid", fgColor=OPEN_FILL)
    cell.protection = Protection(locked=False)


def fill(cell, colour: str) -> None:
    cell.fill = PatternFill("solid", fgColor=colour)


def listing(sheet, words: Sequence[str] | str, cells: Sequence[str]) -> None:
    """A list to pick from: the fixed words, or a range reference such as "'Library materials'!$A$2:$A$21"."""
    formula = words if isinstance(words, str) else '"' + ",".join(words) + '"'
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    sheet.add_data_validation(validation)
    for cell in cells:
        validation.add(cell)


def number_rule(sheet, cells: Sequence[str], whole: bool = False, at_least: float | None = None, above: float | None = None) -> None:
    """A hint from Excel that a cell takes a number (whole or decimal), optionally bounded; the reader does the real refusal."""
    kind = "whole" if whole else "decimal"
    if above is not None:
        validation = DataValidation(type=kind, operator="greaterThan", formula1=str(above), allow_blank=True)
    elif at_least is not None:
        validation = DataValidation(type=kind, operator="greaterThanOrEqual", formula1=str(at_least), allow_blank=True)
    else:
        validation = DataValidation(type=kind, operator="greaterThan", formula1="-1E+300", allow_blank=True)
    sheet.add_data_validation(validation)
    for cell in cells:
        validation.add(cell)


def protect(sheet) -> None:
    sheet.protection.sheet = True
    sheet.protection.formatColumns = False  # widths and heights may still be changed
    sheet.protection.formatRows = False


def printable(sheet, landscape: bool, title_rows: str | None = None) -> None:
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0  # one page wide, as many pages long as it takes
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.6)
    sheet.oddFooter.center.text = "&A, page &P of &N"
    if title_rows:
        sheet.print_title_rows = title_rows


def lines_of(text: str, width: float) -> int:
    """How many lines a text wraps to in a column of this width, by the character count Excel does not fit for us."""
    per_line = max(1, int(width * 1.15))
    return sum(max(1, -(-len(part) // per_line)) for part in text.split("\n"))


def fit_rows(sheet, widths: dict[int, float], first: int, last: int) -> None:
    """Row heights from the longest wrapped text in each row: Excel does not fit them itself for a file it did not write."""
    for r in range(first, last + 1):
        lines = 1
        for col, width in widths.items():
            text = sheet.cell(r, col).value
            if isinstance(text, str) and text and not text.startswith("="):
                lines = max(lines, lines_of(text, width))
        if lines > 1:
            sheet.row_dimensions[r].height = min(409, 13.0 * lines + 4)


def paragraph(sheet, row: int, text: str, last_column: int, width: float, font_: Font | None = None) -> None:
    """One text across the columns 1 to `last_column`, its row fitted to the text."""
    cell = sheet.cell(row, 1, text)
    cell.alignment = wrap()
    if font_ is not None:
        cell.font = font_
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    sheet.row_dimensions[row].height = 13.0 * lines_of(text, width) + 4


def one_font(book) -> None:
    """Every written cell in the house font, keeping its size, weight, slant and colour."""
    for sheet in book.worksheets:
        for cells in sheet.iter_rows():
            for cell in cells:
                if cell.value is not None:
                    f = cell.font
                    cell.font = Font(name=FONT, size=f.sz or 10, bold=f.b, italic=f.i, color=f.color)


def text_of(value) -> str:
    """A cell's value as the text a reader compares: dates in ISO form, numbers as Excel shows them, blanks empty."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def labelled(sheet, label_column: int = 1, value_column: int = 2) -> dict[str, str]:
    """A sheet of labels in one column and values in the next, as a dict of texts."""
    return {text_of(sheet.cell(r, label_column).value): text_of(sheet.cell(r, value_column).value)
            for r in range(1, sheet.max_row + 1) if sheet.cell(r, label_column).value}
